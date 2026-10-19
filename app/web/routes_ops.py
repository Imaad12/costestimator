"""Operational pages: Approvals inbox, Rates workbook (download template, upload, preview, import) and the Settings hub (users, audit, integrity, standards, documents, sources, reference data).

All writes go through `/ui/do/<action>` or the existing import services; nothing here approves or edits data by itself.
"""
from __future__ import annotations

import hashlib
import json
import tempfile
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, Request, UploadFile
from fastapi.responses import HTMLResponse, Response

from ..api.deps import current_actor
from ..api.routes import save_upload
from ..config import get_settings
from ..core.errors import AppError
from ..database import get_db, rows_to_dicts
from ..excel import templates as xl_templates
from ..models.actor import Actor
from ..models.enums import ROLE_VALUES, Role
from ..rate_updates import candidates as rate_candidates
from ..rate_updates import importer as rate_importer
from ..services import audit, auth, commercial, standards
from . import inbox
from .routes import back, page
from .routes_tender import ROOT, _DOC_GROUPS

router = APIRouter()
USER_ROLES = [r for r in ROLE_VALUES if r != Role.AI.value]
SETTINGS_TABS = (("general", "General"), ("users", "Users & roles"), ("audit", "Audit trail"), ("integrity", "Integrity & hashes"), ("standards", "Standards & evidence"),
                 ("documents", "Rules & documents"), ("sources", "Data sources"), ("reference", "Reference data"))


# ------------------------------------------------------------------ plain-language help (the (i) icons)
@router.get("/glossary.json")
def glossary(actor: Actor = Depends(current_actor)):
    from fastapi.responses import JSONResponse
    from . import helptext
    return JSONResponse({"terms": helptext.TERMS, "placeholders": {k: list(v) for k, v in helptext.PLACEHOLDERS.items()}}, headers={"Cache-Control": "private, max-age=600"})


# ------------------------------------------------------------------ approvals inbox
@router.get("/approvals", response_class=HTMLResponse)
def approvals_page(request: Request, group: str = "", entity: str = "", log: int = 0, actor: Actor = Depends(current_actor)):
    db = get_db()
    groups = inbox.build(db, actor, only=group)
    allc = {g["key"]: len(g["items"]) for g in (groups if not group else inbox.build(db, actor))}
    rows = []
    if log or entity:
        rows = rows_to_dicts(db.all("SELECT * FROM approvals " + ("WHERE entity_type=? " if entity else "") + "ORDER BY id DESC LIMIT 300", (entity,) if entity else ()))
    names = {k: n for n, k, _ in inbox.GROUPS}
    return page(request, "approvals.html", actor, groups=groups, tabs=[(k, names[k], c) for k, c in allc.items() if c or k == group], group=group, total=sum(allc.values()),
                log=rows, show_log=bool(log or entity), entity=entity, title="Approvals")


@router.post("/approvals/bulk")
async def approvals_bulk(request: Request, actor: Actor = Depends(current_actor)):
    from .routes import run as run_action
    f = await request.form()
    group, decision, comment = f.get("group", ""), f.get("decision", ""), (f.get("comment") or "").strip()
    nxt = "/approvals" + (f"?group={group}" if group else "")
    sel = set(f.getlist("sel"))
    if decision not in ("approve", "reject"):
        return back(nxt, err="Choose Approve selected or Reject selected.")
    if not sel:
        return back(nxt, err="Select at least one item first.")
    if decision == "reject" and not comment:
        return back(nxt, err="Give a reason to reject the selected items.")
    r = inbox.decide_selected(get_db(), actor, group, sel, decision, comment, run_action)
    msg = f"{'Approved' if decision == 'approve' else 'Rejected'} {r['done']} item(s), each recorded under your name."
    if r["skipped"]:
        msg += f" Skipped {r['skipped']} (your role may not decide them, you raised them, or they need individual input)."
    if r["failed"]:
        msg += f" {r['failed']} refused by the rules: " + "; ".join(r["errors"]) + "."
    return back(nxt, ok=msg) if r["done"] else back(nxt, err=msg)


# ------------------------------------------------------------------ rates workbook
@router.get("/rates/workbook", response_class=HTMLResponse)
def rates_workbook(request: Request, actor: Actor = Depends(current_actor)):
    db = get_db()
    recent = rows_to_dicts(db.all("SELECT id, source_name, description, imported_at, imported_by, status, accepted_rows, rejected_rows, sha256 FROM import_manifests WHERE kind='rates' ORDER BY id DESC LIMIT 15"))
    return page(request, "rates_workbook.html", actor, source_types=rate_candidates.SOURCE_TYPES, register=rows_to_dicts(db.all("SELECT source_id, name FROM sources ORDER BY source_id")),
                recent=recent, counts=rate_candidates.queue_counts(db), preview=None, title="Rates workbook")


@router.get("/rates/workbook/template.xlsx")
def rates_template(actor: Actor = Depends(current_actor)):
    with tempfile.TemporaryDirectory() as d:
        p = xl_templates.rate_workbook(Path(d) / "rates_workbook.xlsx", get_db())
        data = p.read_bytes()
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="rates_workbook_template.xlsx"'})


def _rate_args(source_type: str, source_name: str, source_url: str, source_id: str, description: str) -> dict:
    return dict(source_id=source_id or None, source_name=source_name or None, source_url=source_url or None, source_type=source_type, description=description or "",
                extra_flags=["AWARD_NOT_MARKET_QUOTE"] if source_type == "Public tender award" else None)


def _form_args(f) -> dict:
    """Everything the check / import needs, read from the one form the mapping page posts."""
    g = lambda k, d="": (f.get(k) or d) if hasattr(f, "get") else d          # noqa: E731
    mapping = {fld: int(f[f"map_{fld}"]) for fld, _ in rate_importer.MAP_FIELDS if f.get(f"map_{fld}", "") != ""}
    return dict(stored=Path(g("stored")).name, sheet=g("sheet") or None, header_row=int(g("header_row")) if g("header_row") else None, mapping=mapping or None,
                default_currency=g("default_currency").strip() or None, default_date=g("default_date").strip() or None, default_tax_basis=g("default_tax_basis").strip() or None,
                src=dict(source_name=g("source_name"), source_url=g("source_url"), source_id=g("source_id"), source_type=g("source_type", "Manual dataset"), description=g("description")))


def _rates_page(request: Request, actor: Actor, a: dict, *, filename: str, dry: dict | None, info: dict | None, warn: str | None = None):
    db = get_db()
    recent = rows_to_dicts(db.all("SELECT id, source_name, description, imported_at, imported_by, status, accepted_rows, rejected_rows, sha256 FROM import_manifests WHERE kind='rates' ORDER BY id DESC LIMIT 15"))
    return page(request, "rates_workbook.html", actor, source_types=rate_candidates.SOURCE_TYPES, register=rows_to_dicts(db.all("SELECT source_id, name FROM sources ORDER BY source_id")),
                recent=recent, counts=rate_candidates.queue_counts(db), title="Rates workbook", map_fields=rate_importer.MAP_FIELDS, hint=warn or None,
                preview=dict(result=dry, info=info, stored=a["stored"], file=filename, mapping=a["mapping"] or (info or {}).get("auto") or {}, default_currency=a["default_currency"] or "",
                             default_date=a["default_date"] or "", default_tax_basis=a["default_tax_basis"] or "", **a["src"]))


def _run(db, actor, a: dict, dry_run: bool) -> dict:
    path = get_settings().projects_dir / "_uploads" / a["stored"]
    if not path.exists():
        raise AppError("The uploaded file is no longer available; please upload it again.", code="FILE_GONE")
    return rate_importer.import_rates_file(db, actor, path, dry_run=dry_run, sheet=a["sheet"], mapping=a["mapping"], header_row=a["header_row"],
                                           default_currency=a["default_currency"], default_date=a["default_date"], default_tax_basis=a["default_tax_basis"], **_rate_args(**a["src"]))


@router.post("/rates/workbook/preview", response_class=HTMLResponse)
async def rates_preview(request: Request, file: UploadFile = File(...), source_name: str = Form(""), source_url: str = Form(""), source_id: str = Form(""),
                        source_type: str = Form("Manual dataset"), description: str = Form(""), default_currency: str = Form(""), default_tax_basis: str = Form(""), actor: Actor = Depends(current_actor)):
    db = get_db()
    if not source_name.strip() and not source_id:
        return back("/rates/workbook", err="Say where these prices come from: enter a source name (for example 'Supplier X price list, Sep 2026') or pick a Source Register entry.")
    path = save_upload("", file)
    a = _form_args(dict(stored=path.name, source_name=source_name, source_url=source_url, source_id=source_id, source_type=source_type, description=description, default_currency=default_currency, default_tax_basis=default_tax_basis))
    try:
        info = None if path.suffix.lower() == ".json" else rate_importer.inspect_file(path)
        dry, warn = None, None
        if info is not None and not info["detected"]:
            auto = info["auto"]
            if "original_price" in auto and ({"resource_id", "description"} & set(auto)):          # use the keyword suggestion for the first check; the person confirms or corrects it
                a.update(mapping=dict(auto), header_row=info["header_row"], sheet=info["sheet"])
                warn = "The columns were not recognised exactly. The system has suggested what each column is: check it and correct anything wrong, then press Re-check."
            else:
                warn = "The columns were not recognised. Tell the system which column holds what, below."
        if info is None or info["detected"] or a["mapping"]:
            try:
                dry = _run(db, actor, a, True)
            except AppError as e:
                warn = f"{e.code}: {e.message}"
    except AppError as e:
        return back("/rates/workbook", err=f"{e.code}: {e.message}")
    return _rates_page(request, actor, a, filename=file.filename, dry=dry, info=info, warn=warn)


@router.post("/rates/workbook/remap", response_class=HTMLResponse)
async def rates_remap(request: Request, actor: Actor = Depends(current_actor)):
    f = await request.form()
    a = _form_args(f)
    path = get_settings().projects_dir / "_uploads" / a["stored"]
    try:
        info = rate_importer.inspect_file(path, a["sheet"], a["header_row"])
        dry, warn = None, None
        try:
            if a["mapping"] is None:
                a["mapping"] = info["auto"] or None
            dry = _run(get_db(), actor, a, True)
        except AppError as e:
            warn = f"{e.code}: {e.message}"
    except AppError as e:
        return back("/rates/workbook", err=f"{e.code}: {e.message}")
    info["auto"] = a["mapping"] or info["auto"]
    return _rates_page(request, actor, a, filename=a["stored"], dry=dry, info=info, warn=warn)


@router.post("/rates/workbook/arranged")
async def rates_arranged(request: Request, actor: Actor = Depends(current_actor)):
    """Download the person's own file re-arranged into the standard template (using the mapping they chose), so it can be tidied and re-uploaded without mapping next time."""
    a = _form_args(await request.form())
    path = get_settings().projects_dir / "_uploads" / a["stored"]
    try:
        rows = rate_importer.read_rows(path, a["sheet"], a["mapping"] or rate_importer.inspect_file(path, a["sheet"], a["header_row"])["auto"] or None, a["header_row"])
    except AppError as e:
        return back("/rates/workbook", err=f"{e.code}: {e.message}")
    from openpyxl import load_workbook
    cols = {"resource_id": "Resource ID", "description": "Description (for new resources)", "unit": "Unit", "original_price": "Unit price", "currency": "Currency", "effective_date": "Price date",
            "valid_until": "Valid until", "tax_basis": "Tax basis", "delivery_basis": "Delivery basis", "evidence_snippet": "Evidence / reference"}
    with tempfile.TemporaryDirectory() as d:
        p = xl_templates.rate_workbook(Path(d) / "arranged.xlsx", get_db())
        wb = load_workbook(p)
        ws = wb["Rate Dataset"]
        hdr = {str(c.value): c.column for c in ws[5] if c.value}
        for i, (_, r) in enumerate(rows):
            for f, h in cols.items():
                v = r.get(f)
                if f == "currency" and not v:
                    v = a["default_currency"]
                if f == "tax_basis" and not v:
                    v = a["default_tax_basis"]
                if f == "effective_date" and not v:
                    v = a["default_date"]
                if v not in (None, "") and h in hdr:
                    ws.cell(row=6 + i, column=hdr[h], value=v)
        wb.save(p)
        data = p.read_bytes()
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="rates_arranged_into_template.xlsx"'})


@router.post("/rates/workbook/import")
async def rates_import(request: Request, actor: Actor = Depends(current_actor)):
    a = _form_args(await request.form())
    try:
        r = _run(get_db(), actor, a, False)
    except AppError as e:
        return back("/rates/workbook", err=f"{e.code}: {e.message}")
    return back("/approvals?group=candidates", ok=f"{r['accepted']} price row(s) are now candidates awaiting review; {r['rejected']} rejected (manifest #{r['manifest_id']}, SHA-256 {r['sha256'][:12]}...). "
                "Approve them for publication here, publish, then approve the published rate.")


# ------------------------------------------------------------------ settings hub
def _file_sha(p: Path) -> str:
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _evidence_check(verify: bool) -> list[dict]:
    mf = ROOT / "data" / "standards" / "evidence" / "evidence_manifest.json"
    if not mf.exists():
        return []
    out = []
    for f in json.loads(mf.read_text(encoding="utf-8")).get("files", []):
        p = mf.parent / f["file"]
        state = "not checked"
        if verify:
            state = "file not present on this machine" if not p.exists() else ("MATCH" if _file_sha(p) == f["sha256"] else "MISMATCH")
        out.append(dict(file=f["file"], register_row=f.get("register_row"), sha256=f["sha256"], state=state))
    return out


def _manifest_check(db, verify: bool) -> list[dict]:
    rows = rows_to_dicts(db.all("SELECT id, kind, source_name, sha256, stored_path, imported_at, imported_by, status FROM import_manifests WHERE dry_run=0 ORDER BY id DESC LIMIT 100"))
    for r in rows:
        r["state"] = "not checked"
        if verify and r["stored_path"]:
            p = Path(r["stored_path"])
            r["state"] = "file not present" if not p.exists() else ("MATCH" if r["sha256"] and _file_sha(p) == r["sha256"] else "MISMATCH")
    return rows


@router.get("/settings", response_class=HTMLResponse)
def settings_hub(request: Request, tab: str = "general", verify: int = 0, project: str = "", entity: str = "", doc: str = "", actor: Actor = Depends(current_actor)):
    db, s = get_db(), get_settings()
    tab = tab if tab in dict(SETTINGS_TABS) else "general"
    ctx: dict = dict(tab=tab, tabs=SETTINGS_TABS, title="Settings", is_admin=actor.role == Role.ADMIN.value)
    if tab == "general":
        ctx.update(company=s.company, ai_enabled=s.ai_enabled, environment=getattr(s, "environment", ""), db_path=str(s.db_path),
                   backups=sorted(p.name for p in s.backups_dir.glob("*.zip")) if s.backups_dir.exists() else [],
                   migrations=rows_to_dicts(db.all("SELECT version, applied_at FROM schema_migrations ORDER BY version")))
    elif tab == "users":
        ctx.update(users=auth.list_users(db), roles=USER_ROLES)
    elif tab == "audit":
        sql, par = "SELECT id, ts, actor, actor_role, project_id, entity_type, entity_id, action, comment FROM audit_log WHERE 1=1", []
        if project:
            sql, par = sql + " AND project_id=?", par + [project]
        if entity:
            sql, par = sql + " AND entity_type=?", par + [entity]
        ctx.update(rows=rows_to_dicts(db.all(sql + " ORDER BY id DESC LIMIT 300", par)), etypes=[r["entity_type"] for r in db.all("SELECT DISTINCT entity_type FROM audit_log ORDER BY entity_type")],
                   f_project=project, f_entity=entity, chain=audit.verify_chain(db) if verify else None)
    elif tab == "integrity":
        ctx.update(chain=audit.verify_chain(db), evidence=_evidence_check(bool(verify)), manifests=_manifest_check(db, bool(verify)), verified=bool(verify),
                   boq_imports=rows_to_dicts(db.all("SELECT id, project_id, source_file, sha256, imported_by, imported_at FROM boq_imports ORDER BY id DESC LIMIT 50")))
    elif tab == "standards":
        rows = standards.list_standards(db)
        ctx.update(rows=rows, counts={st: sum(1 for r in rows if r["verification_status"] == st) for st in standards.STATUSES}, todo=standards.acquisition_list(db))
    elif tab == "documents":
        from .routes_tender import render_markdown
        have = {p.stem for p in (ROOT / "docs").glob("*.md")}
        groups = [("Rules", ["DOS_AND_DONTS"])] + [(g, [n for n in names if n in have]) for g, names in _DOC_GROUPS]
        rest = sorted(have - {n for _, ns in groups for n in ns})
        groups.append(("All other documents", rest))
        body = None
        if doc and doc in have:
            body = render_markdown((ROOT / "docs" / f"{doc}.md").read_text(encoding="utf-8"))
        elif not doc:
            doc = "DOS_AND_DONTS"
            body = render_markdown((ROOT / "docs" / "DOS_AND_DONTS.md").read_text(encoding="utf-8"))
        ctx.update(groups=groups, doc=doc, body=body)
    elif tab == "sources":
        ctx.update(sources=rows_to_dicts(db.all("SELECT * FROM sources ORDER BY source_id")), datasets=rows_to_dicts(db.all("SELECT * FROM dataset_imports ORDER BY id DESC LIMIT 20")))
    elif tab == "reference":
        ctx.update(tax=commercial.list_tax_categories(db), fx=rows_to_dicts(db.all("SELECT * FROM fx_rates ORDER BY id DESC")),
                   cost_codes=rows_to_dicts(db.all("SELECT * FROM cost_codes ORDER BY code")))
    return page(request, "settings.html", actor, **ctx)

