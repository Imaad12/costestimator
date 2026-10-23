"""BOQ builder helpers: the pick-list (assemblies + catalogue resources, grouped by trade division) and the two BOQ files made from a person's selection.

Quantities are whatever the person typed - nothing here measures, infers or fills a quantity. A rate is shown only when an Approved rate with a price exists; otherwise the cell stays blank.
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..core.errors import ValidationFailed
from ..database import Database

TYPE_LABELS = {"asm": "Work item (assembly)", "material": "Material", "equipment": "Equipment", "subcontract": "Subcontract package", "plant": "Plant", "labour": "Labour"}


def catalogue(db: Database) -> list[dict]:
    """Divisions (A-L) each with their selectable lines. Key = assembly code or resource id."""
    divs = {r["code"]: dict(code=r["code"], name=r["name"], items=[]) for r in db.all("SELECT code, name FROM cost_codes WHERE level=1 ORDER BY code")}
    for a in db.all("SELECT code, description, unit, cost_code, is_synthetic FROM assemblies a WHERE project_id IS NULL AND status IN ('Approved','Draft') "
                    "AND version=(SELECT MAX(version) FROM assemblies b WHERE b.code=a.code AND b.project_id IS NULL AND b.status IN ('Approved','Draft')) ORDER BY cost_code, code"):
        divs.get((a["cost_code"] or "?")[:1], divs.setdefault("?", dict(code="?", name="Other", items=[])))["items"].append(
            dict(key=a["code"], kind="asm", description=a["description"], unit=a["unit"], cost_code=a["cost_code"], rate=None, synthetic=bool(a["is_synthetic"])))
    for r in db.all("SELECT r.id, r.description, r.unit, r.cost_code, r.resource_type, "
                    "(SELECT unit_price||'|'||currency||'|'||is_synthetic FROM resource_rates x WHERE x.resource_id=r.id AND x.status='Approved' AND x.unit_price IS NOT NULL "
                    " ORDER BY x.version DESC LIMIT 1) AS best FROM resources r WHERE r.is_active=1 ORDER BY r.cost_code, r.id"):
        rate = None
        if r["best"]:
            price, cur, syn = r["best"].split("|")
            rate = dict(price=price, currency=cur, synthetic=syn == "1")
        divs.get((r["cost_code"] or "?")[:1], divs.setdefault("?", dict(code="?", name="Other", items=[])))["items"].append(
            dict(key=r["id"], kind=r["resource_type"], description=r["description"], unit=r["unit"], cost_code=r["cost_code"], rate=rate, synthetic=bool(rate and rate["synthetic"])))
    return [d for d in divs.values() if d["items"]]


def parse_selection(form, db: Database) -> list[dict]:
    """The lines the person gave a quantity to, in catalogue order. A quantity must be a number above zero; anything else is refused with the item named."""
    wanted: dict[str, Decimal] = {}
    for k, v in form.multi_items() if hasattr(form, "multi_items") else form.items():
        if k.startswith("q_") and str(v).strip():
            try:
                q = Decimal(str(v).strip().replace(",", ""))
            except InvalidOperation:
                raise ValidationFailed(f"Quantity '{v}' for {k[2:]} is not a number.", code="BAD_QUANTITY")
            if q <= 0:
                raise ValidationFailed(f"Quantity for {k[2:]} must be greater than zero (leave it empty to leave the item out).", code="BAD_QUANTITY")
            wanted[k[2:]] = q
    out = [dict(item, quantity=wanted[item["key"]], division=d) for d in catalogue(db) for item in d["items"] if item["key"] in wanted]
    for it in out:
        it["division_name"] = f'{it["division"]["code"]}. {it["division"]["name"]}'
    unknown = set(wanted) - {i["key"] for i in out}
    if unknown:
        raise ValidationFailed(f"Not in the catalogue: {', '.join(sorted(unknown))}.", code="UNKNOWN_ITEM")
    if not out:
        raise ValidationFailed("Pick at least one item and give it a quantity.", code="NOTHING_SELECTED")
    return out


def _sections(lines: list[dict]):
    cur, n = None, 0
    for ln in lines:
        if ln["division_name"] != cur:
            cur, n = ln["division_name"], 0
            yield "heading", cur, None
        n += 1
        yield "line", f'{ln["division"]["code"]}.{n:02d}', ln


def write_submission(path: Path, lines: list[dict], *, title: str, project: str | None, prepared_by: str) -> Path:
    """The BOQ document: headings per trade, item, description, unit, quantity, unit rate where an approved rate exists, and a formula amount."""
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ"
    synthetic = any(l["synthetic"] for l in lines)
    unpriced = sum(1 for l in lines if not l["rate"])
    ws["A1"] = title or "Bill of Quantities"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f'{("Project " + project + " - ") if project else ""}Prepared {date.today().isoformat()} by {prepared_by}. DRAFT for review - quantities are as entered by the preparer; the software has not measured or checked them.'
    ws["A3"] = ("CONTAINS SYNTHETIC / DEMO ITEMS OR RATES - NOT A MARKET QUOTE. " if synthetic else "") + (f"{unpriced} of {len(lines)} lines have no approved rate and are left blank: price them before this is submitted." if unpriced else "Every line carries an approved rate.")
    ws["A3"].font = Font(bold=True, color="9C5700" if (synthetic or unpriced) else "1E6B2E")
    for r in (2, 3):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 32
    hdr = 5
    for c, h in enumerate(["Item", "Description", "Unit", "Quantity", "Unit rate", "Currency", "Amount", "Reference"], 1):
        cell = ws.cell(hdr, c, h)
        cell.font, cell.fill = Font(bold=True), PatternFill("solid", fgColor="DDE3F5")
    n, first = hdr, hdr + 1
    for kind, a, ln in _sections(lines):
        n += 1
        if kind == "heading":
            ws.cell(n, 2, a).font = Font(bold=True)
            continue
        ws.cell(n, 1, a)
        ws.cell(n, 2, ln["description"]).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(n, 3, ln["unit"])
        ws.cell(n, 4, float(ln["quantity"])).number_format = "#,##0.00"
        if ln["rate"]:
            ws.cell(n, 5, float(ln["rate"]["price"])).number_format = "#,##0.00"
            ws.cell(n, 6, ln["rate"]["currency"])
            ws.cell(n, 7, f"=D{n}*E{n}").number_format = "#,##0.00"
        ws.cell(n, 8, ln["key"] + (" (SYNTHETIC)" if ln["synthetic"] else ""))
    n += 2
    ws.cell(n, 6, "Total (priced lines)").font = Font(bold=True)
    ws.cell(n, 7, f"=SUM(G{first}:G{n - 2})").number_format = "#,##0.00"
    ws.cell(n, 7).font = Font(bold=True)
    ws.freeze_panes = ws.cell(hdr + 1, 3)
    for c, w in enumerate([9, 62, 8, 12, 13, 9, 15, 26], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    wb.save(path)
    return path


def write_import_file(path: Path, lines: list[dict]) -> Path:
    """The same selection laid out for the BOQ importer (no rates: an imported 'rate' column would be read as the client's rate)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ"
    ws.append(["Item", "Description", "Unit", "Quantity", "Cost code", "Spec ref"])
    for kind, a, ln in _sections(lines):
        if kind == "heading":
            ws.append(["", a])
        else:
            ws.append([a, ln["description"], ln["unit"], float(ln["quantity"]), ln["cost_code"], ln["key"]])
    wb.save(path)
    return path
