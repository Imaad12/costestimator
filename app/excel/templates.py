"""Standalone workbook templates (templates/*.xlsx, *.docx). Same sheet specs as the master => identical headers & IDs."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from ..database import Database
from ..exports.docx_templates import generate_docx_templates
from .builder import write_lists_sheet, write_sheet
from .master import _text_sheet, build_master
from .spec import C, DATE, QTY, SAR, S, TAKEOFF_COLS, TEMPLATE_ROWS, Sheet

RATE_DATASET = Sheet("Rate Dataset", "Rate Dataset (for ingestion)",
                     "Upload rate datasets with full provenance. Rows are STAGED, validated (duplicates, units, currency, dates, conflicts, unusual changes) and published only as 'Pending Review' rates after explicit reviewer action. Rate types: see Lists.",
                     [C("resource_id", "Resource ID", 14), C("description", "Description (for new resources)", 40, wrap=True), C("unit", "Unit", 8, dv="units"),
                      C("unit_price", "Unit price", 14, fmt=SAR), C("currency", "Currency", 9), C("rate_type", "Rate type", 26, dv="RateType"),
                      C("region", "Region", 12, dv="regions"), C("price_date", "Price date", 13, fmt=DATE), C("valid_until", "Valid until", 13, fmt=DATE),
                      C("source_id", "Source ID (Source Register)", 14), C("supplier_id", "Supplier ID", 12), C("tax_basis", "Tax basis", 16),
                      C("delivery_basis", "Delivery basis", 18), C("evidence", "Evidence / reference", 34, wrap=True)], importable=True)
RESOURCE_TAKEOFF = Sheet("Resource Takeoff", "Pre-calculated resource quantities per BOQ item",
                         "Use when the measured take-off already includes resource quantities. Each row prices a BOQ item directly from resource quantity x approved rate (no assembly factors, no extra waste).",
                         [C("item_code", "Item", 12), C("resource_id", "Resource ID", 16), C("description", "Description", 36, wrap=True),
                          C("quantity", "Quantity", 14, fmt=QTY), C("unit", "Unit", 8, dv="units")], importable=True)
TAKEOFF = Sheet("Takeoff", "Measured Take-off (quantities measured by a human)",
                "Enter or paste the quantities you measured. The application imports them as-is and never alters or invents quantities.", TAKEOFF_COLS, importable=True)
RFQ = Sheet("RFQ", "Request for Quotation items",
            "Completed by the supplier (yellow). Return to the procurement team; import with Vendor Quotations.",
            [C("id", "Item", 18, "ro"), C("description", "Description", 46, "ro", wrap=True), C("specification", "Specification", 44, "ro", wrap=True),
             C("quantity", "Qty", 10, "ro", QTY), C("unit", "Unit", 8, "ro"), C("unit_rate", "Your unit rate", 14, fmt=SAR), C("lead_time_days", "Lead time (days)", 12),
             C("comments", "Remarks / deviations", 36, wrap=True)])


def _single(path: Path, db: Database | None, specs: list[Sheet], title: str, intro: str) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    _text_sheet(wb, "Instructions", [(title, "h"), (intro, ""),
                                     ("Yellow = input, grey = formula (locked), blue = read-only. Sheets are protected without a password to prevent accidents.", ""),
                                     ("Imports create DRAFT / PENDING proposals only. Approval statuses typed into a file are ignored; approvals happen in the application.", ""),
                                     ("SYNTHETIC / DEMO figures, where present, are labelled and are NOT market quotes.", "")], width=110)
    for s in specs:
        write_sheet(wb, s, [], meta="BLANK TEMPLATE", template_rows=TEMPLATE_ROWS)
    units = [r["code"] for r in db.all("SELECT code FROM units ORDER BY code")] if db else []
    regions = [r["code"] for r in db.all("SELECT code FROM regions ORDER BY code")] if db else []
    codes = [r["code"] for r in db.all("SELECT code FROM cost_codes ORDER BY code")] if db else []
    write_lists_sheet(wb, units, regions, codes)
    wb.calculation.fullCalcOnLoad = True
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


TEMPLATE_FILES = [
    "estimate_master.xlsx", "boq_import_template.xlsx", "takeoff_import_template.xlsx", "resource_rate_template.xlsx",
    "assembly_library_template.xlsx", "assumption_register_template.xlsx", "rfq_template.xlsx", "vendor_quote_template.xlsx",
    "quote_comparison_template.xlsx", "clarification_register_template.xlsx", "pricing_schedule_template.xlsx",
    "letter_of_offer_template.docx", "estimate_report_template.docx",
]


def generate_all(out_dir: Path, db: Database | None = None) -> list[Path]:
    out_dir = Path(out_dir)
    out: list[Path] = []
    out.append(build_master(out_dir / "estimate_master.xlsx", db, blank=True))
    out.append(_single(out_dir / "boq_import_template.xlsx", db, [S["BOQ Import"]], "BOQ Import Template",
                       "Paste the client BOQ into the yellow columns (column headers may also be mapped in the application's import dialog)."))
    out.append(_single(out_dir / "takeoff_import_template.xlsx", db, [TAKEOFF, RESOURCE_TAKEOFF], "Take-off Import Template",
                       "Sheet 'Takeoff' carries measured quantities; sheet 'Resource Takeoff' optionally carries pre-calculated resource quantities."))
    out.append(_single(out_dir / "resource_rate_template.xlsx", db, [RATE_DATASET, S["Material Library"]], "Resource & Rate Template",
                       "'Rate Dataset' is the ingestion format (with provenance). 'Material Library' shows the library layout; other libraries use the same columns."))
    out.append(_single(out_dir / "assembly_library_template.xlsx", db, [S["Assembly Library"]], "Assembly Library Template",
                       "One row per assembly line. Waste/productivity bases must be explicit; approved versions are immutable."))
    out.append(_single(out_dir / "assumption_register_template.xlsx", db, [S["Assumptions Register"]], "Assumption Register Template",
                       "Effective assumption values and proposed project overrides."))
    out.append(_single(out_dir / "rfq_template.xlsx", db, [RFQ], "RFQ Template", "Request-for-quotation response sheet for suppliers."))
    out.append(_single(out_dir / "vendor_quote_template.xlsx", db, [S["Vendor Quotations"]], "Vendor Quotation Template",
                       "One row per quotation line with header terms repeated on each row."))
    out.append(_single(out_dir / "quote_comparison_template.xlsx", db, [S["Quote Comparison"]], "Quote Comparison Template",
                       "Levelled comparison. Quoted prices and normalised values are always shown separately."))
    out.append(_single(out_dir / "clarification_register_template.xlsx", db, [S["Clarification Register"]], "Clarification Register Template",
                       "Raise clarifications OPEN; resolve only with the client response and evidence."))
    out.append(_single(out_dir / "pricing_schedule_template.xlsx", db, [S["Pricing Schedule"]], "Pricing Schedule Template",
                       "Quantities x unit rates with an engine check column."))
    out += generate_docx_templates(out_dir)
    return out


def rate_workbook(path: Path, db: Database | None = None) -> Path:
    """The rates-entry workbook: the shipped 'Rate Dataset' layout plus a read-only 'Resource List' (ids, descriptions, units - no prices) to copy resource ids from."""
    _single(path, db, [RATE_DATASET], "Rates workbook",
            "Fill the yellow 'Rate Dataset' rows: one price per row with currency, price date and an evidence reference. Use 'Resource List' to find resource ids. "
            "After upload every row is a candidate for review - nothing becomes a rate until a different person approves it.")
    if db:
        from openpyxl import load_workbook
        wb = load_workbook(path)
        ws = wb.create_sheet("Resource List")
        ws.append(["Resource ID", "Description", "Unit", "Type"])
        for r in db.all("SELECT id, description, unit, resource_type FROM resources WHERE is_active=1 ORDER BY id"):
            ws.append([r["id"], r["description"], r["unit"], r["resource_type"]])
        ws.freeze_panes = "A2"
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 18, 70
        wb.save(path)
    return path


PRICE_LIST_HEADERS = ["Resource ID", "Description (for new resources)", "Unit", "Unit price", "Currency", "Rate type", "Region", "Price date", "Valid until",
                      "Source ID (Source Register)", "Supplier ID", "Tax basis", "Delivery basis", "Evidence / reference"]
PRICE_LIST_REFERENCE = ["Trade", "Type", "Category", "Specification (read only)"]


def price_list_workbook(path: Path, db: Database) -> Path:
    """The price list, one sheet per trade division (A-L). One row per active catalogue resource with the price columns left BLANK for a person to fill.
    Every trade sheet has the 'Rate Dataset' headers, so the file uploads as-is; the upload reads every trade sheet, skips rows without a price, and treats the latest upload from a
    source as the source of truth, so the sheets can be filled in parts and uploaded again."""
    import re

    from openpyxl.styles import Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    start = wb.active
    start.title = "Start here"
    divs = {r["code"]: r["name"] for r in db.all("SELECT code, name FROM cost_codes WHERE level=1 ORDER BY code")}
    rows = db.all("SELECT r.id, r.description, r.unit, r.resource_type, r.cost_code, r.category, r.specification FROM resources r WHERE r.is_active=1 ORDER BY r.cost_code, r.resource_type, r.id")
    by_div: dict[str, list] = {}
    for r in rows:
        by_div.setdefault((r["cost_code"] or "?")[:1] or "?", []).append(r)
    yellow, grey, blue = PatternFill("solid", fgColor="FFF8D6"), PatternFill("solid", fgColor="F2F2F2"), PatternFill("solid", fgColor="DDE3F5")
    used: set[str] = set()
    index = []
    for code in sorted(by_div):
        name = divs.get(code, "Other")
        title = re.sub(r"[\[\]:*?/\\]", " ", f"{code} {name}")[:31].strip()
        while title in used:
            title = title[:28] + "..."
        used.add(title)
        ws = wb.create_sheet(title)
        intro = [f"Price list - {code}. {name}",
                 "Fill the yellow cells for the items you have prices for: Unit price + Currency (required, never assumed), then the price date, tax basis and where the price came from. Leave a row blank if you have no price - it stays missing.",
                 "Do not change Resource ID or the grey columns. Fill any sheets, in any order, in as many rounds as you like: upload the workbook each time with the same Source name - the latest upload is the source of truth for the items it prices."]
        for n, t in enumerate(intro, 1):
            ws.cell(n, 1, t).font = Font(bold=(n == 1), size=12 if n == 1 else 10)
            ws.merge_cells(start_row=n, start_column=1, end_row=n, end_column=10)
            ws.cell(n, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[2].height = ws.row_dimensions[3].height = 30
        hdr = 5
        for c, h in enumerate(PRICE_LIST_HEADERS + PRICE_LIST_REFERENCE, 1):
            cell = ws.cell(hdr, c, h)
            cell.font, cell.alignment = Font(bold=True), Alignment(wrap_text=True, vertical="center")
            cell.fill = blue if c <= len(PRICE_LIST_HEADERS) else PatternFill("solid", fgColor="E7E7E7")
        for n, r in enumerate(by_div[code], hdr + 1):
            vals = {1: r["id"], 2: r["description"], 3: r["unit"], 15: f'{r["cost_code"] or ""}', 16: r["resource_type"], 17: r["category"], 18: r["specification"]}
            for c in range(1, len(PRICE_LIST_HEADERS) + len(PRICE_LIST_REFERENCE) + 1):
                cell = ws.cell(n, c, vals.get(c))
                cell.fill = yellow if 4 <= c <= 14 else grey
            ws.cell(n, 4).number_format = "#,##0.00"
            ws.cell(n, 8).number_format = ws.cell(n, 9).number_format = "yyyy-mm-dd"
        last = hdr + len(by_div[code])
        ws.freeze_panes = ws.cell(hdr + 1, 3)
        ws.auto_filter.ref = f"A{hdr}:{get_column_letter(len(PRICE_LIST_HEADERS) + len(PRICE_LIST_REFERENCE))}{last}"
        for c, w in enumerate([14, 46, 8, 14, 10, 24, 12, 13, 13, 16, 12, 14, 18, 34, 12, 12, 18, 40], 1):
            ws.column_dimensions[get_column_letter(c)].width = w
        dv = DataValidation(type="list", formula1='"Included,Excluded"', allow_blank=True)
        dv2 = DataValidation(type="list", formula1='"Verified Supplier Quote,Published Market Price,Official Index or Benchmark,Historical Actual"', allow_blank=True)
        ws.add_data_validation(dv)
        ws.add_data_validation(dv2)
        dv.add(f"L{hdr + 1}:L{last}")
        dv2.add(f"F{hdr + 1}:F{last}")
        index.append((title, name, len(by_div[code])))
    lines = [("Price list - all trades", "h"),
             ("One sheet per trade. Open a trade, type the prices you have (Unit price + Currency), leave the rest blank. Yellow = you fill it, grey = reference, do not change.", ""),
             ("You can fill it in parts: save, upload, come back later, fill more sheets, upload again. Use the SAME Source name each time. For every item you price again the latest upload replaces the earlier price (the earlier one is kept in the history); items you leave blank keep whatever was uploaded before.", ""),
             ("Nothing becomes a rate on upload: each price is a candidate that a different person approves.", ""), ("", "")]
    for n, (t, k) in enumerate(lines, 1):
        start.cell(n, 1, t).font = Font(bold=(k == "h"), size=14 if k == "h" else 10)
        start.cell(n, 1).alignment = Alignment(wrap_text=True, vertical="top")
        start.merge_cells(start_row=n, start_column=1, end_row=n, end_column=4)
        start.row_dimensions[n].height = 20 if k == "h" else 34
    start.append(["Sheet", "Trade", "Items"])
    for c in start[len(lines) + 1]:
        c.font = Font(bold=True)
    for t, n, k in index:
        start.append([t, n, k])
    start.column_dimensions["A"].width, start.column_dimensions["B"].width, start.column_dimensions["C"].width = 32, 48, 10
    wb.save(path)
    return path
