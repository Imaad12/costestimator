"""Generic, spec-driven worksheet writer: headers, validations, formats, freeze panes, filters, protection, totals."""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from typing import Any, Iterable

from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from .spec import FIRST_ROW, HEADER_ROW, KEY_ROW, LISTS, TEMPLATE_ROWS, Col, Sheet

FILL_HDR = PatternFill("solid", fgColor="1F3864")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_CALC = PatternFill("solid", fgColor="E7E6E6")
FILL_RO = PatternFill("solid", fgColor="DDEBF7")
FILL_BAD = PatternFill("solid", fgColor="F8CBAD")
FILL_WARN = PatternFill("solid", fgColor="FFE699")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LEGEND = "Legend: yellow = input | grey = formula (locked) | blue = read-only from the application. Python application is the authoritative calculator."


def cellvalue(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, (datetime, date)):
        return v
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (list, dict)):
        import json
        return json.dumps(v, ensure_ascii=False)
    return v


def _typed(v: Any, col: Col) -> Any:
    """Database TEXT decimals/dates become real numbers/dates in numeric/date columns so Excel formulas work."""
    if not isinstance(v, str) or not col.fmt:
        return v
    from .spec import DATE
    if col.fmt == DATE:
        from ..core.timeutil import parse_date
        d = parse_date(v)
        return d if d else v
    if col.fmt != "@":
        try:
            return float(Decimal(v.replace(",", "")))
        except Exception:
            return v
    return v


def list_ref(name: str) -> str:
    return f"=Lists!${_list_col(name)}$2:${_list_col(name)}$60"


_LIST_ORDER = list(LISTS.keys()) + ["units", "regions", "costcodes"]


def _list_col(name: str) -> str:
    return get_column_letter(_LIST_ORDER.index(name) + 1)


def write_lists_sheet(wb, units: list[str], regions: list[str], costcodes: list[str]) -> Worksheet:
    ws = wb.create_sheet("Lists")
    data = dict(LISTS)
    data["units"], data["regions"], data["costcodes"] = units, regions, costcodes
    for ci, name in enumerate(_LIST_ORDER, 1):
        ws.cell(1, ci, name).font = Font(bold=True)
        for ri, v in enumerate(data[name], 2):
            ws.cell(ri, ci, v)
        ws.column_dimensions[get_column_letter(ci)].width = 22
    ws.sheet_state = "hidden"
    return ws


def _dv_for(col: Col, last_row: int) -> DataValidation | None:
    if not col.dv:
        return None
    if col.dv in ("units", "regions", "costcodes"):
        rng = f"=Lists!${_list_col(col.dv)}$2:${_list_col(col.dv)}$1000"
        dv = DataValidation(type="list", formula1=rng, allow_blank=True, showErrorMessage=True,
                            errorTitle="Invalid value", error="Pick a value from the list.")
        if col.dv == "costcodes":
            dv.showErrorMessage = False   # custom sub-codes may be added by administrators
    else:
        dv = DataValidation(type="list", formula1=list_ref(col.dv), allow_blank=True, showErrorMessage=True,
                            errorTitle="Invalid value", error="Pick a value from the list.")
    return dv


def write_sheet(wb, spec: Sheet, rows: Iterable[dict] | None, *, meta: str = "", template_rows: int = 0, name: str | None = None,
                extra_note: str = "") -> Worksheet:
    ws = wb.create_sheet(name or spec.name)
    ws["A1"] = spec.title
    ws["A1"].font = Font(bold=True, size=14, color="1F3864")
    ws["A2"] = spec.instructions + (" " + extra_note if extra_note else "")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A3"] = (meta + " | " if meta else "") + LEGEND
    ws["A3"].font = Font(italic=True, size=9, color="595959")
    ncols = len(spec.cols)
    last_c = get_column_letter(max(ncols, 6))
    ws.merge_cells(f"A2:{last_c}2")
    ws.row_dimensions[2].height = 48
    ws.row_dimensions[KEY_ROW].hidden = True
    for ci, col in enumerate(spec.cols, 1):
        k = ws.cell(KEY_ROW, ci, col.key)
        h = ws.cell(HEADER_ROW, ci, col.header)
        h.font = Font(bold=True, color="FFFFFF")
        h.fill = FILL_HDR
        h.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        h.border = BORDER
        ws.column_dimensions[get_column_letter(ci)].width = col.width
    ws.row_dimensions[HEADER_ROW].height = 34
    rows = list(rows or [])
    n = max(len(rows), template_rows)
    last = FIRST_ROW + n - 1 if n else FIRST_ROW
    for i in range(n):
        r = FIRST_ROW + i
        data = rows[i] if i < len(rows) else {}
        for ci, col in enumerate(spec.cols, 1):
            c = ws.cell(r, ci)
            if col.kind == "calc" and col.formula:
                c.value = col.formula.format(r=r)
            else:
                c.value = _typed(cellvalue(data.get(col.key)), col)
            c.border = BORDER
            if col.fmt:
                c.number_format = col.fmt
            if col.wrap:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            if col.kind == "input":
                c.fill = FILL_INPUT
                c.protection = Protection(locked=False)
            elif col.kind == "calc":
                c.fill = FILL_CALC
            else:
                c.fill = FILL_RO
    # totals row
    tot_row = last + 2
    if spec.totals:
        ws.cell(tot_row, 1, "TOTAL").font = Font(bold=True)
        for key, fn in spec.totals.items():
            ci = [c.key for c in spec.cols].index(key) + 1
            L = get_column_letter(ci)
            cell = ws.cell(tot_row, ci, f"={fn}({L}{FIRST_ROW}:{L}{last})")
            cell.font = Font(bold=True)
            cell.number_format = spec.cols[ci - 1].fmt or "General"
            cell.fill = FILL_CALC
    # data validations
    for ci, col in enumerate(spec.cols, 1):
        dv = _dv_for(col, last)
        if dv is not None:
            L = get_column_letter(ci)
            dv.add(f"{L}{FIRST_ROW}:{L}{max(last, FIRST_ROW + TEMPLATE_ROWS)}")
            ws.add_data_validation(dv)
        if col.kind == "calc" and col.key == "check":
            L = get_column_letter(ci)
            ws.conditional_formatting.add(f"{L}{FIRST_ROW}:{L}{last}", CellIsRule(operator="equal", formula=['"MISMATCH"'], fill=FILL_BAD))
        if col.key in ("status", "rate_status", "mapping_status", "approval_status", "validation_status"):
            L = get_column_letter(ci)
            ws.conditional_formatting.add(f"{L}{FIRST_ROW}:{L}{last}", FormulaRule(
                formula=[f'OR({L}{FIRST_ROW}="Pending Review",{L}{FIRST_ROW}="Draft",{L}{FIRST_ROW}="Needs Review",{L}{FIRST_ROW}="Suggested")'], fill=FILL_WARN))
            ws.conditional_formatting.add(f"{L}{FIRST_ROW}:{L}{last}", FormulaRule(
                formula=[f'OR({L}{FIRST_ROW}="Rejected",{L}{FIRST_ROW}="Expired")'], fill=FILL_BAD))
        if col.key == "severity":
            L = get_column_letter(ci)
            ws.conditional_formatting.add(f"{L}{FIRST_ROW}:{L}{last}", CellIsRule(operator="equal", formula=['"Critical"'], fill=FILL_BAD))
            ws.conditional_formatting.add(f"{L}{FIRST_ROW}:{L}{last}", CellIsRule(operator="equal", formula=['"High"'], fill=FILL_WARN))
    ws.freeze_panes = ws.cell(FIRST_ROW, 2)
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(ncols)}{max(last, FIRST_ROW)}"
    ws.first_data_row, ws.last_data_row, ws.total_row = FIRST_ROW, last, (tot_row if spec.totals else None)
    ws.sheet_properties.tabColor = "1F3864" if any(c.kind == "input" for c in spec.cols) else "7F7F7F"
    # protect formulas / read-only cells; input cells are unlocked. No password: protection prevents accidents, not tampering.
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.formatColumns = False
    return ws
