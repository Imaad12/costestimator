"""Formula recalculation & verification via LibreOffice (headless) on a COPY of the workbook.

openpyxl does not calculate formulas, so a freshly written workbook has no cached values. We never assume a recalculation
happened: `recalculate_copy` runs LibreOffice, and `verify_workbook` compares the recalculated formula results with the
Python engine's authoritative values (the 'Check' / 'Diff' columns and Project Summary).
"""
from __future__ import annotations

import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from .spec import FIRST_ROW, HEADER_ROW, KEY_ROW


def soffice_path() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


def recalculate_copy(src: Path, out_dir: Path | None = None, timeout: int = 180) -> Path:
    exe = soffice_path()
    if not exe:
        raise RuntimeError("LibreOffice (soffice) not found; cannot recalculate. Excel will recalculate on open (fullCalcOnLoad).")
    out_dir = Path(out_dir or tempfile.mkdtemp(prefix="ksa_recalc_"))
    out_dir.mkdir(parents=True, exist_ok=True)
    profile = Path(tempfile.mkdtemp(prefix="ksa_lo_profile_"))
    cmd = [exe, f"-env:UserInstallation=file://{profile}", "--headless", "--norestore", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
           "--outdir", str(out_dir), str(src)]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
    out = out_dir / Path(src).name
    if r.returncode != 0 or not out.exists():
        raise RuntimeError(f"LibreOffice recalculation failed: {r.stderr or r.stdout}")
    shutil.rmtree(profile, ignore_errors=True)
    return out


@dataclass
class VerifyReport:
    recalculated: bool
    errors: list[str] = field(default_factory=list)       # formula errors (#REF!, #NAME? ...)
    mismatches: list[str] = field(default_factory=list)   # Check columns saying MISMATCH
    formula_count: int = 0
    notes: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return self.recalculated and not self.errors and not self.mismatches


def verify_workbook(path: Path) -> VerifyReport:
    """Recalculate a copy and inspect every formula result."""
    try:
        rec = recalculate_copy(Path(path))
    except RuntimeError as e:
        return VerifyReport(False, notes=[str(e)])
    wb_v = openpyxl.load_workbook(rec, data_only=True)
    wb_f = openpyxl.load_workbook(path, data_only=False)
    rep = VerifyReport(True)
    for ws_f in wb_f.worksheets:
        ws_v = wb_v[ws_f.title]
        for row in ws_f.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    rep.formula_count += 1
                    v = ws_v[c.coordinate].value
                    if isinstance(v, str) and v.startswith("#"):
                        rep.errors.append(f"{ws_f.title}!{c.coordinate}: {v} ({c.value[:60]})")
                    elif v == "MISMATCH":
                        rep.mismatches.append(f"{ws_f.title}!{c.coordinate}")
    return rep
