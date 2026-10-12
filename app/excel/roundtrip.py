"""Re-reading edited workbooks: structure validation, formula verification, and a SAFE import (drafts/proposals only)."""
from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl

from ..approvals import engine as approvals
from ..core.errors import AppError, ValidationFailed
from ..core.money import D_opt
from ..database import Database, jload
from ..estimators.estimate import compute_estimate
from ..imports import boq as boq_import
from ..models.actor import Actor
from ..models.enums import RATE_TYPE_VALUES
from ..services import assemblies as asm_service
from ..services import assumptions as asm_svc
from ..services import audit, boq, indirects, rates, registers, suppliers
from .master import REQUIRED_SHEETS
from .recalc import VerifyReport, verify_workbook
from .spec import FIRST_ROW, HEADER_ROW, KEY_ROW, S


@dataclass
class StructureReport:
    missing_sheets: list[str] = field(default_factory=list)
    header_problems: list[str] = field(default_factory=list)
    overwritten_formulas: list[str] = field(default_factory=list)
    stale: bool = False
    fingerprint_in_file: str | None = None
    fingerprint_now: str | None = None

    @property
    def ok(self) -> bool:
        return not (self.missing_sheets or self.header_problems or self.overwritten_formulas or self.stale)


def read_rows(path: Path | openpyxl.Workbook, sheet: str) -> list[dict]:
    wb = path if isinstance(path, openpyxl.Workbook) else openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    keys = [c.value for c in ws[KEY_ROW]]
    rows = []
    for r in ws.iter_rows(min_row=FIRST_ROW, values_only=True):
        d = {k: v for k, v in zip(keys, r) if k}
        # skip empty and 'TOTAL' footer rows
        if all(v in (None, "") for v in d.values()):
            continue
        first = next(iter(d.values()), None)
        if first == "TOTAL":
            continue
        rows.append(d)
    return rows


def validate_structure(path: Path, db: Database | None = None, project_id: str | None = None) -> StructureReport:
    rep = StructureReport()
    wb_f = openpyxl.load_workbook(path, data_only=False)
    for s in REQUIRED_SHEETS:
        if s not in wb_f.sheetnames:
            rep.missing_sheets.append(s)
    for name, spec in S.items():
        if name not in wb_f.sheetnames:
            continue
        ws = wb_f[name]
        keys = [c.value for c in ws[KEY_ROW]][:len(spec.cols)]
        expect = [c.key for c in spec.cols]
        if keys != expect:
            rep.header_problems.append(f"{name}: column keys differ from template (expected {expect[:3]}..., found {keys[:3]}...)")
            continue
        for ci, col in enumerate(spec.cols, 1):
            if col.kind != "calc" or not col.formula:
                continue
            # a calc column cell that holds a constant where the row is otherwise populated = overwritten formula
            for r in range(FIRST_ROW, ws.max_row + 1):
                v = ws.cell(r, ci).value
                row_has = any(ws.cell(r, k).value not in (None, "") for k in range(1, len(spec.cols) + 1) if k != ci)
                if row_has and ws.cell(r, 1).value != "TOTAL" and not (isinstance(v, str) and v.startswith("=")):
                    rep.overwritten_formulas.append(f"{name}!{ws.cell(r, ci).coordinate} (formula replaced by constant)")
    if db is not None and project_id and "Project Setup" in wb_f.sheetnames:
        ws = wb_f["Project Setup"]
        for r in range(FIRST_ROW, ws.max_row + 1):
            if ws.cell(r, 1).value == "Estimate fingerprint":
                rep.fingerprint_in_file = ws.cell(r, 2).value
        rep.fingerprint_now = compute_estimate(db, project_id).fingerprint
        rep.stale = bool(rep.fingerprint_in_file) and rep.fingerprint_in_file != rep.fingerprint_now
    return rep


def validate_workbook(path: Path, db: Database | None = None, project_id: str | None = None, recalc: bool = True
                      ) -> tuple[StructureReport, VerifyReport | None]:
    st = validate_structure(path, db, project_id)
    vr = verify_workbook(path) if recalc else None
    return st, vr


# ----------------------------------------------------------------------------- import
@dataclass
class ImportReport:
    created: dict[str, int] = field(default_factory=dict)
    ignored: dict[str, int] = field(default_factory=dict)
    errors: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    def add(self, sheet: str, n: int = 1):
        self.created[sheet] = self.created.get(sheet, 0) + n

    def skip(self, sheet: str, n: int = 1):
        self.ignored[sheet] = self.ignored.get(sheet, 0) + n


def _s(v: Any) -> str | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    return str(v).strip()


def import_master(db: Database, actor: Actor, project_id: str, path: Path) -> ImportReport:
    """Import edits from a master workbook. Creates Drafts/proposals only - approval statuses in the file are ignored."""
    approvals.require_author(actor)
    rep = ImportReport()
    st = validate_structure(path)
    if st.missing_sheets or st.header_problems:
        raise ValidationFailed("Workbook structure invalid: " + "; ".join(st.missing_sheets + st.header_problems), code="WORKBOOK_INVALID")
    if st.overwritten_formulas:
        rep.notes.append(f"{len(st.overwritten_formulas)} formula cell(s) were overwritten with constants; they are ignored (the application recalculates).")
    wb = openpyxl.load_workbook(path, data_only=True)

    # ---- BOQ
    boq_rows = read_rows(wb, "BOQ Import")
    if boq_rows:
        try:
            s = boq_import.import_boq(db, actor, project_id, path, sheet="BOQ Import", header_row=HEADER_ROW)
            rep.add("BOQ Import", s["items_new"])
            rep.notes.append(f"BOQ: {s['items_new']} new, {s['items_unchanged']} unchanged, {s['change_requests']} change request(s) pending approval, "
                             f"{s['rejected']} rejected.")
        except AppError as e:
            rep.errors.append(f"BOQ Import: {e.message}")
    # ---- mapping
    for r in read_rows(wb, "Trade and Cost Code Mapping"):
        try:
            it = boq.get_item(db, project_id, r["item_key"])
        except AppError:
            rep.skip("Trade and Cost Code Mapping")
            continue
        try:
            if _s(r.get("cost_code")) and _s(r["cost_code"]) != it["cost_code"]:
                boq.set_cost_code(db, actor, project_id, it["item_key"], _s(r["cost_code"]))
                rep.add("Trade and Cost Code Mapping")
            if _s(r.get("assembly_code")) and _s(r["assembly_code"]) != it["assembly_code"]:
                boq.map_to_assembly(db, actor, project_id, it["item_key"], _s(r["assembly_code"]), source="manual",
                                    comment="mapped via workbook import")
                rep.add("Trade and Cost Code Mapping")
        except AppError as e:
            rep.errors.append(f"Mapping {r['item_key']}: {e.message}")
    # ---- assumptions overrides
    for r in read_rows(wb, "Assumptions Register"):
        if _s(r.get("override_value")) is None:
            continue
        try:
            asm_svc.set_override(db, actor, project_id, r["id"], str(r["override_value"]), _s(r.get("override_reason")) or "")
            rep.add("Assumptions Register")
        except AppError as e:
            rep.errors.append(f"Assumption {r['id']}: {e.message}")
    # ---- indirects
    for r in read_rows(wb, "Indirect Costs"):
        code = _s(r.get("code"))
        if not code or not _s(r.get("method")):
            rep.skip("Indirect Costs")
            continue
        try:
            existing = db.one("SELECT * FROM indirect_items WHERE project_id=? AND code=? AND status<>'Superseded'", (project_id, code))
            vals = dict(category=_s(r.get("category")) or "General", description=_s(r.get("description")) or code, method=r["method"],
                        amount=r.get("amount"), quantity=str(r.get("quantity") or 1), allocation_method=_s(r.get("allocation_method")) or "pro_rata_direct",
                        allocation_target=_s(r.get("allocation_target")), included_in_subcontract=bool(r.get("included_in_subcontract")))
            if existing is None:
                indirects.add_indirect(db, actor, project_id, code=code, **vals)
                rep.add("Indirect Costs")
            else:
                changed = {k: v for k, v in vals.items() if k in ("category", "description", "method", "allocation_method", "allocation_target")
                           and (existing[k] or None) != v}
                if D_opt(vals["amount"]) != D_opt(existing["amount"]):
                    changed["amount"] = vals["amount"]
                if D_opt(vals["quantity"]) != D_opt(existing["quantity"]):
                    changed["quantity"] = vals["quantity"]
                if int(bool(vals["included_in_subcontract"])) != existing["included_in_subcontract"]:
                    changed["included_in_subcontract"] = int(vals["included_in_subcontract"])
                if changed:
                    indirects.revise_indirect(db, actor, project_id, code, **changed)
                    rep.add("Indirect Costs")
        except AppError as e:
            rep.errors.append(f"Indirect {code}: {e.message}")
    # ---- registers (new rows only)
    for r in read_rows(wb, "Clarification Register"):
        if _s(r.get("id")) or not _s(r.get("question")):
            rep.skip("Clarification Register")
            continue
        try:
            registers.add_clarification(db, actor, project_id, r["question"], trade=_s(r.get("trade")), boq_item_key=_s(r.get("boq_item_key")),
                                        doc_ref=_s(r.get("doc_ref")), reason=_s(r.get("reason")), commercial_impact=_s(r.get("commercial_impact")),
                                        cost_exposure=r.get("cost_exposure"), proposed_wording=_s(r.get("proposed_wording")), owner=_s(r.get("owner")))
            rep.add("Clarification Register")
        except AppError as e:
            rep.errors.append(f"Clarification: {e.message}")
    for r in read_rows(wb, "Inclusions and Exclusions"):
        if _s(r.get("id")) or not _s(r.get("description")):
            rep.skip("Inclusions and Exclusions")
            continue
        try:
            registers.add_scope_entry(db, actor, project_id, _s(r.get("kind")) or "", r["description"], trade=_s(r.get("trade")),
                                      boq_item_key=_s(r.get("boq_item_key")), cost_code=_s(r.get("cost_code")), rationale=_s(r.get("rationale")))
            rep.add("Inclusions and Exclusions")
        except AppError as e:
            rep.errors.append(f"Scope entry: {e.message}")
    for r in read_rows(wb, "Risk Register"):
        if _s(r.get("id")) or not _s(r.get("title")):
            rep.skip("Risk Register")
            continue
        try:
            registers.add_risk(db, actor, project_id, r["title"], description=_s(r.get("description")) or "", trade=_s(r.get("trade")),
                               probability=_s(r.get("probability")), impact=_s(r.get("impact")), cost_exposure=r.get("cost_exposure"),
                               mitigation=_s(r.get("mitigation")), owner=_s(r.get("owner")))
            rep.add("Risk Register")
        except AppError as e:
            rep.errors.append(f"Risk: {e.message}")
    for r in read_rows(wb, "Vendor Directory"):
        if _s(r.get("id")) or not _s(r.get("name")):
            rep.skip("Vendor Directory")
            continue
        try:
            suppliers.create_supplier(db, actor, name=r["name"], supplier_type=_s(r.get("supplier_type")) or "supplier", trades=_s(r.get("trades")) or "",
                                      region=_s(r.get("region")) or "KSA-ALL", contact_name=_s(r.get("contact_name")), contact_email=_s(r.get("contact_email")),
                                      contact_phone=_s(r.get("contact_phone")), contact_source=_s(r.get("contact_source")))
            rep.add("Vendor Directory")
        except AppError as e:
            rep.errors.append(f"Vendor {r.get('name')}: {e.message}")
    # ---- rate proposals from libraries
    for sheet in ("Labour Library", "Material Library", "Plant and Equipment Library", "Subcontract Library"):
        for r in read_rows(wb, sheet):
            if D_opt(r.get("new_rate")) is None:
                continue
            rt = _s(r.get("new_rate_type"))
            try:
                if rt not in RATE_TYPE_VALUES or rt == "Missing Rate":
                    raise ValidationFailed(f"Proposed rate type '{rt}' invalid.")
                pd = r.get("new_price_date")
                rates.propose_rate(db, actor, r["id"], unit_price=r["new_rate"], rate_type=rt, currency=_s(r.get("new_currency")) or "SAR",
                                   price_date=pd.date().isoformat() if hasattr(pd, "date") else _s(pd),
                                   valid_until=(r["new_valid_until"].date().isoformat() if hasattr(r.get("new_valid_until"), "date") else _s(r.get("new_valid_until"))),
                                   evidence=_s(r.get("new_evidence")), submit=False)
                rep.add(sheet)
            except AppError as e:
                rep.errors.append(f"{sheet} {r['id']}: {e.message}")
    # ---- assemblies (new draft versions only)
    _import_assemblies(db, actor, read_rows(wb, "Assembly Library"), rep)
    audit.log(db, actor, "workbook_import", Path(path).name, "import", after={"created": rep.created, "errors": len(rep.errors)}, project_id=project_id)
    return rep


def _import_assemblies(db: Database, actor: Actor, rows: list[dict], rep: ImportReport) -> None:
    by: dict[str, list[dict]] = {}
    for r in rows:
        if _s(r.get("assembly_code")):
            by.setdefault(r["assembly_code"], []).append(r)
    for code, lines in by.items():
        def to_line(r):
            crew = []
            for part in (_s(r.get("crew")) or "").split(";"):
                if ":" in part:
                    rid, cnt = part.split(":")
                    crew.append({"resource_id": rid.strip(), "count": float(cnt)})
            d = dict(line_type=r["line_type"], resource_id=_s(r.get("resource_id")), child_assembly_code=_s(r.get("child_assembly_code")),
                     factor=None if D_opt(r.get("factor")) is None else str(r["factor"]), factor_unit=_s(r.get("factor_unit")),
                     factor_includes_waste=int(bool(r.get("factor_includes_waste"))), waste_basis=_s(r.get("waste_basis")) or "none",
                     waste_value=None if D_opt(r.get("waste_value")) is None else str(r["waste_value"]), waste_ref=_s(r.get("waste_ref")),
                     productivity_basis=_s(r.get("productivity_basis")), includes_labour=int(bool(r.get("includes_labour"))),
                     includes_material=int(bool(r.get("includes_material"))), includes_plant=int(bool(r.get("includes_plant"))),
                     utilisation_basis=_s(r.get("utilisation_basis")) or "none",
                     utilisation_value=None if D_opt(r.get("utilisation_value")) is None else str(r["utilisation_value"]), utilisation_ref=_s(r.get("utilisation_ref")))
            if crew:
                d["crew"] = crew
            return {k: v for k, v in d.items() if v is not None}
        try:
            new_lines = [to_line(r) for r in lines]
            try:
                cur = asm_service.get(db, code)
            except AppError:
                first = lines[0]
                asm_service.create_assembly(db, actor, code=code, description=_s(first.get("description")) or code, unit=first["unit"],
                                            cost_code=_s(first.get("cost_code")), lines_=new_lines)
                rep.add("Assembly Library")
                continue
            old = asm_service.lines(db, cur["pk"])

            def norm(ls):
                out = []
                for l in ls:
                    out.append((l["line_type"], l.get("resource_id") or None, l.get("child_assembly_code") or None, str(D_opt(l.get("factor"))),
                                l.get("waste_basis") or "none", str(D_opt(l.get("waste_value"))), l.get("waste_ref") or None,
                                l.get("productivity_basis") or None, int(l.get("factor_includes_waste") or 0)))
                return out
            if norm(old) == norm(new_lines):
                rep.skip("Assembly Library")
                continue
            if cur["status"] in ("Draft", "Rejected", "Pending Review"):
                pk = cur["pk"]
                db.execute("DELETE FROM assembly_lines WHERE assembly_pk=?", (pk,))
                db.update("assemblies", {"status": "Draft"}, "pk=?", (pk,))
            else:
                pk = asm_service.new_version(db, actor, code, revision_note="Imported from workbook")
                db.execute("DELETE FROM assembly_lines WHERE assembly_pk=?", (pk,))
            for l in new_lines:
                asm_service.add_line(db, actor, pk, **l)
            rep.add("Assembly Library")
        except AppError as e:
            rep.errors.append(f"Assembly {code}: {e.message}")
