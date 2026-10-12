from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt


def replace_in_paragraph(p, mapping: dict[str, str]) -> None:
    full = "".join(r.text for r in p.runs)
    if "{{" not in full:
        return
    for k, v in mapping.items():
        token = "{{" + k + "}}"
        if token in full:
            full = full.replace(token, v)
    if p.runs:
        p.runs[0].text = full
        for r in p.runs[1:]:
            r.text = ""


def fill_template(template: Path, mapping: dict[str, str]):
    doc = Document(str(template))
    for p in doc.paragraphs:
        replace_in_paragraph(p, mapping)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_in_paragraph(p, mapping)
    for sec in doc.sections:
        for p in sec.header.paragraphs + sec.footer.paragraphs:
            replace_in_paragraph(p, mapping)
    return doc


def insert_list_after(doc, marker: str, items: list[str], style: str | None = None, empty_text: str = "None.") -> None:
    """Replace the paragraph containing {{marker}} with one paragraph per item."""
    for p in doc.paragraphs:
        if "{{" + marker + "}}" in "".join(r.text for r in p.runs):
            anchor = p
            lines = items or [empty_text]
            anchor.runs[0].text = lines[0] if style is None else lines[0]
            for r in anchor.runs[1:]:
                r.text = ""
            prev = anchor
            for line in lines[1:]:
                new = doc.add_paragraph(line)
                prev._p.addnext(new._p)
                prev = new
            return


def build_simple_template(path: Path, title: str, body_lines: list[str], header_text: str = "") -> None:
    doc = Document()
    doc.add_heading(title, level=0)
    for line in body_lines:
        doc.add_paragraph(line)
    if header_text:
        doc.sections[0].header.paragraphs[0].text = header_text
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
