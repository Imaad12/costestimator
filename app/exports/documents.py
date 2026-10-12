"""Client-facing and internal documents. Preliminary outputs are always watermarked; SUBMISSION READY is only stamped by
`services.estimates.export_submission` after explicit human approval."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..config import get_settings
from ..core.money import ZERO, fmt, to_text
from ..core.timeutil import today_iso
from ..estimators.estimate import EstimateResult
from ..exports.docx_util import fill_template, insert_list_after
from ..services import projects

SAR = '#,##0.00;[Red]-#,##0.00'
HDR = PatternFill("solid", fgColor="1F3864")
WARN = PatternFill("solid", fgColor="FCE4D6")
RED = PatternFill("solid", fgColor="F8CBAD")
GREY = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="BFBFBF")


def banner_text(est: EstimateResult, submission_ready: bool) -> str:
    if submission_ready:
        t = "SUBMISSION READY - approved by an authorised human reviewer"
    elif est.blockers or est.state.startswith("Preliminary"):
        t = "PRELIMINARY - NOT FOR SUBMISSION"
    else:
        t = "DRAFT - NOT YET APPROVED FOR SUBMISSION"
    if est.project.get("is_demo"):
        t += " | DEMONSTRATION PROJECT - SYNTHETIC DATA - NOT A MARKET QUOTE"
    elif est.flags.get("synthetic"):
        t += " | CONTAINS SYNTHETIC DATA"
    if est.flags.get("provisional"):
        t += " | CONTAINS PROVISIONAL (PROJECT-APPROVED) INPUTS"
    return t


def client_pricing_schedule(est: EstimateResult, path: Path | None = None, *, submission_ready: bool = False) -> Path:
    """Client-facing schedule: quantities, rates and amounts only (no internal cost, markup or allocation detail)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Pricing Schedule"
    sch, c = est.schedule, est.commercial
    p = est.project
    ws["A1"] = f"PRICING SCHEDULE - {p['name']} ({p['id']})"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = banner_text(est, submission_ready)
    ws["A2"].font = Font(bold=True, color="C00000")
    ws["A3"] = f"Currency: {p['currency']} | Prices exclude tax unless stated | Generated {today_iso()}"
    r = 5
    for ci, h in enumerate(["Item", "Description", "Unit", "Quantity", "Unit rate", "Amount", "Note"], 1):
        cell = ws.cell(r, ci, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HDR
    ws.freeze_panes = "A6"
    first = r + 1
    if sch is None:
        ws.cell(first, 1, "Pricing schedule unavailable: commercial policy unresolved.")
    else:
        for l in sch.lines:
            r += 1
            note = {"Excluded": "Excluded - see offer exclusions", "Unpriced": "UNPRICED - unresolved", "Provisional": "Provisional / allowance"}.get(l.status, "")
            if "PROVISIONAL_INPUTS" in l.flags:
                note = (note + "; " if note else "") + "provisional inputs"
            vals = [l.item_code or l.item_key, l.description, l.unit, l.quantity, l.unit_rate, l.amount, note]
            for ci, v in enumerate(vals, 1):
                cell = ws.cell(r, ci, None if v is None else (float(v) if hasattr(v, "quantize") else v))
                if ci in (4, 5, 6):
                    cell.number_format = SAR if ci > 4 else "#,##0.000"
                if l.status == "Unpriced":
                    cell.fill = RED
            ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        for s in sch.separate_lines:
            r += 1
            for ci, v in enumerate([s["code"], s["description"], "ls", 1, float(s["price"]), float(s["price"]), "Indirect / preliminaries"], 1):
                cell = ws.cell(r, ci, v)
                if ci in (5, 6):
                    cell.number_format = SAR
        last = r
        r += 2
        ws.cell(r, 5, "Total before tax").font = Font(bold=True)
        ws.cell(r, 6, f"=SUM(F{first}:F{last})").number_format = SAR
        ws.cell(r, 7, "Complete" if sch.complete else "INCOMPLETE - unresolved items remain")
        if c:
            for cat, v in c.tax_by_category.items():
                r += 1
                ws.cell(r, 5, f"Tax {cat} @ {v['rate_pct']}%" if v["rate_pct"] is not None else f"Tax {cat} (rate unresolved)")
                ws.cell(r, 6, None if v["tax"] is None else float(v["tax"])).number_format = SAR
            r += 1
            ws.cell(r, 5, "Total including tax").font = Font(bold=True)
            ws.cell(r, 6, f"=F{r - len(c.tax_by_category) - 1}+" + "+".join(f"F{r - i}" for i in range(1, len(c.tax_by_category) + 1))).number_format = SAR
    for col, w in zip("ABCDEFG", (12, 60, 8, 12, 14, 16, 34)):
        ws.column_dimensions[col].width = w
    ws.auto_filter.ref = f"A5:G{max(r, 6)}"
    ws.protection.sheet = True
    ws.protection.password = ""
    out = path or (projects.output_dir(p["id"]) / f"pricing_schedule_{p['id']}{'_FINAL' if submission_ready else '_PRELIMINARY'}.xlsx")
    wb.save(out)
    return out


def estimate_report_docx(est: EstimateResult, path: Path | None = None, *, submission_ready: bool = False, version_label: str = "-",
                         assumptions: list[dict] | None = None) -> Path:
    tmpl = get_settings().templates_dir / "estimate_report_template.docx"
    if not tmpl.exists():
        from .docx_templates import generate_docx_templates
        generate_docx_templates(get_settings().templates_dir)
    p, c, d = est.project, est.commercial, est.direct
    mapping = dict(status_banner=banner_text(est, submission_ready), project_name=p["name"], project_id=p["id"], client=p.get("client") or "",
                   version=version_label, date=today_iso(), state=est.state)
    doc = fill_template(tmpl, mapping)
    summ = [f"Direct cost: {fmt(d.direct_cost_total)} {p['currency']}", f"Indirect cost: {fmt(est.indirect.total)}"]
    if c:
        summ += [f"Contingency: {fmt(c.contingency)}", f"Risk allowances: {fmt(c.risk)}", f"Escalation: {fmt(c.escalation)}",
                 f"Total cost: {fmt(c.total_cost)}", f"Commercial method: {c.method_label}", f"Selling price before discount: {fmt(c.selling_price_pre_discount)}",
                 f"Discount: {fmt(c.discount)}", f"Net price before tax: {fmt(c.net_price)} (commercial rounding {fmt(c.rounding_delta)})",
                 f"Tax: {fmt(c.tax_total)}", f"Final payable: {fmt(c.final_payable)}"]
        if est.schedule:
            summ.append(f"Pricing schedule total: {fmt(est.schedule.total_pre_tax)} (rounding variance {fmt(est.schedule.rounding_variance)})")
    else:
        summ.append("Commercial calculation unavailable (policy unresolved).")
    summ.append(f"Completeness: {est.completeness['items_complete']}/{est.completeness['items_total']} BOQ items fully priced"
                + (f" ({est.completeness['pct_items_complete']}%)" if est.completeness["pct_items_complete"] is not None else " (no priced scope)"))
    summ.append(f"Indicative (unapproved-rate) exposure: {fmt(est.completeness['indicative_exposure'])}; lines with unknown exposure: "
                f"{est.completeness['unknown_exposure_lines']}")
    insert_list_after(doc, "SUMMARY", summ)
    recon = []
    if c:
        for t in c.trace:
            recon.append(f"{t.name}: {fmt(t.amount)} [{t.formula}]")
    insert_list_after(doc, "RECON_TABLE", recon, empty_text="Not available.")
    exc = [f"[{e['severity']}] {e['code']}: {e['message']}" for e in est.exceptions if e["severity"] != "Low"][:200]
    insert_list_after(doc, "EXCEPTIONS", exc + [f"BLOCKER: {b}" for b in est.blockers], empty_text="No exceptions.")
    insert_list_after(doc, "ASSUMPTIONS", [f"{a['id']} {a['name']}: {a['value']} {a.get('unit') or ''} [{a['status']}]" for a in (assumptions or [])],
                      empty_text="None.")
    out = path or (projects.output_dir(p["id"]) / f"estimate_report_{p['id']}{'_FINAL' if submission_ready else '_PRELIMINARY'}.docx")
    doc.save(str(out))
    return out


def reconciliation_report_xlsx(run: dict, project: dict, path: Path | None = None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reconciliation"
    ws["A1"] = f"RECONCILIATION / AUDIT REPORT - {project['name']} ({project['id']}) - run #{run['id']} at {run['ts']}"
    ws["A1"].font = Font(bold=True, size=13)
    s = run["summary"]
    ws["A2"] = f"Open by severity: {s.get('open_by_severity')} | Blocking: {s.get('blocking')} | Estimate state: {s.get('state')}"
    heads = ["Severity", "Rule", "Title", "Detail", "Entity", "Financial impact (SAR)", "Owner", "Status", "Resolution / exception reason"]
    for ci, h in enumerate(heads, 1):
        c = ws.cell(4, ci, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HDR
    fills = {"Critical": RED, "High": WARN}
    for ri, f in enumerate(run["findings"], 5):
        vals = [f["severity"], f["rule_code"], f["title"], f["detail"], f"{f['entity_type'] or ''}:{f['entity_id'] or ''}",
                None if f["financial_impact"] is None else float(f["financial_impact"]), f["owner"], f["status"], f.get("exception_reason") or ""]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(ri, ci, v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if ci == 6:
                c.number_format = SAR
            if f["severity"] in fills and f["status"] == "Open":
                c.fill = fills[f["severity"]]
    for col, w in zip("ABCDEFGHI", (10, 28, 36, 80, 24, 18, 20, 18, 40)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:I{max(5, 4 + len(run['findings']))}"
    out = path or (projects.output_dir(project["id"]) / f"reconciliation_report_{project['id']}_run{run['id']}.xlsx")
    wb.save(out)
    return out
