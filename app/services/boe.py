"""Basis of Estimate (BoE) generator. Assembles the FACTS the system holds - scope, sources, factors, assumptions, indirects, commercial policy, programme,
risks, classification, client-profile status, exceptions and approvals - into a document, and leaves every judgement section as an explicit
[TO BE COMPLETED BY ESTIMATOR] placeholder. Nothing is invented and nothing is presented as approved that is not."""
from __future__ import annotations

import hashlib
import json
from decimal import Decimal
from pathlib import Path

from ..approvals import engine as approvals
from ..core.money import fmt
from ..core.timeutil import today_iso
from ..database import Database, rows_to_dicts
from ..estimators.estimate import compute_estimate
from ..models.actor import Actor
from . import analysis, governed, mto, profiles, projects, workspaces

PLACEHOLDER = "[TO BE COMPLETED BY ESTIMATOR]"
HUMAN_SECTIONS = [
    ("Design basis and scope narrative", "Describe the design basis, the extent of the engineering received and any design still to be developed."),
    ("Construction methodology and sequencing", "Describe the construction approach, major methods, key resources and constraints assumed."),
    ("Productivity and labour basis narrative", "Explain the productivity norms applied and their source (company actuals, client norms, benchmarks)."),
    ("Procurement and subcontracting strategy", "Describe package strategy, long-lead items, vendor/subcontractor selection and any client-nominated items."),
    ("Site conditions, access and logistics", "Describe the site, access, laydown, climate and logistics assumptions."),
    ("Risks and opportunities commentary", "Comment on key risks, how exposure was assessed and what the contingency covers."),
]


def build_boe(db: Database, project_id: str, est=None) -> dict:
    est = est or compute_estimate(db, project_id)
    p = est.project
    inp = est.inputs
    imports = rows_to_dicts(db.all("SELECT revision, source_file, sha256, imported_by, imported_at, summary_json FROM boq_imports WHERE project_id=? ORDER BY id", (project_id,)))
    methods: dict[str, int] = {}
    for i in est.direct.items.values():
        methods[i.pricing_method] = methods.get(i.pricing_method, 0) + 1
    # rate evidence
    rate_rows = []
    if est.direct.used_rates:
        q = ",".join("?" * len(est.direct.used_rates))
        rate_rows = rows_to_dicts(db.all(f"SELECT id, resource_id, rate_type, status, source_id, price_date, valid_until, is_synthetic, confidence FROM resource_rates WHERE id IN ({q})",
                                         tuple(sorted(est.direct.used_rates))))
    by_type: dict[str, int] = {}
    by_source: dict[str, int] = {}
    for r in rate_rows:
        by_type[r["rate_type"]] = by_type.get(r["rate_type"], 0) + 1
        by_source[r["source_id"] or "(none recorded)"] = by_source.get(r["source_id"] or "(none recorded)", 0) + 1
    asm = [dict(code=a.code, version=a.version, status=a.status, synthetic=a.is_synthetic) for a in inp.assemblies.values()]
    params = {}
    for kind in ("productivity_factor", "mto_factor", "plant_hire_terms", "calendar", "classification_scheme", "maturity_input"):
        rows = governed.effective(db, project_id, kind)
        if rows:
            params[kind] = [dict(code=r["code"], version=r["version"], scope=r["scope"], source=r["source"], approved_by=r["approved_by"]) for r in rows]
    assumptions_used = []
    for aid in sorted(est.direct.used_assumptions):
        a = inp.assumptions.get(aid)
        assumptions_used.append(dict(id=aid, name=a.name if a else "", value=a.value if a else None, status=a.status if a else None))
    for k, a in inp.engine_keys.items():
        assumptions_used.append(dict(id=a.id, name=a.name, value=a.value, status=a.status, engine_key=k))
    scope = rows_to_dicts(db.all("SELECT id, kind, description, status FROM scope_register WHERE project_id=? ORDER BY id", (project_id,)))
    clar = rows_to_dicts(db.all("SELECT id, question, status FROM clarifications WHERE project_id=? ORDER BY id", (project_id,)))
    risks = rows_to_dicts(db.all("SELECT id, title, probability, impact, cost_exposure, prob_pct, cost_likely, status FROM risks WHERE project_id=? ORDER BY id", (project_id,)))
    ind = [dict(code=l.code, description=l.description, amount=l.amount) for l in est.indirect.lines]
    sched = rows_to_dicts(db.all("SELECT task_code, name, duration_days, predecessors FROM schedule_tasks WHERE project_id=? ORDER BY id", (project_id,)))
    c = est.commercial
    cls = profiles.classification_report(db, project_id, est=est)
    comp = profiles.compliance_report(db, project_id, est=est)
    mt = mto.status_report(db, project_id)
    appr = rows_to_dicts(db.all("SELECT entity_type, COUNT(*) n FROM approvals WHERE project_id=? GROUP BY entity_type ORDER BY entity_type", (project_id,)))
    banner = []
    if p["is_demo"]:
        banner.append("DEMONSTRATION PROJECT - SYNTHETIC DATA - NOT A MARKET QUOTE")
    if est.state != "Complete":
        banner.append(f"ESTIMATE STATE: {est.state.upper()} - NOT FOR SUBMISSION")
    if est.direct.synthetic_used:
        banner.append("SYNTHETIC RATES IN USE")
    if (p.get("notes") or "").strip():                                   # the project's own human-entered label/limitation statement
        banner.append(f"PROJECT NOTE: {p['notes'].strip()}")
    return dict(
        generated=today_iso(), project=dict(id=p["id"], name=p["name"], client=p.get("client"), location=p.get("location"), region=p.get("region"), currency=p["currency"],
                                            tender_ref=p.get("tender_ref"), design_stage=p.get("design_stage"), drawing_revision=p.get("drawing_revision"),
                                            spec_revision=p.get("spec_revision"), as_of=est.as_of.isoformat()),
        banner=banner, estimate=dict(state=est.state, fingerprint=est.fingerprint, totals=est.totals(), blockers=est.blockers, completeness=est.completeness),
        scope=dict(boq_imports=[dict(r, summary=json.loads(r["summary_json"] or "{}")) for r in imports], pricing_methods=methods, register=scope, clarifications=clar,
                   statement="Quantities are measured and entered by people; the system never measures drawings or completes quantities. Derived quantities (take-off templates) are listed under 'Take-off derivations'."),
        pricing=dict(rate_types=by_type, rate_sources=by_source, rates_used=len(rate_rows), synthetic_rates=sum(1 for r in rate_rows if r["is_synthetic"]),
                     provisional_rates=sum(1 for r in rate_rows if r["status"].startswith("Provisional")), unpriced_items=est.completeness["items_unpriced"],
                     oldest_price_date=min((r["price_date"] for r in rate_rows if r["price_date"]), default=None),
                     tax=est.tax_status, fx=[f for f in inp.fx if f.get("status") == "Approved"][:20]),
        assemblies=asm, parameters=params, assumptions=assumptions_used, take_off=dict(applications=mt["applications"], stale=len(mt["stale"])),
        indirects=dict(lines=ind, total=est.indirect.total, duration_months=est.indirect.duration_months, duration_status=est.indirect.duration_status),
        programme=dict(tasks=sched),
        commercial=None if not c else dict(policy=c.policy_name, status=c.policy_status, method=c.method_label, direct=c.direct, indirect=c.indirect, contingency=c.contingency,
                                          risk=c.risk, escalation=c.escalation, total_cost=c.total_cost, margin=c.margin_amount, net_price=c.net_price, tax_total=c.tax_total),
        risks=risks, classification=cls, compliance=dict(linked=comp["linked"], statement=comp["statement"], profile=comp.get("profile"), counts=comp.get("counts")),
        exceptions=[dict(code=e["code"], severity=e["severity"], message=e["message"]) for e in est.exceptions[:200]], exceptions_total=len(est.exceptions),
        approvals=appr, human_sections=[dict(title=t, guidance=g, text=PLACEHOLDER) for t, g in HUMAN_SECTIONS])


def _md_table(rows: list[dict], cols: list[str]) -> str:
    if not rows:
        return "_None recorded._\n"
    out = ["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
    for r in rows:
        out.append("| " + " | ".join(str(r.get(c, "") if r.get(c) is not None else "") .replace("|", "/") for c in cols) + " |")
    return "\n".join(out) + "\n"


def render_markdown(b: dict) -> str:
    p, e = b["project"], b["estimate"]
    L = [f"# Basis of Estimate - {p['id']} {p['name']}", ""]
    for x in b["banner"]:
        L += [f"> **{x}**", ""]
    L += [f"Generated {b['generated']} from the estimating database. Estimate state: **{e['state']}**. Fingerprint `{e['fingerprint'][:16]}`.", ""]
    L += ["## 1. Project", _md_table([{"item": k, "value": v} for k, v in p.items()], ["item", "value"])]
    L += ["## 2. Scope and quantities", b["scope"]["statement"], "", "**BOQ imports**", _md_table(b["scope"]["boq_imports"], ["revision", "source_file", "imported_at", "sha256"]),
          "**Pricing methods**", _md_table([{"method": k, "items": v} for k, v in sorted(b["scope"]["pricing_methods"].items())], ["method", "items"]),
          "**Scope register**", _md_table(b["scope"]["register"], ["id", "kind", "status", "description"]),
          "**Clarifications**", _md_table(b["scope"]["clarifications"], ["id", "status", "question"])]
    pr = b["pricing"]
    L += ["## 3. Pricing basis", f"Rates used: {pr['rates_used']} ({pr['synthetic_rates']} synthetic, {pr['provisional_rates']} provisional). Unpriced items: {pr['unpriced_items']}. "
          f"Oldest price date: {pr['oldest_price_date'] or 'n/a'}.", "", _md_table([{"rate_type": k, "count": v} for k, v in pr["rate_types"].items()], ["rate_type", "count"]),
          _md_table([{"source": k, "rates": v} for k, v in pr["rate_sources"].items()], ["source", "rates"])]
    L += ["## 4. Assemblies, factors and assumptions", "**Assemblies used**", _md_table(b["assemblies"], ["code", "version", "status", "synthetic"])]
    for k, rows in b["parameters"].items():
        L += [f"**Governed parameters - {k}**", _md_table(rows, ["code", "version", "scope", "source", "approved_by"])]
    L += ["**Assumptions in force**", _md_table(b["assumptions"], ["id", "name", "value", "status"])]
    L += ["## 5. Take-off derivations", _md_table(b["take_off"]["applications"], ["item_key", "template", "driver_qty", "driver_unit", "pricing_status", "flags"]),
          f"Stale applications: {b['take_off']['stale']}.", ""]
    ind = b["indirects"]
    L += ["## 6. Indirects and duration", f"Duration: {ind['duration_months']} months ({ind['duration_status']}). Indirect total: {fmt(ind['total'])}.", "",
          _md_table(ind["lines"], ["code", "description", "amount"])]
    L += ["## 7. Programme", _md_table(b["programme"]["tasks"], ["task_code", "name", "duration_days", "predecessors"])]
    cm = b["commercial"]
    if cm:
        L += ["## 8. Commercial", _md_table([{"item": k, "value": fmt(v) if isinstance(v, Decimal) else v} for k, v in cm.items()], ["item", "value"])]
    else:
        L += ["## 8. Commercial", "No computable commercial policy.", ""]
    L += ["## 9. Risk register", _md_table(b["risks"], ["id", "title", "probability", "impact", "cost_exposure", "prob_pct", "cost_likely", "status"])]
    c = b["classification"]
    L += ["## 10. Classification and maturity",
          (f"Scheme: {c['scheme']['name']} ({c['scheme']['source_document']}). Definition maturity {c['definition_pct']}% ({c.get('definition_basis')}). Class: **{c['class_id']}**."
           if c.get("class_id") else "No classification determined. " + " ".join(c["notes"])), ""]
    L += ["## 11. Client profile", b["compliance"]["statement"], ""]
    L += ["## 12. Exceptions", f"{b['exceptions_total']} exception(s); first {len(b['exceptions'])} shown.", _md_table(b["exceptions"], ["severity", "code", "message"])]
    L += ["## 13. Approvals recorded", _md_table(b["approvals"], ["entity_type", "n"])]
    L += ["## 14. Sections requiring the estimator's judgement", "_The system does not write these. Complete them before issue._", ""]
    for h in b["human_sections"]:
        L += [f"### {h['title']}", f"_{h['guidance']}_", "", h["text"], ""]
    return "\n".join(L)


def render_docx(b: dict, path: Path) -> None:
    from docx import Document
    md = render_markdown(b)
    doc = Document()
    for line in md.split("\n"):
        if line.startswith("# "):
            doc.add_heading(line[2:], 0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], 1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], 2)
        elif line.startswith("|"):
            continue
        elif line.startswith("> "):
            doc.add_paragraph().add_run(line[2:].replace("**", "")).bold = True
        elif line.strip():
            doc.add_paragraph(line.replace("**", "").replace("_", "") if line.startswith(("**", "_")) else line.replace("**", ""))
    # tables rendered properly from the structured data (markdown tables are skipped above)
    doc.add_heading("Tables", 1)
    for title, rows, cols in (("BOQ imports", b["scope"]["boq_imports"], ["revision", "source_file", "imported_at"]),
                              ("Scope register", b["scope"]["register"], ["id", "kind", "status", "description"]),
                              ("Assumptions in force", b["assumptions"], ["id", "name", "value", "status"]),
                              ("Indirects", b["indirects"]["lines"], ["code", "description", "amount"]),
                              ("Exceptions", b["exceptions"], ["severity", "code", "message"])):
        doc.add_heading(title, 2)
        if not rows:
            doc.add_paragraph("None recorded.")
            continue
        t = doc.add_table(rows=1, cols=len(cols))
        t.style = "Table Grid"
        for i, cname in enumerate(cols):
            t.rows[0].cells[i].text = cname
        for r in rows:
            cells = t.add_row().cells
            for i, cname in enumerate(cols):
                v = r.get(cname)
                cells[i].text = "" if v is None else str(v)
    doc.save(str(path))


def generate_boe(db: Database, actor: Actor, project_id: str) -> dict:
    approvals.require_author(actor)
    projects.require_active(db, project_id)
    est = compute_estimate(db, project_id)
    b = build_boe(db, project_id, est)
    md = render_markdown(b)
    n = (db.scalar("SELECT COUNT(*) FROM analysis_runs WHERE project_id=? AND kind='boe'", (project_id,)) or 0) + 1
    d = workspaces.area_dir(project_id, "estimate", db) / "basis_of_estimate"
    d.mkdir(exist_ok=True)
    stem = f"BoE_{project_id}_draft{n:02d}"
    (d / f"{stem}.md").write_text(md, encoding="utf-8")
    render_docx(b, d / f"{stem}.docx")
    rid = analysis.save_run(db, actor, project_id, "boe", {"draft": n}, {"file_md": f"{stem}.md", "file_docx": f"{stem}.docx", "sha256": hashlib.sha256(md.encode()).hexdigest(),
                                                                         "estimate_state": est.state, "fingerprint": est.fingerprint}, inputs_hash=est.fingerprint)
    return dict(run_id=rid, directory=str(d), files=[f"{stem}.md", f"{stem}.docx"], estimate_state=est.state, placeholders=len(b["human_sections"]))
