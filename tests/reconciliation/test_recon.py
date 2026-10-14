from decimal import Decimal

import pytest

from app.core.errors import ApprovalBlocked, PermissionDenied, ValidationFailed
from app.estimators.estimate import compute_estimate
from app.imports import boq as bi
from app.reconciliation import engine as recon
from app.reconciliation.rules import load_rules
from app.services import assumptions, boq, estimates, indirects, offers, projects, registers, rates
from app.services.rates import create_resource

from tests.conftest import ADMIN, COMMERCIAL, ESTIMATOR, PROCUREMENT
from tests.helpers import approved_assembly, approved_rate, build_mini, write_boq


@pytest.fixture()
def mini(seeded, tmp_path):
    return seeded, build_mini(seeded, tmp_path), tmp_path


def findings(db, pid, **kw):
    r = recon.run_reconciliation(db, ESTIMATOR, pid, **kw)
    return recon.get_run(db, r["run_id"])["findings"]


def codes(fs, status="Open"):
    return {f["rule_code"] for f in fs if f["status"] == status}


def accept_adopted_assumptions(db, pid, fs):
    """Unapproved adopted (unused) assumptions are High: a Commercial Approver must accept them as a recorded exception - they never pass silently."""
    for f in fs:
        if f["rule_code"] == "UNAPPROVED_ADOPTED_ASSUMPTION" and f["status"] == "Open":
            recon.accept_exception(db, COMMERCIAL, f["id"], "unused mandatory assumptions: basis to be completed after award (test)")


def test_clean_project_has_no_blocking_findings(mini):
    db, pid, _ = mini
    fs = findings(db, pid)
    blocking = [f for f in fs if f["status"] == "Open" and f["severity"] in ("Critical", "High")]
    assert {f["rule_code"] for f in blocking} == {"UNAPPROVED_ADOPTED_ASSUMPTION"}, blocking       # O-2: nothing else, and this one needs a person's recorded exception
    assert recon.final_gate(db, ESTIMATOR, pid)["blocked"]
    accept_adopted_assumptions(db, pid, fs)
    assert not recon.final_gate(db, ESTIMATOR, pid)["blocked"]
    assert {"OFFER_MISSING", "SYNTHETIC_DEMO"} <= codes(fs)


def test_unpriced_item_and_missing_rate_are_critical(mini):
    db, pid, tmp = mini
    create_resource(db, ESTIMATOR, resource_type="material", description="no rate", unit="m", id="NR-1", is_synthetic=True)
    approved_assembly(db, "A-NR", "m", [dict(line_type="material", resource_id="NR-1", factor="1")])
    bi.import_boq(db, ESTIMATOR, pid, write_boq(tmp / "more.csv", [("2.1", "Item w/o rate", "m", "10", "", ""), ("2.2", "Item w/o mapping", "m", "3", "", "")]))
    boq.map_to_assembly(db, ESTIMATOR, pid, "2.1", "A-NR")
    fs = findings(db, pid)
    sev = {(f["rule_code"], f["entity_id"]): f["severity"] for f in fs}
    assert sev[("UNPRICED_ITEM", "2.1")] == "Critical" and sev[("UNPRICED_ITEM", "2.2")] == "Critical"
    assert sev[("MISSING_RATE", "NR-1")] == "Critical"
    assert recon.final_gate(db, ESTIMATOR, pid)["blocked"]


def test_unapproved_assumption_in_use_flagged(mini):
    db, pid, _ = mini
    create_resource(db, ESTIMATOR, resource_type="material", description="w", unit="m3", id="W-1", is_synthetic=True)
    approved_rate(db, "W-1", "5")
    approved_assembly(db, "A-W", "m3", [dict(line_type="material", resource_id="W-1", factor="1", waste_basis="assumption", waste_ref="ASMP-C-010")])
    db.execute("UPDATE boq_items SET assembly_code='A-W', pricing_method='assembly' WHERE project_id=? AND item_key='1.1'", (pid,))
    # assumption has no value & is unapproved: the line stays unpriced and the cause is reported
    fs = findings(db, pid)
    assert ("UNAPPROVED_ASSUMPTION", "ASMP-C-010") in {(f["rule_code"], f["entity_id"]) for f in fs}
    f = next(f for f in fs if f["rule_code"] == "UNAPPROVED_ASSUMPTION" and f["entity_id"] == "ASMP-C-010")
    assert f["severity"] == "Critical"
    with pytest.raises(Exception) as e:                                       # O-2: an unapproved assumption in use can never be waived, even by a Commercial Approver
        recon.accept_exception(db, COMMERCIAL, f["id"], "attempt to waive")
    assert getattr(e.value, "code", "") == "EXCEPTION_NOT_ALLOWED"
    assumptions.set_override(db, ESTIMATOR, pid, "ASMP-C-010", "5", "synthetic")
    fs = findings(db, pid)
    assert any(f["rule_code"] == "UNAPPROVED_ASSUMPTION" and "Pending Review" in f["detail"] for f in fs)           # value present but not approved
    assumptions.approve_override(db, COMMERCIAL, pid, "ASMP-C-010")
    assert "UNAPPROVED_ASSUMPTION" not in codes(findings(db, pid))


def test_expired_rate_flagged_and_not_used(mini):
    db, pid, _ = mini
    with pytest.raises(ValidationFailed) as e:                                   # an already-lapsed rate cannot be approved
        r = rates.propose_rate(db, ESTIMATOR, "DEMO-MAT-001", unit_price="110", rate_type="Synthetic Demonstration Rate", price_date="2026-01-01",
                               valid_until="2026-02-01", is_synthetic=True, evidence="e", submit=True)
        rates.approve_rate(db, PROCUREMENT, r)
    assert e.value.code == "RATE_ALREADY_EXPIRED"
    db.execute("UPDATE resource_rates SET valid_until='2026-01-01' WHERE resource_id='DEMO-MAT-001' AND status='Approved'")   # time passes
    fs = findings(db, pid)
    assert "UNPRICED_ITEM" in codes(fs) and "EXPIRED_RATE" in codes(fs)
    assert rates.expire_stale_rates(db, actor="system") == 1
    assert db.scalar("SELECT status FROM resource_rates WHERE resource_id='DEMO-MAT-001' AND unit_price='100'") == "Expired"


def test_quotation_expiry_behind_used_rate(mini):
    db, pid, _ = mini
    from app.procurement import quotes
    from app.services import suppliers
    sid = suppliers.create_supplier(db, ESTIMATOR, name="S", is_synthetic=True)
    q = quotes.create_quotation(db, ESTIMATOR, supplier_id=sid, project_id=pid, quote_date="2026-01-01", valid_until="2026-02-01", tax_treatment="excl",
                                lines=[dict(description="c", unit_rate="90", unit="m3", quantity="1")], is_synthetic=True)
    lid = db.scalar("SELECT id FROM quotation_lines WHERE quotation_id=?", (q,))
    r = rates.propose_rate(db, ESTIMATOR, "DEMO-MAT-001", unit_price="90", rate_type="Synthetic Demonstration Rate", price_date="2026-01-01", quotation_line_id=lid,
                           is_synthetic=True, evidence="e", submit=True)
    rates.approve_rate(db, PROCUREMENT, r)
    db.execute("UPDATE resource_rates SET valid_until=NULL WHERE id=?", (r,))       # rate itself carries no expiry, but the quotation behind it has expired
    fs = findings(db, pid)
    assert ("QUOTATION_EXPIRED", q) in {(f["rule_code"], f["entity_id"]) for f in fs}


def test_duplicate_items_and_cost_in_multiple_packages(mini):
    db, pid, tmp = mini
    bi.import_boq(db, ESTIMATOR, pid, write_boq(tmp / "dup.csv", [("9.1", "Concrete", "m3", "10", "", ""), ("9.2", "Concrete", "m3", "10", "", "")]))
    fs = findings(db, pid)
    assert any(f["rule_code"] == "POSSIBLE_DUPLICATE_ITEM" and "9.1" in f["detail"] and "9.2" in f["detail"] for f in fs)
    db.execute("UPDATE boq_items SET cost_code='D.01', quantity='10' WHERE project_id=? AND item_key='1.2'", (pid,))   # same resource+quantity under another trade
    for k in ("9.1", "9.2"):
        boq.map_to_assembly(db, ESTIMATOR, pid, k, "DEMO-CONC")
    fs = findings(db, pid)
    assert any(f["rule_code"] == "POSSIBLE_DUPLICATE_ITEM" and f["entity_type"] == "resource" and "several packages" in f["detail"] for f in fs)


def test_double_count_subcontract_is_critical(mini):
    db, pid, _ = mini
    create_resource(db, ESTIMATOR, resource_type="subcontract", description="sub", unit="m3", id="SUB-X", is_synthetic=True)
    approved_rate(db, "SUB-X", "50")
    approved_assembly(db, "A-DBL", "m3", [dict(line_type="subcontract", resource_id="SUB-X", factor="1", includes_labour=1),
                                          dict(line_type="labour", resource_id="DEMO-LAB-001", factor="1", productivity_basis="paid_hours_per_unit", explicitly_separate=1)])
    db.execute("UPDATE assembly_lines SET explicitly_separate=0 WHERE assembly_pk=(SELECT pk FROM assemblies WHERE code='A-DBL')")   # bypass approval-time check
    db.execute("UPDATE boq_items SET assembly_code='A-DBL' WHERE project_id=? AND item_key='1.1'", (pid,))
    fs = findings(db, pid)
    assert any(f["rule_code"] == "DOUBLE_COUNT" and f["severity"] == "Critical" for f in fs)


def test_scope_mismatches_exclusion_vs_priced(mini):
    db, pid, _ = mini
    exc = registers.add_scope_entry(db, ESTIMATOR, pid, "exclusion", "Concrete B excluded", boq_item_key="1.2", rationale="by client")
    registers.approve_scope_entry(db, ESTIMATOR, exc)
    assert ("PRICED_BUT_EXCLUDED", exc) in {(f["rule_code"], f["entity_id"]) for f in findings(db, pid)}
    db.execute("UPDATE boq_items SET pricing_method='excluded', assembly_code=NULL WHERE project_id=? AND item_key='1.2'", (pid,))
    fs = findings(db, pid)
    assert "PRICED_BUT_EXCLUDED" not in codes(fs) and "EXCLUSION_UNAPPROVED" not in codes(fs)
    db.execute("UPDATE scope_register SET status='Draft'")
    assert "EXCLUSION_UNAPPROVED" in codes(findings(db, pid))                       # excluded without an APPROVED exclusion


def test_missing_boq_items_in_revision_and_pending_changes(mini):
    db, pid, tmp = mini
    bi.import_boq(db, ESTIMATOR, pid, write_boq(tmp / "r2.csv", [("1.1", "Concrete", "m3", "12", "", "")]), revision="B")
    fs = findings(db, pid)
    assert ("BOQ_ITEM_MISSING_IN_REVISION", "1.2") in {(f["rule_code"], f["entity_id"]) for f in fs}
    assert "BOQ_CHANGE_PENDING" in codes(fs)
    assert recon.final_gate(db, ESTIMATOR, pid)["blocked"]


def test_offer_and_pricing_schedule_discrepancies(mini):
    db, pid, _ = mini
    o = offers.generate_offer(db, ESTIMATOR, pid, payment_terms="SYNTHETIC")
    assert "OFFER_MISSING" not in codes(findings(db, pid)) and "OFFER_STALE" not in codes(findings(db, pid))
    inc = registers.add_scope_entry(db, ESTIMATOR, pid, "inclusion", "Extra works", cost_code="Z9")
    registers.approve_scope_entry(db, ESTIMATOR, inc)
    assert "OFFER_STALE" in codes(findings(db, pid))                                # scope changed after the offer
    o2 = offers.generate_offer(db, ESTIMATOR, pid)
    assert "OFFER_STALE" not in codes(findings(db, pid))
    assert any(f["rule_code"] == "OFFER_COMMITMENT_NOT_PRICED" for f in findings(db, pid))      # inclusion for cost code Z9 has no priced scope
    db.execute("UPDATE boq_items SET quantity='11' WHERE project_id=? AND item_key='1.1'", (pid,))
    fs = findings(db, pid)
    st = [f for f in fs if f["rule_code"] == "OFFER_STALE"]
    assert st and st[0]["severity"] == "Critical" and st[0]["financial_impact"]
    with pytest.raises(Exception):
        offers.approve_offer(db, COMMERCIAL, pid, o2["version_no"])                 # cannot approve a stale offer


def test_totals_that_do_not_reconcile_are_caught(mini, monkeypatch):
    db, pid, _ = mini
    est = compute_estimate(db, pid)
    est.schedule.lines[0].allocated_indirect += Decimal("50")                         # corrupt the allocation
    fs = recon.collect_findings(db, pid, est)
    assert any(f.rule_code == "TOTALS_DO_NOT_RECONCILE" for f in fs)
    est2 = compute_estimate(db, pid)
    est2.schedule.total_pre_tax += Decimal("500")                                      # schedule drifts from the internal estimate
    fs2 = recon.collect_findings(db, pid, est2)
    assert any(f.rule_code in ("SCHEDULE_RECON", "SCHEDULE_RECON_MAJOR") for f in fs2)


def test_exceptions_policy_and_roles(mini):
    db, pid, _ = mini
    bi.import_boq(db, ESTIMATOR, pid, write_boq(mini[2] / "rej.csv", [("3.1", "x", "m", "TBA", "", "")]))
    fs = findings(db, pid)
    rej = next(f for f in fs if f["rule_code"] == "BOQ_ROWS_REJECTED")
    assert rej["severity"] == "Medium"
    with pytest.raises(ValidationFailed):
        recon.accept_exception(db, ESTIMATOR, rej["id"], "short")                       # documented reason required
    recon.accept_exception(db, ESTIMATOR, rej["id"], "Row 'TBA' is a placeholder for a later addendum")
    fs2 = findings(db, pid)
    assert next(f for f in fs2 if f["rule_code"] == "BOQ_ROWS_REJECTED")["status"] == "Accepted Exception"      # persists across runs by fingerprint
    ex = next(f for f in fs2 if f["rule_code"] == "BOQ_ROWS_REJECTED")
    assert ex["exception_reason"] and ex["exception_by"] == "est"
    # Critical / non-excusable rules can never be waived, even by an administrator
    create_resource(db, ESTIMATOR, resource_type="material", description="nr", unit="m", id="NR-2", is_synthetic=True)
    approved_assembly(db, "A-NR2", "m", [dict(line_type="material", resource_id="NR-2", factor="1")])
    bi.import_boq(db, ESTIMATOR, pid, write_boq(mini[2] / "m2.csv", [("4.1", "nr", "m", "1", "", "")]))
    boq.map_to_assembly(db, ESTIMATOR, pid, "4.1", "A-NR2")
    crit = next(f for f in findings(db, pid) if f["severity"] == "Critical")
    for who in (ADMIN, COMMERCIAL):
        with pytest.raises(PermissionDenied) as e:
            recon.accept_exception(db, who, crit["id"], "I really want this waived please")
        assert e.value.code == "EXCEPTION_NOT_ALLOWED"


def test_high_exception_needs_commercial_role_and_configurable_rules(mini):
    db, pid, _ = mini
    cid = registers.add_clarification(db, ESTIMATOR, pid, "Which concrete grade?", boq_item_key="1.1", cost_exposure="5000")
    f = next(f for f in findings(db, pid) if f["rule_code"] == "CLARIFICATION_OPEN")
    assert f["severity"] == "High"
    with pytest.raises(PermissionDenied) as e:
        recon.accept_exception(db, ESTIMATOR, f["id"], "client will answer after award per addendum")
    assert e.value.code == "EXCEPTION_ROLE"
    recon.accept_exception(db, COMMERCIAL, f["id"], "client will answer after award per addendum")
    accept_adopted_assumptions(db, pid, findings(db, pid))
    assert not recon.final_gate(db, ESTIMATOR, pid)["blocked"]
    # severities are configurable
    db.insert("settings", dict(key="recon_rules_override", value='{"rules": {"CLARIFICATION_OPEN": {"severity": "Critical", "exception_allowed": false}}}'))
    r = load_rules(db)["rules"]["CLARIFICATION_OPEN"]
    assert r["severity"] == "Critical" and r["exception_allowed"] is False
    assert recon.final_gate(db, ESTIMATOR, pid)["blocked"]                           # Critical now; the old exception no longer applies


def test_version_difference_and_report(mini, tmp_path):
    db, pid, _ = mini
    estimates.freeze_version(db, ESTIMATOR, pid, "v1")
    db.execute("UPDATE boq_items SET quantity='20' WHERE project_id=? AND item_key='1.1'", (pid,))
    estimates.freeze_version(db, ESTIMATOR, pid, "v2")
    fs = findings(db, pid)
    assert "VERSION_DIFFERENCE" in codes(fs)
    from app.exports.documents import reconciliation_report_xlsx
    run = recon.get_run(db, project_id=pid)
    out = reconciliation_report_xlsx(run, projects.get_project(db, pid), tmp_path / "rec.xlsx")
    import openpyxl
    ws = openpyxl.load_workbook(out).active
    assert ws["A4"].value == "Severity" and ws.max_row >= 5 and ws.freeze_panes == "A5"
