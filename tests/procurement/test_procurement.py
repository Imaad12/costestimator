from datetime import date, timedelta
from decimal import Decimal

import openpyxl
import pytest

from app.core.errors import ApprovalBlocked, Conflict, PermissionDenied, ValidationFailed
from app.estimators.estimate import compute_estimate
from app.imports import boq as bi
from app.procurement import leveling, needs, quotes, rfq, selection
from app.services import boq, commercial, projects, rates, suppliers
from app.services.rates import create_resource

from tests.conftest import ADMIN, COMMERCIAL, ESTIMATOR, PROCUREMENT, VIEWER
from tests.helpers import approved_assembly, approved_rate, write_boq


@pytest.fixture()
def scen(seeded, tmp_path):
    db = seeded
    create_resource(db, ESTIMATOR, resource_type="material", description="Pipe 50mm (SYNTHETIC)", unit="m", id="P-1", cost_code="G.01", is_synthetic=True,
                    specification="Black steel, ASTM A53 grade B (to confirm)")
    create_resource(db, ESTIMATOR, resource_type="equipment", description="Pump set (SYNTHETIC)", unit="nr", id="P-2", cost_code="G.03", is_synthetic=True)
    create_resource(db, ESTIMATOR, resource_type="material", description="Priced item (SYNTHETIC)", unit="m", id="P-3", cost_code="C.01", is_synthetic=True)
    for rid, u in (("P-1", "m"), ("P-2", "nr")):
        rates.create_missing_placeholder(db, rid, unit=u)
    approved_rate(db, "P-3", "5")
    approved_assembly(db, "A-P1", "m", [dict(line_type="material", resource_id="P-1", factor="1"), dict(line_type="material", resource_id="P-3", factor="1")], cost_code="G.01")
    approved_assembly(db, "A-P2", "nr", [dict(line_type="equipment", resource_id="P-2", factor="1")], cost_code="G.03")
    pid = projects.create_project(db, ESTIMATOR, name="Proc (SYNTHETIC)", location="SYNTHETIC SITE", is_demo=True)
    bi.import_boq(db, ESTIMATOR, pid, write_boq(tmp_path / "b.csv", [("1", "Pipework", "m", "100", "", ""), ("2", "Pumps", "nr", "2", "", "")]))
    boq.map_to_assembly(db, ESTIMATOR, pid, "1", "A-P1")
    boq.map_to_assembly(db, ESTIMATOR, pid, "2", "A-P2")
    r = needs.detect_needs(db, ESTIMATOR, pid)
    sa = suppliers.create_supplier(db, ESTIMATOR, name="DEMO Supplier A [SYNTHETIC]", is_synthetic=True)
    sb = suppliers.create_supplier(db, ESTIMATOR, name="DEMO Supplier B [SYNTHETIC]", is_synthetic=True)
    items = {i["resource_id"]: i for i in needs.register(db, pid)}
    return dict(db=db, pid=pid, sa=sa, sb=sb, items=items, pkg=items["P-1"]["package_id"])


def test_missing_rates_become_procurement_requirements(scen):
    db, items = scen["db"], scen["items"]
    assert set(items) == {"P-1", "P-2"}
    p1 = items["P-1"]
    assert p1["reason"].startswith("Missing rate") and Decimal(p1["quantity"]) == 100 and p1["unit"] == "m" and p1["estimated_exposure"] is None   # unknown, not 0
    assert p1["specification"].startswith("Black steel") and p1["delivery_location"] == "SYNTHETIC SITE" and p1["rfq_status"] == "Not Issued"
    assert p1["package_id"] == "PKG-%s-G" % scen["pid"]
    again = needs.detect_needs(db, ESTIMATOR, scen["pid"])
    assert again["created"] == 0 and again["updated"] == 2                                  # idempotent


def test_unapproved_and_expired_rates_detected_with_indicative_exposure(scen):
    db, pid = scen["db"], scen["pid"]
    rates.propose_rate(db, ESTIMATOR, "P-1", unit_price="40", rate_type="Synthetic Demonstration Rate", price_date="2026-09-01", is_synthetic=True, evidence="x", submit=True)
    needs.detect_needs(db, ESTIMATOR, pid)
    it = next(i for i in needs.register(db, pid) if i["resource_id"] == "P-1")
    assert it["reason_code"] == "UNAPPROVED_RATE" and Decimal(it["estimated_exposure"]) == Decimal("4000")      # 100 m x 40, shown as exposure only
    assert compute_estimate(db, pid).direct.direct_cost_total == Decimal("500")                               # only P-3 (100 x 5) is in the total
    # the old MISSING item closes (replaced by the more specific reason)
    assert not any(i["reason_code"] == "MISSING_RATE" and i["resource_id"] == "P-1" for i in needs.register(db, pid))


def test_rfq_generation_content_and_exports(scen, tmp_path):
    db = scen["db"]
    rid = rfq.create_rfq(db, ESTIMATOR, scen["pid"], scen["pkg"], supplier_id=scen["sa"], response_deadline="2026-11-15")
    r = rfq.get_rfq(db, rid)
    for must in ("Pipe 50mm", "100", "SYNTHETIC SITE", "2026-11-15", "validity", "Payment terms", "unloading", "datasheets", "deviations", "SAR"):
        assert must.lower() in r["body"].lower(), must
    assert "@" not in r["body"].split("Kind regards")[0]                                  # never invents emails
    assert r["status"] == "Draft" and len(r["items"]) == 2
    assert db.scalar("SELECT rfq_status FROM procurement_items WHERE resource_id='P-1'") == "Draft RFQ"
    x = rfq.export_xlsx(db, rid, tmp_path / "r.xlsx")
    ws = openpyxl.load_workbook(x)["RFQ"]
    rows = {ws.cell(i, 1).value: ws.cell(i, 4).value for i in range(6, 8)}
    assert rows[scen["items"]["P-1"]["id"]] == "100" and rows[scen["items"]["P-2"]["id"]] == "2"
    assert rfq.export_docx(db, rid, tmp_path / "r.docx").stat().st_size > 5000
    assert rfq.email_text(db, rid).startswith("Subject: Request for Quotation")
    rfq.mark_issued(db, ESTIMATOR, rid)
    assert rfq.get_rfq(db, rid)["status"] == "Issued (manually)"                          # human-recorded; nothing is ever sent
    with pytest.raises(PermissionDenied):
        rfq.create_rfq(db, VIEWER, scen["pid"], scen["pkg"])


def test_arabic_rfq_template_and_unknown_language(scen):
    db = scen["db"]
    rid = rfq.create_rfq(db, ESTIMATOR, scen["pid"], scen["pkg"], language="ar")
    assert "طلب عرض سعر" in rfq.get_rfq(db, rid)["subject"]
    with pytest.raises(ValidationFailed):
        rfq.create_rfq(db, ESTIMATOR, scen["pid"], scen["pkg"], language="fr")


def test_supplier_contacts_need_a_source(seeded):
    with pytest.raises(ValidationFailed) as e:
        suppliers.create_supplier(seeded, ESTIMATOR, name="X", contact_email="a@b.example")
    assert e.value.code == "CONTACT_SOURCE_REQUIRED"
    sid = suppliers.create_supplier(seeded, ESTIMATOR, name="Y", contact_email="a@b.example", contact_source="Supplier's public catalogue, 2026-10-01")
    assert suppliers.get_supplier(seeded, sid)["contact_source"]


def mkquote(scen, supplier, p1_rate, p2_rate=None, *, tax="excl", rate_tax=None, valid_days=30, ccy="SAR", delivery="Yes", unit1="m", disc=None, **kw):
    db = scen["db"]
    it1, it2 = scen["items"]["P-1"]["id"], scen["items"]["P-2"]["id"]
    lines = [dict(procurement_item_id=it1, description="Pipe", quantity="100", unit=unit1, unit_rate=p1_rate, delivery_included=delivery, discount_pct=disc, compliance="Compliant")]
    if p2_rate:
        lines.append(dict(procurement_item_id=it2, description="Pump", quantity="2", unit="nr", unit_rate=p2_rate, compliance="Compliant"))
    return quotes.create_quotation(db, ESTIMATOR, supplier_id=supplier, project_id=scen["pid"], package_id=scen["pkg"], reference="Q-1", quote_date=date.today().isoformat(),
                                   valid_until=(date.today() + timedelta(days=valid_days)).isoformat(), currency=ccy, tax_treatment=tax, tax_rate_pct=rate_tax, lines=lines,
                                   is_synthetic=True, **kw)


def approve_q(db, qid):
    quotes.submit_quotation(db, ESTIMATOR, qid)
    quotes.approve_quotation(db, PROCUREMENT, qid, evidence="checked against synthetic document")


def test_quote_import_from_completed_rfq_sheet(scen, tmp_path):
    db = scen["db"]
    rid = rfq.create_rfq(db, ESTIMATOR, scen["pid"], scen["pkg"], supplier_id=scen["sa"])
    x = rfq.export_xlsx(db, rid, tmp_path / "r.xlsx")
    wb = openpyxl.load_workbook(x)
    ws = wb["RFQ"]
    row = next(i for i in range(6, 8) if ws.cell(i, 1).value == scen["items"]["P-1"]["id"])
    ws.cell(row, 6).value, ws.cell(row, 7).value, ws.cell(row, 8).value = 42.5, 14, "Equal approved"
    wb.save(x)
    qid = quotes.create_quotation(db, ESTIMATOR, supplier_id=scen["sa"], project_id=scen["pid"], package_id=scen["pkg"], quote_date="2026-10-01", valid_until="2026-12-01",
                                  tax_treatment="excl", is_synthetic=True)
    res = quotes.import_quotation_lines(db, ESTIMATOR, qid, x)
    assert res["added"] == 1
    ln = quotes.get_quotation(db, qid)["lines"][0]
    assert ln["procurement_item_id"] == scen["items"]["P-1"]["id"] and ln["unit_rate"] == "42.5" and ln["lead_time_days"] == "14" and ln["comments"] == "Equal approved"
    wb2 = openpyxl.load_workbook(x)
    wb2["RFQ"].cell(row, 6).value = "abc"
    wb2.save(x)
    qid2 = quotes.create_quotation(db, ESTIMATOR, supplier_id=scen["sb"], project_id=scen["pid"], package_id=scen["pkg"], quote_date="2026-10-01", tax_treatment="excl")
    r2 = quotes.import_quotation_lines(db, ESTIMATOR, qid2, x)
    assert r2["added"] == 0 and "not numeric" in r2["skipped"][0]["reason"]
    assert quotes.get_quotation(db, qid)["status"] == "Draft"                                  # imported quotes are never auto-approved


def test_quote_validity_and_expiry(scen):
    db = scen["db"]
    q_ok = mkquote(scen, scen["sa"], "40")
    q_old = mkquote(scen, scen["sb"], "39", valid_days=-5)
    assert quotes.validity(quotes.get_quotation(db, q_ok)) == "Valid" and quotes.validity(quotes.get_quotation(db, q_old)) == "Expired"
    assert quotes.validity({"valid_until": None}) == "No validity stated"
    approve_q(db, q_old)
    assert quotes.mark_expired(db) == 1 and quotes.get_quotation(db, q_old)["status"] == "Expired"
    lv = leveling.level_package(db, scen["pkg"])
    expired = next(q for q in lv["quotations"] if q["quotation_id"] == q_old)
    assert expired["validity"] == "Expired" and not expired["like_for_like"]


def test_quote_edit_rules_and_approval(scen):
    db = scen["db"]
    q = mkquote(scen, scen["sa"], "40")
    with pytest.raises(PermissionDenied):
        quotes.submit_quotation(db, VIEWER, q)
    quotes.submit_quotation(db, ESTIMATOR, q)
    with pytest.raises(PermissionDenied):
        quotes.approve_quotation(db, ESTIMATOR, q)                                             # estimator cannot verify supplier quotes
    quotes.approve_quotation(db, PROCUREMENT, q, evidence="ok")
    with pytest.raises(Conflict):
        quotes.add_line(db, ESTIMATOR, q, description="late line", unit_rate="1")               # approved = immutable
    with pytest.raises(ValidationFailed):
        quotes.add_line(db, ESTIMATOR, mkquote(scen, scen["sb"], "41"), unit="furlong", unit_rate="1")
    ai = quotes.create_quotation(db, ESTIMATOR, supplier_id=scen["sa"], project_id=scen["pid"], package_id=scen["pkg"], quote_date="2026-10-01",
                                 extracted_by_ai=True, lines=[dict(description="x", unit_rate="1")])
    quotes.submit_quotation(db, ESTIMATOR, ai)
    with pytest.raises(ValidationFailed) as e:
        quotes.approve_quotation(db, PROCUREMENT, ai)
    assert e.value.code == "AI_EXTRACTION_UNVERIFIED"


def test_leveling_normalisation_tax_currency_units(scen):
    db = scen["db"]
    a = mkquote(scen, scen["sa"], "40", "1000", tax="excl")                                     # 4000 + 2000 = 6000
    b = mkquote(scen, scen["sb"], "36.8", "920", tax="incl", rate_tax="15")                     # incl 15%: 3680/1.15=3200 ; 1840/1.15=1600 => 4800
    for q in (a, b):
        approve_q(db, q)
    lv = leveling.level_package(db, scen["pkg"])
    qa = next(x for x in lv["quotations"] if x["quotation_id"] == a)
    qb = next(x for x in lv["quotations"] if x["quotation_id"] == b)
    assert qa["quoted_total"] == Decimal("6000.00") and qa["normalised_total"] == Decimal("6000.00")
    assert qb["quoted_total"] == Decimal("5520.00") and qb["normalised_total"] == Decimal("4800.00")       # actual quote vs adjusted comparison value
    assert qa["like_for_like"] and qb["like_for_like"] and lv["suggested_lowest_like_for_like"] == b and "Advisory" in lv["advisory"]
    # a quotation without a stated tax treatment cannot be normalised
    c = mkquote(scen, scen["sa"], "10", "100", tax=None)
    lv = leveling.level_package(db, scen["pkg"])
    qc = next(x for x in lv["quotations"] if x["quotation_id"] == c)
    assert qc["normalised_total"] is None and not qc["like_for_like"] and any("Tax treatment" in n for n in qc["notes"])
    # foreign currency needs an approved FX rate; the original currency is preserved
    d = mkquote(scen, scen["sb"], "10", "250", ccy="USD")
    qd = next(x for x in leveling.level_package(db, scen["pkg"])["quotations"] if x["quotation_id"] == d)
    assert qd["currency"] == "USD" and qd["normalised_total"] is None and qd["quoted_total"] == Decimal("1500.00")
    fx = commercial.propose_fx(db, ESTIMATOR, "USD", "SAR", "3.75", "2026-01-01", "SAMA peg (SYNTHETIC test)")
    assert next(x for x in leveling.level_package(db, scen["pkg"])["quotations"] if x["quotation_id"] == d)["normalised_total"] is None   # pending FX is not used
    commercial.approve_fx(db, PROCUREMENT, fx)
    qd = next(x for x in leveling.level_package(db, scen["pkg"])["quotations"] if x["quotation_id"] == d)
    assert qd["normalised_total"] == Decimal("5625.00")                                                  # 1500 x 3.75
    # unit conversion: quoted per km but required per m
    e = mkquote(scen, scen["sa"], "40000", "1000", unit1="km")
    qe = next(x for x in leveling.level_package(db, scen["pkg"])["quotations"] if x["quotation_id"] == e)
    assert qe["normalised_total"] == Decimal("6000.00")                                                  # 100 m x (40000/1000)=4000 + 2000
    f = mkquote(scen, scen["sa"], "10", "100", unit1="m2")
    qf = next(x for x in leveling.level_package(db, scen["pkg"])["quotations"] if x["quotation_id"] == f)
    assert qf["unit_issues"] and not qf["like_for_like"]


def test_leveling_missing_items_quantity_difference_discount_and_adjustments(scen):
    db = scen["db"]
    part = mkquote(scen, scen["sa"], "30", None, disc="10")                                    # only P-1, 10% discount: 100 x 30 x 0.9 = 2700
    approve_q(db, part)
    lv = leveling.level_package(db, scen["pkg"])
    q = lv["quotations"][0]
    assert q["coverage"] == "1/2" and q["missing_items"] == [scen["items"]["P-2"]["id"]] and not q["like_for_like"] and q["normalised_total"] == Decimal("2700.00")
    assert lv["suggested_lowest_like_for_like"] is None                                         # an incomplete bid is never presented as the cheapest like-for-like
    adj = leveling.add_adjustment(db, ESTIMATOR, part, "missing_item", "1800", "Add pump at the median of other bids (SYNTHETIC)")
    assert leveling.level_package(db, scen["pkg"])["quotations"][0]["normalised_total"] == Decimal("2700.00")        # pending adjustments are not applied
    with pytest.raises(PermissionDenied):
        leveling.approve_adjustment(db, ESTIMATOR, adj)
    leveling.approve_adjustment(db, PROCUREMENT, adj)
    q = leveling.level_package(db, scen["pkg"])["quotations"][0]
    assert q["normalised_total"] == Decimal("4500.00") and q["quoted_total"] == Decimal("2700.00")
    with pytest.raises(ValidationFailed):
        leveling.add_adjustment(db, ESTIMATOR, part, "other", "5", "")


def test_supplier_selection_needs_human_approval_and_reasons(scen):
    db = scen["db"]
    a = mkquote(scen, scen["sa"], "40", "1000")                                                 # 6000
    b = mkquote(scen, scen["sb"], "30", "900")                                                  # 4800 (cheapest, like-for-like)
    for q in (a, b):
        approve_q(db, q)
    with pytest.raises(ValidationFailed) as e:
        selection.propose_selection(db, ESTIMATOR, scen["pkg"], a, "preferred supplier")
    assert e.value.code == "NOT_LOWEST_REASON_REQUIRED"
    s = selection.propose_selection(db, ESTIMATOR, scen["pkg"], a, "better lead time and local support", not_lowest_reason="B has 10 week lead time vs A 3 weeks")
    assert db.scalar("SELECT status FROM quote_selections WHERE id=?", (s,)) == "Pending Review"
    with pytest.raises(PermissionDenied):
        selection.decide_selection(db, ESTIMATOR, s, True)
    with pytest.raises(ValidationFailed):
        selection.publish_rates(db, ESTIMATOR, s)                                               # not approved yet
    selection.decide_selection(db, PROCUREMENT, s, True, "agreed")
    assert db.scalar("SELECT quotation_status FROM procurement_items WHERE resource_id='P-1'") == "Selected"
    # an unapproved quotation cannot be selected at all
    c = mkquote(scen, scen["sb"], "20", "500")
    with pytest.raises(ValidationFailed) as e:
        selection.propose_selection(db, ESTIMATOR, scen["pkg"], c, "cheapest")
    assert e.value.code == "QUOTE_NOT_APPROVED"


def test_high_value_selection_requires_commercial_approver(scen, settings):
    db = scen["db"]
    settings.approval_thresholds = {"supplier_selection_commercial_approver_sar": "5000"}
    a = mkquote(scen, scen["sa"], "40", "1000")                                                 # 6000 >= 5000
    approve_q(db, a)
    s = selection.propose_selection(db, ESTIMATOR, scen["pkg"], a, "only bidder")
    with pytest.raises(PermissionDenied) as e:
        selection.decide_selection(db, PROCUREMENT, s, True)
    assert e.value.code == "THRESHOLD_ROLE_REQUIRED"
    selection.decide_selection(db, COMMERCIAL, s, True)


def test_rate_publication_requires_rate_approval_and_links_quotation(scen):
    db, pid = scen["db"], scen["pid"]
    a = mkquote(scen, scen["sa"], "40", "1000")
    approve_q(db, a)
    s = selection.propose_selection(db, ESTIMATOR, scen["pkg"], a, "only bidder")
    selection.decide_selection(db, PROCUREMENT, s, True)
    ids = selection.publish_rates(db, ESTIMATOR, s)
    assert len(ids) == 2
    r = rates.get_rate(db, ids[0])
    assert r["status"] == "Pending Review" and r["quotation_line_id"] and r["supplier_id"] == scen["sa"] and r["price_date"] and r["valid_until"]
    assert r["rate_type"] == "Synthetic Demonstration Rate" and r["is_synthetic"] == 1                    # synthetic quotes stay synthetic
    with pytest.raises(Conflict):
        selection.publish_rates(db, ESTIMATOR, s)                                                    # only once
    assert compute_estimate(db, pid).direct.direct_cost_total == Decimal("500")                      # not usable until the rates are approved
    for i in ids:
        rates.approve_rate(db, PROCUREMENT, i)
    est = compute_estimate(db, pid)
    assert est.direct.direct_cost_total == Decimal("500") + Decimal("4000") + Decimal("2000")
    assert [i for i in needs.register(db, pid)] and needs.detect_needs(db, ESTIMATOR, pid)["open"] == 0   # needs close once rates are approved
    # tax-inclusive quote: the published rate is derived pre-tax
    scen2 = dict(scen)


def test_published_rate_from_tax_inclusive_quote_is_pre_tax(scen):
    db = scen["db"]
    b = mkquote(scen, scen["sb"], "46", "1150", tax="incl", rate_tax="15")
    approve_q(db, b)
    s = selection.propose_selection(db, ESTIMATOR, scen["pkg"], b, "only bidder")
    selection.decide_selection(db, PROCUREMENT, s, True)
    ids = selection.publish_rates(db, ESTIMATOR, s)
    assert sorted(Decimal(rates.get_rate(db, i)["unit_price"]) for i in ids) == [Decimal("40"), Decimal("1000")]
    no_tax = mkquote(scen, scen["sa"], "10", tax=None)
    approve_q(db, no_tax)
    s2 = selection.propose_selection(db, ESTIMATOR, scen["pkg"], no_tax, "x", not_lowest_reason="incomplete tax data accepted for test")
    selection.decide_selection(db, PROCUREMENT, s2, True)
    with pytest.raises(ValidationFailed) as e:
        selection.publish_rates(db, ESTIMATOR, s2)
    assert e.value.code == "TAX_BASIS_UNCLEAR"
