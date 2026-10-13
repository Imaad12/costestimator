"""Regression tests for defects found by the capability audit (docs/ESTIMATING_CAPABILITY_MATRIX.md)."""
from __future__ import annotations

import ast
import time
from pathlib import Path

import pytest

from app.services import project_index

ROOT = Path(__file__).resolve().parent.parent.parent


def _minimal_pdf(text: str | None) -> bytes:
    content = (f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()) if text else b""
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    x = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1) + b"".join(b"%010d 00000 n \n" % o for o in offs)
    return out + b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, x)


def test_pypdf_is_a_declared_dependency_so_pdfs_are_extracted_on_a_default_install():
    assert "pypdf" in (ROOT / "pyproject.toml").read_text(encoding="utf-8")
    import pypdf  # noqa: F401  (the test environment must satisfy the declared dependency)


def test_real_pdf_docx_and_xlsx_text_is_extracted_and_image_only_pdf_is_reported(tmp_path):
    (tmp_path / "spec.pdf").write_bytes(_minimal_pdf("Concrete grade C30 cover 50 mm"))
    txt, why = project_index.extract_text(tmp_path / "spec.pdf")
    assert why is None and "Concrete grade C30 cover 50 mm" in txt and "=== page 1 ===" in txt
    (tmp_path / "scan.pdf").write_bytes(_minimal_pdf(None))
    txt, why = project_index.extract_text(tmp_path / "scan.pdf")
    assert txt is None and "no extractable text" in why and "OCR" in why                    # scanned PDFs are reported, never invented
    import docx
    d = docx.Document()
    d.add_paragraph("Clause 3.1 the Contractor shall comply")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Value"
    d.save(tmp_path / "s.docx")
    txt, why = project_index.extract_text(tmp_path / "s.docx")
    assert why is None and "Contractor shall comply" in txt and "Item | Value" in txt
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(["a", "b"])
    wb.active.append([1, 2])
    wb.save(tmp_path / "x.xlsx")
    txt, why = project_index.extract_text(tmp_path / "x.xlsx")
    assert why is None and "1 | 2" in txt


def test_estimate_computation_scales_roughly_linearly_with_boq_size(tmp_path, seeded):
    """Regression: an O(n^2) per-item tax-category lookup made a 20,000-item estimate take ~14 s instead of ~3 s."""
    from app.estimators.estimate import compute_estimate
    from app.imports import boq as bi
    from app.services import boq, projects
    from app.services.rates import create_resource
    from tests.conftest import ESTIMATOR
    from tests.helpers import approved_assembly, approved_rate, std_policy, write_boq
    db = seeded
    create_resource(db, ESTIMATOR, resource_type="material", description="DEMO mat", unit="m3", cost_code="C.01", id="DEMO-MAT-001", is_synthetic=True)
    approved_rate(db, "DEMO-MAT-001", "100")
    approved_assembly(db, "DEMO-A", "m3", [dict(line_type="material", resource_id="DEMO-MAT-001", factor="1")])

    def build(n: int) -> str:
        pid = projects.create_project(db, ESTIMATOR, name=f"scale {n} (SYNTHETIC)", is_demo=True, tax_category="VAT-STD")
        rows = [(f"{i // 100 + 1}.{i % 100 + 1}", f"Item {i}", "m3", "5", "", "") for i in range(n)]
        bi.import_boq(db, ESTIMATOR, pid, write_boq(tmp_path / f"b{n}.csv", rows))
        for i in range(n):
            boq.map_to_assembly(db, ESTIMATOR, pid, f"{i // 100 + 1}.{i % 100 + 1}", "DEMO-A")
        std_policy(db, pid)
        return pid

    def best_of(pid: str, k: int = 2) -> float:
        ts = []
        for _ in range(k):
            t = time.perf_counter()
            compute_estimate(db, pid)
            ts.append(time.perf_counter() - t)
        return min(ts)

    small, large = build(2000), build(12000)
    t_small, t_large = best_of(small), best_of(large)
    # linear => ratio ~6; the old quadratic code gave ~15-20. Generous bound against timer noise.
    assert t_large / max(t_small, 1e-3) < 11, (t_small, t_large)


FORBIDDEN_IMPORTS = {"smtplib", "ftplib", "poplib", "imaplib", "telnetlib", "xmlrpc", "paramiko", "selenium", "playwright"}


def test_application_has_no_autonomous_submission_channel():
    """The suite only produces files for a human to send: no e-mail/FTP/portal-automation libraries, and the only HTTP clients are
    the read-only source fetcher (GET) and the optional AI provider."""
    offenders = []
    for f in sorted((ROOT / "app").rglob("*.py")):
        tree = ast.parse(f.read_text(encoding="utf-8"))
        for n in ast.walk(tree):
            mods = []
            if isinstance(n, ast.Import):
                mods = [a.name.split(".")[0] for a in n.names]
            elif isinstance(n, ast.ImportFrom) and n.module:
                mods = [n.module.split(".")[0]]
            offenders += [f"{f.relative_to(ROOT)}:{n.lineno} imports {m}" for m in mods if m in FORBIDDEN_IMPORTS]
    assert not offenders, offenders
    src = (ROOT / "app" / "sources" / "base.py").read_text(encoding="utf-8")
    assert 'method="POST"' not in src and "method='POST'" not in src and "data=" not in src.split("def _once", 1)[1].split("def ", 1)[0].replace("max_bytes", "")


def test_non_sar_project_currency_needs_an_approved_fx_rate_then_converts(tmp_path, seeded):
    """Multi-currency capability (audit record CUR-02): a SAR rate in a USD project is unpriced until an APPROVED SAR->USD rate exists."""
    from decimal import Decimal
    from app.estimators.estimate import compute_estimate
    from app.imports import boq as bi
    from app.services import boq, commercial, projects
    from app.services.rates import create_resource
    from tests.conftest import ESTIMATOR, PROCUREMENT
    from tests.helpers import approved_assembly, approved_rate, std_policy, write_boq
    db = seeded
    create_resource(db, ESTIMATOR, resource_type="material", description="DEMO mat", unit="m3", cost_code="C.01", id="DEMO-MAT-001", is_synthetic=True)
    approved_rate(db, "DEMO-MAT-001", "100")
    approved_assembly(db, "DEMO-A", "m3", [dict(line_type="material", resource_id="DEMO-MAT-001", factor="1")])
    pid = projects.create_project(db, ESTIMATOR, name="USD project (SYNTHETIC)", is_demo=True, currency="USD", tax_category="VAT-STD")
    bi.import_boq(db, ESTIMATOR, pid, write_boq(tmp_path / "u.csv", [("1", "x", "m3", "10", "", "")]))
    boq.map_to_assembly(db, ESTIMATOR, pid, "1", "DEMO-A")
    std_policy(db, pid)
    est = compute_estimate(db, pid)
    assert est.direct.direct_cost_total == 0 and any(e["code"] == "FX_MISSING" for e in est.exceptions) and est.state.startswith("Preliminary")
    pending = commercial.propose_fx(db, ESTIMATOR, "SAR", "USD", "0.30", "2026-09-01", "SYNTHETIC test rate")
    assert compute_estimate(db, pid).direct.direct_cost_total == 0                              # a Pending Review FX rate is not used
    commercial.approve_fx(db, PROCUREMENT, pending)
    est = compute_estimate(db, pid)
    assert est.direct.currency == "USD" and est.direct.direct_cost_total == Decimal("300.00")   # 10 m3 x SAR 100 x 0.30


def test_what_if_assumption_override_changes_the_estimate_by_the_expected_amount_and_leaves_stored_data_alone(tmp_path, seeded):
    """Audit record EST-09: the 'show effect' analysis was only render-tested; this proves the numbers."""
    from decimal import Decimal
    from app.estimators.estimate import compute_estimate
    from app.imports import boq as bi
    from app.services import boq, projects
    from app.services.rates import create_resource
    from tests.conftest import ESTIMATOR
    from tests.helpers import approve_assumption, approved_assembly, approved_rate, write_boq
    db = seeded
    create_resource(db, ESTIMATOR, resource_type="material", description="DEMO mat", unit="m3", cost_code="C.01", id="DEMO-MAT-001", is_synthetic=True)
    approved_rate(db, "DEMO-MAT-001", "100")
    approve_assumption(db, "ASMP-C-003", "5")                                              # wastage 5 %
    approved_assembly(db, "DEMO-W", "m3", [dict(line_type="material", resource_id="DEMO-MAT-001", factor="1", waste_basis="assumption", waste_ref="ASMP-C-003")])
    pid = projects.create_project(db, ESTIMATOR, name="What-if (SYNTHETIC)", is_demo=True, tax_category="VAT-STD")
    bi.import_boq(db, ESTIMATOR, pid, write_boq(tmp_path / "w.csv", [("1", "x", "m3", "10", "", "")]))
    boq.map_to_assembly(db, ESTIMATOR, pid, "1", "DEMO-W")
    base = compute_estimate(db, pid).direct.direct_cost_total
    assert base == Decimal("1050.00")                                                      # 10 m3 x 1.05 x 100
    what_if = compute_estimate(db, pid, assumption_overrides={"ASMP-C-003": "6"}).direct.direct_cost_total
    assert what_if == Decimal("1060.00")                                                   # 10 m3 x 1.06 x 100
    assert compute_estimate(db, pid).direct.direct_cost_total == base                      # the override was not persisted anywhere


def test_takeoff_resource_lines_are_priced_from_approved_rates_and_missing_rates_stay_unpriced(tmp_path, seeded):
    """Audit record SCO-06: only the import of takeoff resource lines was tested; this proves the pricing."""
    from decimal import Decimal
    from app.estimators.estimate import compute_estimate
    from app.imports import boq as bi
    from app.services.rates import create_resource
    from app.services import projects
    from tests.conftest import ESTIMATOR
    from tests.helpers import approved_rate, write_boq
    db = seeded
    create_resource(db, ESTIMATOR, resource_type="material", description="DEMO rebar", unit="kg", id="TK-1", is_synthetic=True)
    create_resource(db, ESTIMATOR, resource_type="material", description="DEMO tie wire (no rate)", unit="kg", id="TK-2", is_synthetic=True)
    approved_rate(db, "TK-1", "2")
    pid = projects.create_project(db, ESTIMATOR, name="Takeoff (SYNTHETIC)", is_demo=True, tax_category="VAT-STD")
    bi.import_boq(db, ESTIMATOR, pid, write_boq(tmp_path / "t.csv", [("1", "Rebar", "t", "2", "", "")]))
    tk = write_boq(tmp_path / "tk.csv", [("1", "TK-1", "2100", "kg"), ("1", "TK-2", "30", "kg")], header=("Item", "Resource ID", "Quantity", "Unit"))
    assert bi.import_takeoff_resources(db, ESTIMATOR, pid, tk)["lines_added"] == 2
    est = compute_estimate(db, pid)
    assert est.direct.direct_cost_total == Decimal("4200.00")                        # 2,100 kg x SAR 2; no waste or factors added
    item = est.direct.items["1"]
    assert not item.complete and item.unpriced_lines == 1                           # the line without a rate is unpriced, not zero
    assert any(e.code in ("MISSING_RATE", "RATE_MISSING") or "MISSING" in e.code for e in est.direct.exceptions)


def test_indirect_that_overlaps_priced_preliminaries_is_flagged(tmp_path, seeded):
    """Audit record PRE-01: the POSSIBLE_DUPLICATE_INDIRECT rule was configured but never exercised by a test."""
    from app.imports import boq as bi
    from app.reconciliation import engine as recon
    from app.services import boq, indirects, projects
    from app.services.rates import create_resource
    from tests.conftest import COMMERCIAL, ESTIMATOR
    from tests.helpers import approve_vat, approved_assembly, approved_rate, std_policy, write_boq
    db = seeded
    create_resource(db, ESTIMATOR, resource_type="subcontract", description="DEMO security service", unit="month", cost_code="A.04", id="DEMO-SUB-SEC", is_synthetic=True)
    approved_rate(db, "DEMO-SUB-SEC", "1000")
    approved_assembly(db, "DEMO-SEC", "month", [dict(line_type="subcontract", resource_id="DEMO-SUB-SEC", factor="1")], cost_code="A.04", description="Site security guards")
    pid = projects.create_project(db, ESTIMATOR, name="Prelims (SYNTHETIC)", is_demo=True, tax_category="VAT-STD")
    bi.import_boq(db, ESTIMATOR, pid, write_boq(tmp_path / "p.csv", [("A.1", "Site security guards 24h", "month", "12", "", "")]))
    boq.map_to_assembly(db, ESTIMATOR, pid, "A.1", "DEMO-SEC")
    boq.set_cost_code(db, ESTIMATOR, pid, "A.1", "A.04")
    std_policy(db, pid)
    approve_vat(db)
    indirects.add_indirect(db, ESTIMATOR, pid, code="IND-SEC", category="Security", description="Site security", method="fixed", amount="5000", is_synthetic=True)
    indirects.submit_indirect(db, ESTIMATOR, pid, "IND-SEC")
    indirects.approve_indirect(db, COMMERCIAL, pid, "IND-SEC")
    run = recon.run_reconciliation(db, ESTIMATOR, pid)
    codes = {r["rule_code"] for r in db.all("SELECT rule_code FROM recon_findings WHERE run_id=?", (run["run_id"] if isinstance(run, dict) and "run_id" in run else db.scalar("SELECT MAX(id) FROM recon_runs"),))}
    assert "POSSIBLE_DUPLICATE_INDIRECT" in codes
