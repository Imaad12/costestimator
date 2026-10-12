import re
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from app.core.errors import ValidationFailed
from app.estimators.estimate import compute_estimate
from app.excel import roundtrip
from app.excel.master import REQUIRED_SHEETS, build_master
from app.excel.recalc import soffice_path, verify_workbook
from app.excel.spec import FIRST_ROW, HEADER_ROW, KEY_ROW, S
from app.excel.templates import TEMPLATE_FILES, generate_all
from app.services import assumptions, boq, indirects, projects, registers

from tests.conftest import COMMERCIAL, ESTIMATOR
from tests.helpers import build_mini

needs_lo = pytest.mark.skipif(soffice_path() is None, reason="LibreOffice not installed")


@pytest.fixture(scope="module")
def blank(tmp_path_factory):
    from app.database import Database
    from app.services import refdata
    import app.config as cfg
    db = Database(":memory:")
    db.migrate()
    refdata.seed_reference(db)
    p = tmp_path_factory.mktemp("wb") / "blank.xlsx"
    build_master(p, db, blank=True)
    return openpyxl.load_workbook(p), p, db


def test_required_sheets_present_and_ordered(blank):
    wb, _, _ = blank
    assert wb.sheetnames[:27] == REQUIRED_SHEETS and len(REQUIRED_SHEETS) == 27
    assert wb["Lists"].sheet_state == "hidden"
    assert wb.calculation.fullCalcOnLoad


def test_headers_keys_freeze_filters_for_every_spec_sheet(blank):
    wb, _, _ = blank
    for name, spec in S.items():
        ws = wb[name]
        assert [c.value for c in ws[HEADER_ROW]][:len(spec.cols)] == [c.header for c in spec.cols], name
        assert [c.value for c in ws[KEY_ROW]][:len(spec.cols)] == [c.key for c in spec.cols], name
        assert ws.row_dimensions[KEY_ROW].hidden
        assert ws.freeze_panes == "B6", name
        assert ws.auto_filter.ref and ws.auto_filter.ref.startswith("A5:"), name
        assert ws.protection.sheet, name


def test_data_validations_cover_status_type_unit_and_lists(blank):
    wb, _, _ = blank
    def dv_cols(ws):
        out = set()
        for dv in ws.data_validations.dataValidation:
            for r in str(dv.sqref).split():
                out.add(re.match(r"[A-Z]+", r).group())
        return out
    ws = wb["Direct Cost Build-Up"]
    assert wb["Assembly Library"].data_validations.dataValidation
    al = wb["Assembly Library"]
    cols = {c.key: openpyxl.utils.get_column_letter(i + 1) for i, c in enumerate(S["Assembly Library"].cols)}
    dvc = dv_cols(al)
    assert {cols["line_type"], cols["waste_basis"], cols["productivity_basis"], cols["unit"]} <= dvc
    lib = wb["Material Library"]
    lc = {c.key: openpyxl.utils.get_column_letter(i + 1) for i, c in enumerate(S["Material Library"].cols)}
    assert {lc["new_rate_type"], lc["unit"]} <= dv_cols(lib)
    lists = wb["Lists"]
    header = [c.value for c in lists[1]]
    assert "RateType" in header and "Status" in header
    status_vals = [lists.cell(r, header.index("Status") + 1).value for r in range(2, 12)]
    assert "Provisional — Approved for This Estimate" in status_vals and "Pending Review" in status_vals
    rt = [lists.cell(r, header.index("RateType") + 1).value for r in range(2, 12)]
    assert "Verified Supplier Quote" in rt and "Missing Rate" in rt and "Synthetic Demonstration Rate" in rt
    assert "m2" in [lists.cell(r, header.index("units") + 1).value for r in range(2, 60)]


def test_formulas_present_reference_existing_sheets_and_cells(blank):
    wb, _, _ = blank
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n += 1
                    for ref in re.findall(r"'([^']+)'!", c.value):
                        assert ref in wb.sheetnames, (ws.title, c.coordinate, ref)
    assert n > 1000
    pr = wb["Pricing Schedule"]
    assert pr["F6"].value == '=IF(OR(D6="",E6=""),"",ROUND(D6*E6,2))'
    cc = wb["Commercial Calculation"]
    assert "ROUND(IF(B17=\"markup\"" in cc["B20"].value and cc["B11"].value.startswith("=ROUND(")
    assert wb["Project Summary"]["B6"].value.startswith("='Direct Cost Build-Up'!K")


def test_formatting_input_vs_calculated_cells_and_number_formats(blank):
    wb, _, _ = blank
    ws = wb["Pricing Schedule"]
    assert ws["F6"].fill.fgColor.rgb.endswith("E7E6E6") and ws["F6"].protection.locked          # grey formula cell is locked
    ws2 = wb["BOQ Import"]
    assert ws2["B6"].fill.fgColor.rgb.endswith("FFF2CC") and not ws2["B6"].protection.locked   # yellow input is editable
    assert "#,##0.00" in ws["F6"].number_format and "#,##0.000" in ws["D6"].number_format
    assert ws["A1"].font.bold and ws.column_dimensions["B"].width >= 40


def test_data_dictionary_lists_every_column(blank):
    wb, _, _ = blank
    ws = wb["Data Dictionary"]
    rows = [r for r in ws.iter_rows(min_row=FIRST_ROW, values_only=True) if r[0]]
    assert len(rows) == sum(len(S[n].cols) for n in REQUIRED_SHEETS if n in S)
    assert ("BOQ Import", "Quantity", "quantity", "Input (yellow)") == tuple(next(r for r in rows if r[2] == "quantity" and r[0] == "BOQ Import")[:4])


def test_standalone_templates_generated_and_openable(tmp_path):
    from app.database import Database
    from app.services import refdata
    db = Database(":memory:")
    db.migrate()
    refdata.seed_reference(db)
    out = generate_all(tmp_path, db)
    assert {p.name for p in out} == set(TEMPLATE_FILES)
    for p in out:
        if p.suffix == ".xlsx":
            wb = openpyxl.load_workbook(p)
            assert wb.sheetnames[0] in ("Instructions", "Read Me") and len(wb.sheetnames) >= 2
        else:
            from docx import Document
            assert "{{" in " ".join(par.text for par in Document(str(p)).paragraphs)
    assert "Rate Dataset" in openpyxl.load_workbook(tmp_path / "resource_rate_template.xlsx").sheetnames
    # shipped templates in the repo are the same structure
    shipped = Path(__file__).resolve().parents[2] / "templates"
    assert {p.name for p in shipped.iterdir()} >= set(TEMPLATE_FILES)


@pytest.fixture()
def mini_wb(seeded, tmp_path):
    pid = build_mini(seeded, tmp_path)
    p = build_master(tmp_path / "m.xlsx", seeded, pid)
    return seeded, pid, p, tmp_path


def test_populated_workbook_matches_engine_and_ids(mini_wb):
    db, pid, p, _ = mini_wb
    est = compute_estimate(db, pid)
    wb = openpyxl.load_workbook(p)
    ps = wb["Pricing Schedule"]
    amounts = [ps.cell(r, 7).value for r in range(6, 8)]
    assert sum(Decimal(str(a)) for a in amounts) + Decimal(str(wb["Pricing Schedule"].cell(8, 7).value or 0)) == est.schedule.total_pre_tax    # incl. separate lines (none here)
    dc = wb["Direct Cost Build-Up"]
    assert {dc.cell(r, 5).value for r in range(6, 10)} == {"DEMO-MAT-001", "DEMO-LAB-001"}
    assert wb["Material Library"]["A6"].value == "DEMO-MAT-001" and wb["Material Library"]["G6"].value == 100
    setup = {wb["Project Setup"].cell(r, 1).value: wb["Project Setup"].cell(r, 2).value for r in range(6, 24)}
    assert setup["Project ID"] == pid and setup["Estimate fingerprint"] == est.fingerprint and setup["Demonstration project"] == "Yes"
    boq_ws = wb["BOQ Import"]
    assert [boq_ws.cell(r, 1).value for r in (6, 7)] == ["1.1", "1.2"] and boq_ws.cell(7, 4).value == "m3"


@needs_lo
def test_libreoffice_recalculation_has_no_errors_or_mismatches(mini_wb):
    db, pid, p, _ = mini_wb
    rep = verify_workbook(p)
    assert rep.recalculated and rep.formula_count > 50
    assert rep.errors == [] and rep.mismatches == [], (rep.errors, rep.mismatches)


@needs_lo
def test_edited_formula_or_value_is_detected_as_mismatch(mini_wb):
    db, pid, p, tmp = mini_wb
    wb = openpyxl.load_workbook(p)
    wb["Commercial Calculation"]["B9"] = 50          # someone changes the contingency % input: formulas no longer match the engine
    wb.save(tmp / "tampered.xlsx")
    rep = verify_workbook(tmp / "tampered.xlsx")
    assert rep.mismatches and any("Commercial Calculation" in m or "Project Summary" in m for m in rep.mismatches)


@needs_lo
def test_blank_template_recalculates_cleanly(blank):
    _, p, _ = blank
    rep = verify_workbook(p)
    assert rep.errors == [] and rep.recalculated


def test_structure_validation_detects_tampering(mini_wb):
    db, pid, p, tmp = mini_wb
    assert roundtrip.validate_structure(p, db, pid).ok
    wb = openpyxl.load_workbook(p)
    wb["Pricing Schedule"]["F6"] = 12345                                     # formula overwritten by a constant
    del wb["Risk Register"]
    wb["BOQ Import"].cell(KEY_ROW, 2).value = "renamed_key"
    wb.save(tmp / "bad.xlsx")
    st = roundtrip.validate_structure(tmp / "bad.xlsx")
    assert "Risk Register" in st.missing_sheets and st.header_problems and any("Pricing Schedule!F6" in x for x in st.overwritten_formulas)
    with pytest.raises(ValidationFailed) as e:
        roundtrip.import_master(db, ESTIMATOR, pid, tmp / "bad.xlsx")
    assert e.value.code == "WORKBOOK_INVALID"
    # stale: the data changed after the workbook was exported
    db.execute("UPDATE boq_items SET quantity='99' WHERE project_id=? AND item_key='1.1'", (pid,))
    assert roundtrip.validate_structure(p, db, pid).stale


def test_roundtrip_edits_become_drafts_and_approval_columns_are_ignored(mini_wb):
    db, pid, p, tmp = mini_wb
    wb = openpyxl.load_workbook(p)
    # 1. new BOQ row appended in the yellow input area
    ws = wb["BOQ Import"]
    ws["A8"], ws["B8"], ws["D8"], ws["E8"] = "1.3", "Extra concrete (added in Excel)", "m3", 7
    # 2. an assumption override proposed + a forged approval status
    wa = wb["Assumptions Register"]
    keys = [c.value for c in wa[KEY_ROW]]
    row = FIRST_ROW
    wa.cell(row, keys.index("override_value") + 1).value = "9"
    wa.cell(row, keys.index("override_reason") + 1).value = "excel what-if"
    aid = wa.cell(row, 1).value
    wa.cell(row, keys.index("status") + 1).value = "Approved"
    # 3. register rows
    wc = wb["Clarification Register"]
    wc["B6"], wc["J6"] = "Please confirm concrete grade", "Estimator"
    wk = wb["Risk Register"]
    wk["B6"] = "Weather delay"
    # 4. rate proposal on a library row
    wl = wb["Material Library"]
    kl = [c.value for c in wl[KEY_ROW]]
    wl.cell(6, kl.index("new_rate") + 1).value = 118
    wl.cell(6, kl.index("new_rate_type") + 1).value = "Estimator-Entered Allowance"
    wl.cell(6, kl.index("new_evidence") + 1).value = "allowance per PM"
    wl.cell(6, kl.index("rate_status") + 1).value = "Approved (forged)"
    # 5. indirect edit
    wi = wb["Indirect Costs"]
    ki = [c.value for c in wi[KEY_ROW]]
    wi.cell(6, ki.index("amount") + 1).value = 750
    wb.save(tmp / "edited.xlsx")
    rep = roundtrip.import_master(db, ESTIMATOR, pid, tmp / "edited.xlsx")
    assert rep.errors == [], rep.errors
    assert boq.get_item(db, pid, "1.3")["quantity"] == "7" and boq.get_item(db, pid, "1.3")["description"].startswith("Extra concrete")
    # existing items are unchanged (identical re-import), no duplicates
    assert db.scalar("SELECT COUNT(*) FROM boq_items WHERE project_id=? AND item_type='item'", (pid,)) == 3
    ra = assumptions.resolve(db, pid, aid)
    assert ra.value == "9" and ra.status == "Pending Review"                      # forged 'Approved' ignored
    assert len(registers.list_clarifications(db, pid)) == 1 and len(registers.list_risks(db, pid)) == 1
    new_rate = db.one("SELECT * FROM resource_rates WHERE resource_id='DEMO-MAT-001' AND unit_price='118'")
    assert new_rate["status"] == "Draft" and new_rate["rate_type"] == "Estimator-Entered Allowance"
    cur = db.one("SELECT * FROM indirect_items WHERE project_id=? AND code='IND-1' ORDER BY version DESC", (pid,))
    assert cur["amount"] == "750" and cur["status"] == "Draft"                    # approved v1 is superseded by a new DRAFT version, not edited
    assert db.scalar("SELECT status FROM indirect_items WHERE project_id=? AND code='IND-1' AND version=1", (pid,)) == "Approved"
    assert compute_estimate(db, pid).indirect.total == Decimal("500.00")          # approved v1 stays in force until the revision is approved
    # a changed quantity in the file becomes a pending change request, not an overwrite
    wb = openpyxl.load_workbook(tmp / "edited.xlsx")
    wb["BOQ Import"]["E6"] = 11
    wb.save(tmp / "edited2.xlsx")
    roundtrip.import_master(db, ESTIMATOR, pid, tmp / "edited2.xlsx")
    assert boq.get_item(db, pid, "1.1")["quantity"] == "10" and any(c["new_value"] == "11" for c in boq.pending_changes(db, pid))


def test_round_trip_into_a_fresh_project_reproduces_the_boq(mini_wb):
    db, pid, p, tmp = mini_wb
    new = projects.create_project(db, ESTIMATOR, name="Copy")
    rep = roundtrip.import_master(db, ESTIMATOR, new, p)
    a = {i["item_key"]: (i["description"], i["unit"], i["quantity"]) for i in boq.list_items(db, pid)}
    b = {i["item_key"]: (i["description"], i["unit"], i["quantity"]) for i in boq.list_items(db, new)}
    assert a == b and rep.created.get("BOQ Import") == 2
    it = boq.get_item(db, new, "1.1")
    assert it["assembly_code"] == "DEMO-CONC" and it["mapping_status"] == "Approved"   # mapping travels via the mapping sheet
    assert db.scalar("SELECT COUNT(*) FROM approvals WHERE entity_type='boq_mapping' AND project_id=?", (new,)) == 2   # recorded as the importing estimator's decision
