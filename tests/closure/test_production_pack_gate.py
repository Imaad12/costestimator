"""A pack can only carry a production label through the real approval chain, and even then not while it has BLOCKED documents or placeholders.

Fixture records are ordinary non-synthetic TEST FIXTURES inside the test database; they are not market data."""
from __future__ import annotations

import json
from pathlib import Path

import pytest

from app.core.errors import ApprovalBlocked
from app.services import estimates, readiness, tender
from app.tender import model, packs
from tests.approvals.test_actor_isolation import finish, mk_users, real_project


def _ids(db, pk):
    return {r["doc_code"]: dict(r) for r in db.all("SELECT * FROM generated_documents WHERE pack_pk=?", (pk,))}


def test_human_chain_alone_does_not_make_the_pack_submission_ready(seeded, tmp_path):
    U = mk_users(seeded)
    pid = real_project(seeded, tmp_path, U)
    from app.services import projects
    projects.update_project(seeded, U["h-est"], pid, client="FIXTURE CLIENT (not a real customer)")
    vid, ap = finish(seeded, tmp_path, U, pid)
    assert readiness.evaluate(seeded, pid, vid, ap)["production_ready"]                      # the approval chain itself verifies ...
    r = packs.generate_pack(seeded, U["h-est"], pid, out_root=tmp_path / "out", issue_date="2027-01-01", version_id=vid)
    docs = _ids(seeded, r["pack_pk"])
    assert r["classification"] == "not_ready" and r["production_ready"] is False                      # ... but the pack has BLOCKED documents and placeholders
    assert r["banner"].startswith("NOT SUBMISSION READY - the approval chain verifies but the pack has BLOCKED documents")
    assert all(d["status"] in ("BLOCKED", "Draft", "Not Applicable") for d in docs.values()) and "Approved for Submission" not in {d["status"] for d in docs.values()}
    man = json.loads((Path(r["out_dir"]) / "manifest.json").read_text())
    assert man["production_ready"] is False and man["blocked_documents"]
    d15 = model.parse_back(next(Path(r["out_dir"]).glob("*D15*.docx")))["text"]
    assert "Production ready" in d15 and "pack: documents BLOCKED" in d15                      # the readiness record states the pack's status
    v = packs.validate_pack(seeded, r["pack_pk"])
    assert v["ok"], v["errors"]                                                               # an honest not-ready pack validates; a mislabelled one would not
    # document approval by a registered human is recorded, but it cannot lift a blocked/not-ready pack
    from tests.conftest import COMMERCIAL
    ap_id = packs.approve_document(seeded, U["h-com"], pid, f"{pid}-D41")
    assert seeded.one("SELECT actor_type, environment FROM approvals WHERE id=?", (ap_id,))["actor_type"] == "human"
    r2 = packs.generate_pack(seeded, U["h-est"], pid, out_root=tmp_path / "out", issue_date="2027-01-01", version_id=vid)
    assert r2["production_ready"] is False and _ids(seeded, r2["pack_pk"])["D41"]["status"] == "Draft"


def test_test_actor_cannot_approve_documents_on_a_real_project_in_production(seeded, tmp_path):
    U = mk_users(seeded)
    from app.services import auth
    auth.create_user(seeded, "t-com", "Test Commercial", "Commercial Approver", "pw", actor_type="test")
    T = auth.get_actor(seeded, "t-com")
    pid = real_project(seeded, tmp_path, U)
    vid, ap = finish(seeded, tmp_path, U, pid)
    r = packs.generate_pack(seeded, U["h-est"], pid, out_root=tmp_path / "out", issue_date="2027-01-01", version_id=vid)
    with pytest.raises(ApprovalBlocked) as e:
        packs.approve_document(seeded, T, pid, f"{pid}-D02")
    assert e.value.code == "NON_HUMAN_ACTOR"
    # a copied / hand-inserted document approval row has no audit entry and never counts
    seeded.insert("approvals", dict(project_id=pid, entity_type="tender_document", entity_id=f"{pid}-D02", version="1", reviewer="h-com", reviewer_role="Commercial Approver", action="approve",
                                    ts="2027-01-01T00:00:00", actor_type="human", environment="production"))
    assert packs._doc_approved(seeded, pid, f"{pid}-D02") is False                                # the audit-less row is not an approval
    row = seeded.one("SELECT * FROM approvals WHERE entity_type='tender_document' ORDER BY id DESC LIMIT 1")
    probs, _ = readiness.check_approval(seeded, row, "doc", entity_type="tender_document", entity_id=f"{pid}-D02")
    assert any("no matching audit-log entry" in p for p in probs)


def test_doc_status_mapping_never_approves_without_the_gate():
    assert packs.doc_status("production", True, False, True) == "Approved for Submission"
    assert packs.doc_status("production", True, False, False) == "For Review"
    assert packs.doc_status("production", True, True, True) == "BLOCKED"
    assert packs.doc_status("demonstration", False, False, True) == "Preliminary (DEMONSTRATION)"
    assert packs.doc_status("not_ready", False, False, True) == "Draft"


def test_validator_refuses_unsupported_claims_in_a_non_production_pack(seeded, tmp_path):
    U = mk_users(seeded)
    pid = real_project(seeded, tmp_path, U)
    vid, ap = finish(seeded, tmp_path, U, pid)
    r = packs.generate_pack(seeded, U["h-est"], pid, out_root=tmp_path / "out", issue_date="2027-01-01", version_id=vid)
    d02 = next(Path(r["out_dir"]).glob("*D02*.docx"))
    from docx import Document
    d = Document(str(d02))
    d.add_paragraph("This tender is submission ready and approved for submission and fully compliant.")
    d.save(str(d02))
    v = packs.validate_pack(seeded, r["pack_pk"])
    assert not v["ok"] and any("claims submission readiness" in e for e in v["errors"]) and any("claims approval for submission" in e for e in v["errors"])
    assert any("checksum" in e for e in v["errors"])                                           # the edit is also visible as a checksum break
