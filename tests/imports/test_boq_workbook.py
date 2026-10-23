"""A whole multi-sheet BOQ workbook in one upload: cover/summary skipped, stacked two-line headings, swapped QTY/UNIT headings, lettered sub-items, total lines, real-world unit spellings.
The workbook is built here from SYNTHETIC rows laid out like a real client BOQ; the expected counts are written down by hand, not read back from the importer."""
from openpyxl import Workbook

from app.imports import boq as bi
from app.services import projects

from tests.conftest import ESTIMATOR
from tests.imports.test_boq_import import pid  # noqa: F401


def _header(ws, extra=()):
    ws.append(["PROJECT   : TEST WAREHOUSE (SYNTHETIC)"]); ws.append(["OWNER      :  FICTIONAL OWNER CO"]); ws.append(["LOCATION : NOWHERE, KSA"])
    ws.append(["ITEM", "DESCRIPTION", "BOQ"]); ws.append([None, None, "QTY.", "UNIT", "PRICE", "AMOUNT"])


def make(path):
    wb = Workbook()
    cover = wb.active; cover.title = "COVER"
    cover.append(["BILL OF QUANTITIES"]); cover.append(["TEST WAREHOUSE"])
    summ = wb.create_sheet("SUMMARY")
    summ.append(["PROJECT   : TEST WAREHOUSE (SYNTHETIC)"]); summ.append(["ITEM", "DESCRIPTION", "TOTAL AMOUNT"]); summ.append([None, None, "SAR"]); summ.append(["A", "PLUMBING", 0])
    pl = wb.create_sheet("PLUMBING"); _header(pl)
    pl.append([None, "DIVISION 1A- MECHANICAL"]); pl.append(["1A.1", "SANITARY PIPING"]); pl.append([None, "Narrative text about the glue to be used"])
    pl.append(["a.", "160 mm dia pipe", 246, "LM"]); pl.append(["b.", "110 mm dia pipe", 300, "LM"]); pl.append(["c.", "50 mm dia pipe", "-", "LM"])
    pl.append(["1A.2", "SANITARY COMPONENTS"]); pl.append(["a.", "Floor drain", 30, "pcs"]); pl.append(["b.", "Manhole", 1, "LOT"])
    pl.append([None, "To Bill Summary", None, None, None, 0])
    fa = wb.create_sheet("FIRE ALARM")                                  # headings say UNIT then QTY, data is the other way round
    fa.append(["PROJECT   : TEST WAREHOUSE (SYNTHETIC)"]); fa.append(["ITEM", "DESCRIPTION", "UNIT", "QTY", "UNIT PRICE (SR)", "TOTAL AMOUNT (SR)"])
    fa.append(["1.1", "Control panel", 1, "PSC"]); fa.append(["1.2", "Repeater panel", 2, "PSC"]); fa.append(["2.1", "Detector", 28, "PSC"]); fa.append(["2.2", "Sounder", 5, "No"])
    wb.save(path)
    return path


def test_analyse_says_what_each_sheet_is(tmp_path):
    a = bi.analyse_workbook(make(tmp_path / "w.xlsx"))
    roles = {s["name"]: s for s in a["sheets"]}
    assert roles["COVER"]["role"] == "skip" and roles["SUMMARY"]["role"] == "skip" and "Totals only" in roles["SUMMARY"]["reason"]
    assert roles["PLUMBING"]["role"] == "boq" and roles["PLUMBING"]["items"] == 4            # '-' is not a number, the total line has no quantity
    assert roles["FIRE ALARM"]["role"] == "boq" and roles["FIRE ALARM"]["items"] == 4
    assert any("swapped" in w for w in roles["FIRE ALARM"]["warnings"])
    assert a["importable"] == ["PLUMBING", "FIRE ALARM"]
    assert a["project_info"] == {"name": "TEST WAREHOUSE (SYNTHETIC)", "client": "FICTIONAL OWNER CO", "location": "NOWHERE, KSA"}


def test_import_workbook_keeps_every_quantity_and_keys_are_unique_across_sheets(tmp_path, seeded, pid):
    db = seeded
    res = bi.import_workbook(db, ESTIMATOR, pid, make(tmp_path / "w.xlsx"))
    assert res["items_new"] == 8 + 0 and res["errors"] == []                                   # 4 plumbing + 4 fire alarm lines with a number; hand-counted
    assert res["rejected"] == 1                                                                 # the '-' quantity is refused with a reason, not guessed
    rows = {r["item_key"]: r for r in db.all("SELECT * FROM boq_items WHERE project_id=? AND item_type='item'", (pid,))}
    assert rows["PLUMBING/1A.1.a"]["quantity"] == "246" and rows["PLUMBING/1A.2.a"]["quantity"] == "30"       # lettered sub-items stay distinct under their parent
    assert rows["PLUMBING/1A.1.a"]["unit"] == "m" and rows["PLUMBING/1A.2.b"]["unit"] == "ls" and rows["FIRE ALARM/1.1"]["unit"] == "nr"      # LM, LOT, PSC
    assert rows["FIRE ALARM/2.1"]["quantity"] == "28" and rows["FIRE ALARM/2.1"]["trade"] == "FIRE ALARM"
    assert all(r["unit_status"] == "ok" for r in rows.values())
    assert db.scalar("SELECT COUNT(*) FROM boq_rejected_rows WHERE project_id=?", (pid,)) == 1
    again = bi.import_workbook(db, ESTIMATOR, pid, tmp_path / "w.xlsx")                          # the same file again changes nothing
    assert again["items_new"] == 0 and again["items_unchanged"] == 8 and again["change_requests"] == 0
    assert db.scalar("SELECT COUNT(*) FROM boq_items WHERE project_id=? AND status_in_revision<>'Current'", (pid,)) == 0
