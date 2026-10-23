"""BOQ builder (pick items -> BOQ document / project lines) and the full price-list workbook. SYNTHETIC data."""
from __future__ import annotations

import io
import re

import openpyxl

from app.database import get_db
from tests.integration.test_api_ui import client  # noqa: F401  (shared fixture)
from tests.integration.test_web_tender_pages import _login


def test_price_list_lists_every_active_resource_with_blank_prices_and_uploads_as_is(client):
    db = get_db()
    _login(client, "proc")
    r = client.get("/rates/workbook/price-list.xlsx")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    ws = openpyxl.load_workbook(io.BytesIO(r.content))["Rate Dataset"]
    hdr = {str(c.value): c.column for c in ws[5] if c.value}
    rows = [row for row in ws.iter_rows(min_row=6, values_only=True) if row[0]]
    assert len(rows) == db.scalar("SELECT COUNT(*) FROM resources WHERE is_active=1") > 0
    assert all(row[hdr["Unit price"] - 1] is None and row[hdr["Currency"] - 1] is None for row in rows)       # no price is ever supplied
    assert {"material", "labour", "plant", "equipment", "subcontract"} <= {row[hdr["Type"] - 1] for row in rows}
    ws.cell(row=6, column=hdr["Unit price"], value=11.5); ws.cell(row=6, column=hdr["Currency"], value="SAR")
    buf = io.BytesIO(); ws.parent.save(buf)
    chk = client.post("/rates/workbook/preview", data={"source_name": "Price list test (SYNTHETIC)", "source_type": "Quotation"},
                      files={"file": ("pl.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert chk.status_code == 200 and "Listed rows with no price" in chk.text          # unfilled rows are skipped and said so, not rejected one by one


def test_builder_page_lists_catalogue_and_builds_a_boq_with_only_approved_rates(client):
    db = get_db()
    _login(client, "est")
    page = client.get("/boq-builder")
    assert page.status_code == 200 and "BOQ builder" in page.text and 'name="q_MAT-A-001"' in page.text
    assert client.post("/boq-builder/build", data={"title": "x"}, follow_redirects=False).headers["location"].startswith("/boq-builder?err=")        # nothing picked
    bad = client.post("/boq-builder/build", data={"q_MAT-A-001": "abc"}, follow_redirects=False)
    assert "not+a+number" in bad.headers["location"]
    neg = client.post("/boq-builder/build", data={"q_MAT-A-001": "-2"}, follow_redirects=False)
    assert "greater+than+zero" in neg.headers["location"]
    r = client.post("/boq-builder/build", data={"title": "Package A", "mode": "download", "q_MAT-A-001": "12.5", "q_MAT-A-002": "", "q_NOPE-1": ""})
    assert r.status_code == 200 and "Package_A.xlsx" in r.headers["content-disposition"]
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    cells = [[c for c in row] for row in ws.iter_rows(values_only=True)]
    line = next(c for c in cells if c[7] and str(c[7]).startswith("MAT-A-001"))
    assert line[3] == 12.5 and line[4] is None and line[6] is None                      # quantity as typed; no approved rate -> rate and amount stay blank
    assert any("have no approved rate" in str(c[0]) for c in cells)
    assert len([c for c in cells if c[7]]) == 2                                          # header + the one selected line; the empty quantity was left out


def test_builder_adds_lines_to_a_project_boq(client):
    db = get_db()
    _login(client, "est")
    made = client.post("/ui/do/create_project", data={"name": "Builder test (SYNTHETIC)", "client": "FICTIONAL", "_next": "/projects/{result}/tower"}, follow_redirects=False)
    pid = made.headers["location"].split("/")[2]
    r = client.post("/boq-builder/build", data={"title": "From builder", "mode": "project", "project": pid, "q_MAT-A-001": "7"}, follow_redirects=False)
    assert r.headers["location"].startswith(f"/projects/{pid}/boq?ok="), r.headers["location"]
    row = db.one("SELECT quantity, unit, spec_ref FROM boq_items WHERE project_id=? AND spec_ref='MAT-A-001'", (pid,))
    assert row is not None and row["quantity"] in ("7", "7.0")
    no_proj = client.post("/boq-builder/build", data={"mode": "project", "q_MAT-A-001": "1"}, follow_redirects=False)
    assert "Choose+the+project" in no_proj.headers["location"]


def test_boq_upload_reads_a_workbook_and_creates_the_project_from_its_details(client, tmp_path):
    from tests.imports.test_boq_workbook import make
    f = make(tmp_path / "w.xlsx").read_bytes()
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    _login(client, "est")
    assert "BOQ upload" in client.get("/boq-upload").text
    r = client.post("/boq-upload/analyse", data={"project": ""}, files={"file": ("w.xlsx", f, mime)})
    assert r.status_code == 200 and "Totals only" in r.text and "swapped" in r.text and 'value="TEST WAREHOUSE (SYNTHETIC)"' in r.text
    stored = re.search(r'name="stored" value="([^"]+)"', r.text).group(1)
    nothing = client.post("/boq-upload/import", data={"stored": stored, "project": "", "new_name": "X"}, follow_redirects=False)
    assert "Tick+at+least+one+sheet" in nothing.headers["location"]
    done = client.post("/boq-upload/import", data={"stored": stored, "project": "", "new_name": "Test Warehouse (Synthetic)", "new_client": "Fictional Owner Co", "sheet": ["PLUMBING", "FIRE ALARM"]}, follow_redirects=False)
    loc = done.headers["location"]
    assert re.match(r"/projects/PRJ-\d+/boq\?ok=Imported\+8\+new", loc), loc
    pid = loc.split("/")[2]
    assert get_db().scalar("SELECT COUNT(*) FROM boq_items WHERE project_id=? AND item_type='item'", (pid,)) == 8
    assert client.post("/boq-upload/analyse", data={"project": ""}, files={"file": ("w.pdf", b"%PDF", "application/pdf")}, follow_redirects=False).headers["location"].startswith("/boq-upload?err=")
