"""Approvals inbox, Rates workbook (template, check, import) and the Settings hub (users, audit, integrity, standards, documents). SYNTHETIC data."""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

from app.database import get_db
from app.models.actor import Actor
from app.services import auth, rates
from tests.integration.test_api_ui import client  # noqa: F401  (shared fixture)
from tests.integration.test_web_tender_pages import _login

PROC = Actor("proc", "Procurement Reviewer", "proc")


def _pending_rate(db):
    rid = rates.create_resource(db, PROC, resource_type="material", description="Inbox test grout (SYNTHETIC)", unit="kg", cost_code="C.07")
    r = rates.propose_rate(db, PROC, rid, unit_price="12.50", rate_type="Synthetic Demonstration Rate", price_date="2026-10-01", is_synthetic=True, evidence="SYNTHETIC test")
    rates.submit_rate(db, PROC, r)
    return rid, r


def test_inbox_shows_buttons_only_to_roles_that_may_decide_and_never_on_own_work(client):
    db = get_db()
    rid, r = _pending_rate(db)
    _login(client, "proc")
    own = client.get("/approvals").text
    assert "You raised this" in own and 'action="/ui/do/approve_rate"' not in own          # four-eyes: the author sees no approve button
    _login(client, "view")
    v = client.get("/approvals").text
    assert 'action="/ui/do/approve_rate"' not in v and "cannot decide this" in v
    _login(client, "com")
    page = client.get("/approvals").text
    assert 'action="/ui/do/approve_rate"' in page and f'name="rate_id" value="{r}"' in page and 'action="/ui/do/reject_rate"' in page
    assert client.get("/approvals?group=rates").status_code == 200
    for href in set(re.findall(r'href="(/(?:libraries|assemblies|assumptions|projects|rates)[^"]*)"', page)):          # every link the inbox offers must resolve
        assert client.get(href).status_code == 200, href
    ok = client.post("/ui/do/approve_rate", data={"rate_id": r, "comments": "checked", "_next": "/approvals"}, follow_redirects=False)
    assert "ok=" in ok.headers["location"], ok.headers["location"]
    assert db.scalar("SELECT status FROM resource_rates WHERE id=?", (r,)) == "Approved"
    assert 'action="/ui/do/approve_rate"' not in client.get("/approvals").text
    assert "approve" in client.get("/approvals?log=1").text.lower()


def test_rates_workbook_template_check_and_import(client):
    db = get_db()
    rid = rates.create_resource(db, PROC, resource_type="material", description="Workbook test cement (SYNTHETIC)", unit="kg", cost_code="C.07")
    n_rates = db.scalar("SELECT COUNT(*) FROM resource_rates")
    _login(client, "proc")
    assert "Rates workbook" in client.get("/rates/workbook").text
    t = client.get("/rates/workbook/template.xlsx")
    assert t.status_code == 200 and t.headers["content-type"].startswith("application/vnd.openxmlformats")
    import io
    wb = openpyxl.load_workbook(io.BytesIO(t.content))
    assert "Rate Dataset" in wb.sheetnames and "Resource List" in wb.sheetnames
    assert any(row[0].value == rid for row in wb["Resource List"].iter_rows(min_row=2))
    ws = wb["Rate Dataset"]
    hdr = {str(c.value): c.column for c in ws[5] if c.value}
    for col, v in (("Resource ID", rid), ("Unit", "kg"), ("Unit price", 9.75), ("Currency", "SAR"), ("Price date", "2026-10-01"), ("Evidence / reference", "SYNTHETIC test quote")):
        ws.cell(row=6, column=hdr[col], value=v)
    ws.cell(row=7, column=hdr["Resource ID"], value="NO-SUCH-RES"); ws.cell(row=7, column=hdr["Unit price"], value=5); ws.cell(row=7, column=hdr["Currency"], value="SAR")
    buf = io.BytesIO(); wb.save(buf)
    files = {"file": ("filled.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    no_src = client.post("/rates/workbook/preview", files=files, data={"source_type": "Quotation"}, follow_redirects=False)
    assert "err=" in no_src.headers["location"]                                          # a source is mandatory
    pv = client.post("/rates/workbook/preview", files=files, data={"source_name": "SYNTHETIC supplier list", "source_type": "Quotation"})
    assert pv.status_code == 200 and "Rejected rows" in pv.text and "does not exist" in pv.text
    assert db.scalar("SELECT COUNT(*) FROM rate_candidates") == 0                       # preview stores nothing
    stored = re.search(r'name="stored" value="([^"]+)"', pv.text).group(1)
    imp = client.post("/rates/workbook/import", data={"stored": stored, "source_name": "SYNTHETIC supplier list", "source_type": "Quotation"}, follow_redirects=False)
    assert "ok=" in imp.headers["location"] and "group=candidates" in imp.headers["location"], imp.headers["location"]
    row = db.one("SELECT status, original_price, source_name FROM rate_candidates")
    assert row["status"] == "Pending Review" and str(row["original_price"]) in ("9.75", "9.750000") and row["source_name"] == "SYNTHETIC supplier list"
    assert db.scalar("SELECT COUNT(*) FROM resource_rates") == n_rates                        # an import never creates a rate
    again = client.post("/rates/workbook/import", data={"stored": stored, "source_name": "SYNTHETIC supplier list", "source_type": "Quotation"}, follow_redirects=False)
    assert "err=" in again.headers["location"]                                          # the identical file is refused
    _login(client, "com")
    assert "review_rate_candidate" in client.get("/approvals?group=candidates").text


def test_settings_hub_tabs_render_and_user_admin_rules(client):
    _login(client, "admin")
    for tab in ("general", "users", "audit", "integrity", "standards", "documents", "sources", "reference"):
        r = client.get(f"/settings?tab={tab}")
        assert r.status_code == 200 and "Settings" in r.text, tab
    import html
    assert "Don't approve your own work" in html.unescape(client.get("/settings?tab=documents").text)
    assert "Recompute and compare" in client.get("/settings?tab=integrity").text and client.get("/settings?tab=integrity&verify=1").status_code == 200
    assert "Verify hash chain" in client.get("/settings?tab=audit").text and "Chain verified" in client.get("/settings?tab=audit&verify=1").text
    mk = client.post("/ui/do/create_user", data={"username": "newbie", "display_name": "New B", "role": "Estimator", "password": "longenough1", "_next": "/settings?tab=users"}, follow_redirects=False)
    assert "ok=" in mk.headers["location"]
    assert "newbie" in client.get("/settings?tab=users").text
    short = client.post("/ui/do/reset_user_password", data={"username": "newbie", "password": "short", "_next": "/settings?tab=users"}, follow_redirects=False)
    assert "err=" in short.headers["location"]
    ok = client.post("/ui/do/reset_user_password", data={"username": "newbie", "password": "another-long-pw", "_next": "/settings?tab=users"}, follow_redirects=False)
    assert "ok=" in ok.headers["location"]
    client.cookies.clear()
    assert client.post("/login", data={"username": "newbie", "password": "another-long-pw"}, follow_redirects=False).status_code == 303
    assert client.post("/login", data={"username": "newbie", "password": "longenough1"}, follow_redirects=False).status_code != 303      # old password no longer works
    _login(client, "admin")
    me = client.post("/ui/do/set_user_active", data={"username": "admin", "active": "0", "_next": "/settings?tab=users"}, follow_redirects=False)
    assert "err=" in me.headers["location"]                                                                                              # cannot lock yourself out
    demote = client.post("/ui/do/set_user_role", data={"username": "admin", "role": "Viewer", "_next": "/settings?tab=users"}, follow_redirects=False)
    assert "LAST_ADMIN" in demote.headers["location"]
    off = client.post("/ui/do/set_user_active", data={"username": "newbie", "active": "0", "_next": "/settings?tab=users"}, follow_redirects=False)
    assert "ok=" in off.headers["location"]
    client.cookies.clear()
    assert client.post("/login", data={"username": "newbie", "password": "another-long-pw"}, follow_redirects=False).status_code != 303   # deactivated
    assert get_db().one("SELECT 1 FROM audit_log WHERE entity_type='user' AND entity_id='newbie' AND action='password_reset'")
    _login(client, "est")
    deny = client.post("/ui/do/set_user_role", data={"username": "view", "role": "Administrator", "_next": "/settings?tab=users"}, follow_redirects=False)
    assert "err=" in deny.headers["location"]
    assert "/ui/do/reset_user_password" not in client.get("/settings?tab=users").text                                    # no management controls for non-admins


def test_integrity_tab_detects_a_tampered_evidence_file(client, tmp_path, monkeypatch):
    import app.web.routes_ops as ops
    ev = tmp_path / "ev"; ev.mkdir()
    (ev / "a.pdf").write_bytes(b"original")
    import hashlib, json
    (ev / "evidence_manifest.json").write_text(json.dumps({"files": [{"file": "a.pdf", "sha256": hashlib.sha256(b"original").hexdigest(), "register_row": "STD-X"}]}))
    monkeypatch.setattr(ops, "ROOT", tmp_path / "root")
    (tmp_path / "root" / "data" / "standards").mkdir(parents=True); (tmp_path / "root" / "data" / "standards" / "evidence").symlink_to(ev)
    _login(client, "admin")
    assert "MATCH" in client.get("/settings?tab=integrity&verify=1").text
    (ev / "a.pdf").write_bytes(b"tampered")
    assert "MISMATCH" in client.get("/settings?tab=integrity&verify=1").text


def test_placeholders_are_capitalised_in_templates():
    bad = []
    for f in Path("app/web/templates").glob("*.html"):
        for m in re.finditer(r'placeholder="([a-z][^"]{0,30})', f.read_text(encoding="utf-8")):
            bad.append((f.name, m.group(1)))
    assert not bad, bad[:5]


def test_rates_workbook_column_mapping_for_a_file_that_does_not_follow_the_template(client):
    import io
    import json as _json
    db = get_db()
    rid = rates.create_resource(db, PROC, resource_type="material", description="Mapping test sand (SYNTHETIC)", unit="m3", cost_code="C.07")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Our prices"
    ws.append(["Supplier price list - SYNTHETIC"]); ws.append([])
    ws.append(["Mat Code", "Item name", "Measure", "Cost per unit", "Quote ref"])
    ws.append([rid, "Sand", "m3", 31.5, "Q-1"]); ws.append(["NOPE", "Gravel", "m3", 40, "Q-2"]); ws.append([rid, "Sand again", "m3", 0, "Q-3"])
    buf = io.BytesIO(); wb.save(buf)
    _login(client, "proc")
    files = {"file": ("theirs.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    src = {"source_name": "SYNTHETIC supplier", "source_type": "Quotation"}
    pv = client.post("/rates/workbook/preview", files=files, data=src)
    assert pv.status_code == 200 and "Check the suggested column mapping" in pv.text and "not recognised exactly" in pv.text and "Mat Code" in pv.text
    assert "Import " not in re.sub(r"<[^>]+>", " ", pv.text).split("Start again")[0].split("Re-check with this mapping")[-1]     # nothing to import yet
    stored = re.search(r'name="stored" value="([^"]+)"', pv.text).group(1)
    form = {**src, "stored": stored, "sheet": "Our prices", "header_row": "3", "map_resource_id": "0", "map_description": "1", "map_unit": "2", "map_original_price": "3", "map_source_url": "4"}
    nocur = client.post("/rates/workbook/remap", data=form)
    assert "currency is required" in nocur.text                                                    # never assumed
    ok = client.post("/rates/workbook/remap", data={**form, "default_currency": "sar"})
    assert ok.status_code == 200 and "Import 1 row(s) for review" in ok.text and "does not exist" in ok.text and "price must be greater than zero" in ok.text
    partial = client.post("/rates/workbook/remap", data={**src, "stored": stored, "sheet": "Our prices", "header_row": "3", "map_unit": "2"})
    assert "Map at least the unit price" in partial.text
    n = db.scalar("SELECT COUNT(*) FROM rate_candidates")
    imp = client.post("/rates/workbook/import", data={**form, "default_currency": "sar"}, follow_redirects=False)
    assert "ok=" in imp.headers["location"], imp.headers["location"]
    c = db.one("SELECT original_price, original_currency, source_url, status FROM rate_candidates ORDER BY id DESC LIMIT 1")
    assert db.scalar("SELECT COUNT(*) FROM rate_candidates") == n + 1 and c["original_currency"] == "SAR" and c["source_url"] == "Q-1" and c["status"] == "Pending Review"
    man = _json.loads(db.scalar("SELECT parsing_json FROM import_manifests WHERE kind='rates' ORDER BY id DESC LIMIT 1"))
    assert man["mapping"]["original_price"] == 3 and man["default_currency"] == "sar" and man["header_row"] == 3            # the mapping is part of the provenance


def test_superuser_is_a_generated_password_test_principal_and_self_approval_stays_opt_in(client, capsys, monkeypatch):
    import argparse
    from app import cli
    from app.config import get_settings
    monkeypatch.delenv("KSA_SUPERUSER_PASSWORD", raising=False)
    assert cli.cmd_superuser(argparse.Namespace(password=None, reset=False)) == 0
    out = capsys.readouterr().out
    pw = re.search(r"superuser\s+/\s+(\S+)", out).group(1)
    assert pw != "superuser-dev-pw" and len(pw) >= 10                                              # generated, shown once - never a published default
    u = get_db().one("SELECT role, actor_type FROM users WHERE username='superuser'")
    assert (u["role"], u["actor_type"]) == ("Administrator", "test")                               # a test principal can never make an export submission-ready
    assert cli.cmd_superuser(argparse.Namespace(password=None, reset=False)) == 0 and "already exists" in capsys.readouterr().out
    assert client.post("/login", data={"username": "superuser", "password": pw}, follow_redirects=False).status_code == 303
    db = get_db()
    rid, r = _pending_rate(db)
    # self-approval off (default): the Administrator cannot approve a rate that same login raised
    adm = Actor("superuser", "Administrator", "Super User")
    rid2 = rates.create_resource(db, adm, resource_type="material", description="Own work (SYNTHETIC)", unit="kg", cost_code="C.07")
    mine = rates.propose_rate(db, adm, rid2, unit_price="5", rate_type="Synthetic Demonstration Rate", price_date="2026-10-01", is_synthetic=True, evidence="SYNTHETIC")
    rates.submit_rate(db, adm, mine)
    page = client.get("/approvals?group=rates").text
    assert "You raised this" in page and "Self-approval is switched on" not in page
    monkeypatch.setattr(get_settings(), "allow_self_approval", True)                              # only the person running the server can switch it on
    page = client.get("/approvals?group=rates").text
    assert "Self-approval is switched on" in page and f'name="rate_id" value="{mine}"' in page and "You raised this" not in page


def _messy_file(rid, tax=None):
    import io
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Prices"
    ws.append(["Mat Code", "Item name", "Measure", "Cost per unit", "Quote ref"] + (["VAT"] if tax else []))
    ws.append([rid, "Sand", "m3", 31.5, "Q-1"] + ([tax] if tax else []))
    buf = io.BytesIO(); wb.save(buf)
    return {"file": ("messy.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


def test_rate_candidate_needs_vat_confirmed_inbox_explains_it_and_the_whole_flow_reaches_an_approved_rate(client):
    db = get_db()
    rid = rates.create_resource(db, PROC, resource_type="material", description="VAT flow sand (SYNTHETIC)", unit="m3", cost_code="C.07")
    _login(client, "proc")
    src = {"source_name": "SYNTHETIC supplier", "source_type": "Quotation", "default_currency": "SAR"}
    pv = client.post("/rates/workbook/preview", files=_messy_file(rid), data=src)
    assert "Check the suggested column mapping" in pv.text and 'value="original_price" selected' in pv.text and 'value="unit" selected' in pv.text
    stored = re.search(r'name="stored" value="([^"]+)"', pv.text).group(1)
    form = {k: v for k, v in src.items()} | {"stored": stored, "sheet": "Prices", "header_row": "1", "map_resource_id": "0", "map_description": "1", "map_unit": "2", "map_original_price": "3", "map_source_url": "4"}
    chk = client.post("/rates/workbook/remap", data=form)
    assert "do not say whether the price includes VAT" in chk.text                                  # warned before importing
    arr = client.post("/rates/workbook/arranged", data=form)
    assert arr.status_code == 200 and "rates_arranged" in arr.headers["content-disposition"]
    import io
    ws = openpyxl.load_workbook(io.BytesIO(arr.content))["Rate Dataset"]
    assert [c.value for c in ws[6]][:5] == [rid, "Sand", "m3", 31.5, "SAR"]                          # the person's data, in the standard layout
    imp = client.post("/rates/workbook/import", data=form, follow_redirects=False)
    assert "ok=" in imp.headers["location"]
    ref = db.scalar("SELECT candidate_ref FROM rate_candidates ORDER BY id DESC LIMIT 1")
    _login(client, "com")
    page = client.get("/approvals?group=candidates").text
    assert "VAT treatment must be confirmed" in page and 'name="tax_basis"' in page and 'value="approve"' not in page.split(ref)[1].split("qitem")[0]
    blocked = client.post("/ui/do/review_rate_candidate", data={"candidate": ref, "decision": "approve", "_next": "/approvals"}, follow_redirects=False)
    assert "NOT_PUBLISHABLE" in blocked.headers["location"]                                          # the server refuses regardless of the form
    fixed = client.post("/ui/do/edit_rate_candidate", data={"candidate": ref, "tax_basis": "Excluded", "_next": "/approvals"}, follow_redirects=False)
    assert "ok=" in fixed.headers["location"], fixed.headers["location"]
    assert 'name="decision" value="approve"' in client.get("/approvals?group=candidates").text
    for act, extra in (("review_rate_candidate", {"decision": "approve"}), ("publish_rate_candidate", {})):
        r = client.post(f"/ui/do/{act}", data={"candidate": ref, **extra, "_next": "/approvals"}, follow_redirects=False)
        assert "ok=" in r.headers["location"], (act, r.headers["location"])
    rate = db.one("SELECT id, status, unit_price FROM resource_rates WHERE resource_id=? ORDER BY id DESC LIMIT 1", (rid,))
    assert rate["status"] == "Pending Review"                                                       # published, still needs its own approval
    _login(client, "admin")
    assert "ok=" in client.post("/ui/do/approve_rate", data={"rate_id": rate["id"], "_next": "/approvals"}, follow_redirects=False).headers["location"]
    assert db.scalar("SELECT status FROM resource_rates WHERE id=?", (rate["id"],)) == "Approved"


def test_vat_stated_for_the_whole_file_removes_the_per_row_blocker(client):
    db = get_db()
    rid = rates.create_resource(db, PROC, resource_type="material", description="VAT default sand (SYNTHETIC)", unit="m3", cost_code="C.07")
    _login(client, "proc")
    src = {"source_name": "SYNTHETIC supplier", "source_type": "Quotation", "default_currency": "SAR", "default_tax_basis": "Excluded"}
    pv = client.post("/rates/workbook/preview", files=_messy_file(rid), data=src)
    stored = re.search(r'name="stored" value="([^"]+)"', pv.text).group(1)
    form = src | {"stored": stored, "sheet": "Prices", "header_row": "1", "map_resource_id": "0", "map_unit": "2", "map_original_price": "3", "map_source_url": "4", "map_description": "1"}
    assert "do not say whether" not in client.post("/rates/workbook/remap", data=form).text
    client.post("/rates/workbook/import", data=form, follow_redirects=False)
    assert db.scalar("SELECT tax_basis FROM rate_candidates ORDER BY id DESC LIMIT 1") == "Excluded"


def test_errors_are_styled_escaped_and_never_a_bare_500(client, monkeypatch):
    import app.web.routes as web
    _login(client, "est")
    r = client.get("/no-such-page")
    assert r.status_code == 404 and "Page not found" in r.text and "text/html" in r.headers["content-type"]
    assert client.get("/api/no-such-endpoint").headers["content-type"].startswith("application/json")          # API clients still get JSON
    pid = client.post("/ui/do/create_project", data={"name": "Err test (SYNTHETIC)", "_next": "/projects/{result}"}, follow_redirects=False).headers["location"].split("?")[0].split("/")[2]
    r = client.post("/ui/do/update_project", data={"project_id": pid, "_next": "/"}, follow_redirects=False)
    assert "NO_CHANGES" in r.headers["location"]                                                                # was a database syntax error (HTTP 500)

    def boom(*a, **k):
        raise RuntimeError("secret internal detail <script>")
    monkeypatch.setattr(web, "run", boom)
    r = client.post("/ui/do/update_project", data={"project_id": pid, "client": "x", "_next": "/"}, follow_redirects=False)
    assert r.status_code == 303 and "UNEXPECTED_ERROR" in r.headers["location"] and "secret" not in r.headers["location"]       # logged, not leaked
    client.get("/")                                                                                              # still alive


def test_logo_favicon_and_fonts_are_served_without_external_assets(client):
    assert client.get("/favicon.ico").headers["content-type"].startswith("image/svg+xml")
    assert client.get("/static/logo.svg").status_code == 200
    login = client.get("/login").text
    assert "/static/logo.svg" in login and 'rel="icon"' in login
    assert "fonts.googleapis" not in login and "https://" not in re.sub(r"<a [^>]*>", "", login)               # CSP: nothing external is loaded


def test_a_database_rule_violation_is_a_friendly_message_not_a_server_error(client):
    _login(client, "est")
    pid = client.post("/ui/do/create_project", data={"name": "Rule test (SYNTHETIC)", "_next": "/projects/{result}"}, follow_redirects=False).headers["location"].split("?")[0].split("/")[2]
    r = client.post("/ui/do/create_norm", data={"id": "N-QA", "trade": "Civil", "activity": "A", "output_unit": "m3", "basis": "crew_output_per_crew_day", "value": "5", "crew": "{}",
                                                  "evidence_status": "QA test", "_next": f"/projects/{pid}/workforce"}, follow_redirects=False)
    assert r.status_code == 303 and "NORM_EVIDENCE" in r.headers["location"]
