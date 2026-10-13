"""Public-source pipeline: providers (mocked - no internet), manual imports, manifests, candidates, normalisation, publication.

All prices below are SYNTHETIC test values, not market data.  The CCI fixture is the real KAPSARC/GASTAT response saved in
data/public_data/gastat_cci/ (it is an INDEX)."""
from __future__ import annotations

import csv
import json
import urllib.error
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

import pytest

from app.config import BASE_DIR
from app.core.errors import Conflict, PermissionDenied, ValidationFailed
from app.estimators.rate_resolution import ELIGIBLE, UNAPPROVED, resolve_rate
from app.rate_updates import candidates as cand
from app.rate_updates import importer, index_import, normalise, refresh
from app.services import commercial, estimates, rates, snapshots
from app.services.rates import create_resource
from app.sources import providers
from app.sources.base import AccessPolicy, FetchResult, HttpFetcher
from tests.conftest import AI, COMMERCIAL, ESTIMATOR, PROCUREMENT, VIEWER
from tests.helpers import approve_vat, approved_rate, build_mini

CCI = BASE_DIR / "data" / "public_data" / "gastat_cci" / "kapsarc_cci_2026-04.records.json"
TODAY = date.today().isoformat()


class FakeFetcher:
    """Never touches the network. `routes` maps url-substring -> FetchResult-factory."""

    def __init__(self, routes=None, default=None):
        self.calls: list[str] = []
        self.routes = routes or {}
        self.default = default

    def fetch(self, url, policy):
        self.calls.append(url)
        for k, v in self.routes.items():
            if k in url:
                return v(url) if callable(v) else v
        if self.default:
            return self.default(url) if callable(self.default) else self.default
        raise AssertionError(f"unexpected network request in test: {url}")


def ok(url, body: str | bytes, ct="application/json"):
    return FetchResult(url, "OK", 200, body.encode() if isinstance(body, str) else body, ct)


@pytest.fixture()
def res(seeded):
    create_resource(seeded, ESTIMATOR, resource_type="material", description="Ready-mix concrete grade C30 supply", unit="m3",
                    cost_code="C.01", id="T-MAT-001")
    create_resource(seeded, ESTIMATOR, resource_type="material", description="Steel reinforcement bar B500", unit="kg", cost_code="C.01", id="T-MAT-002")
    return seeded


def write_csv(path: Path, rows, header=("Resource ID", "Description", "Unit", "Unit Price", "Currency", "Price Date", "VAT", "Delivery", "Source URL")):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    return path


# ---------------------------------------------------------------------------------------------- index (GASTAT CCI) import
def test_real_cci_fixture_imports_as_index_observations_never_as_rates(seeded):
    dry = index_import.import_index_file(seeded, ESTIMATOR, CCI, dataset="gastat-cci", source_id="SRC-002", source_name="KAPSARC/GASTAT", dry_run=True)
    assert dry["dry_run"] and dry["records"] == 39 and dry["rejected"] == 0
    assert seeded.scalar("SELECT COUNT(*) FROM index_observations") == 0                       # dry run stores no observations
    r = index_import.import_index_file(seeded, ESTIMATOR, CCI, dataset="gastat-cci", source_id="SRC-002", source_name="KAPSARC/GASTAT",
                                       source_url="https://datasource.kapsarc.org/...", description="test")
    assert r["new"] == 39 and r["unchanged"] == 0
    row = seeded.one("SELECT * FROM index_observations WHERE series='Labor Costs' AND measure='Index' AND period='2026-04'")
    assert row["value"] == "104.8951126515095" and row["unit"] == "index points"                  # exact published text
    assert seeded.scalar("SELECT COUNT(*) FROM rate_candidates") == 0 and seeded.scalar("SELECT COUNT(*) FROM resource_rates WHERE unit_price IS NOT NULL") == 0
    m = seeded.one("SELECT * FROM import_manifests WHERE id=?", (r["manifest_id"],))
    assert m["sha256"] and m["accepted_rows"] == 39 and m["status"] == "Imported" and Path(m["manifest_path"]).exists()
    with pytest.raises(Conflict) as e:                                                           # never silently replace a dataset
        index_import.import_index_file(seeded, ESTIMATOR, CCI, dataset="gastat-cci", source_id="SRC-002", source_name="x")
    assert e.value.code == "ALREADY_IMPORTED"


def test_index_revision_keeps_both_values_and_bad_rows_are_rejected_with_reasons(seeded, tmp_path):
    index_import.import_index_file(seeded, ESTIMATOR, CCI, dataset="gastat-cci", source_id="SRC-002", source_name="x")
    revised = {"results": [{"time_period": "2026-04", "measure": "Index", "section_and_group": "Labor Costs", "index": 105.5},
                           {"time_period": "not-a-date", "measure": "Index", "section_and_group": "Labor Costs", "index": 1},
                           {"time_period": "2026-05", "measure": "Index", "section_and_group": "Labor Costs", "index": "abc"},
                           {"time_period": "2026-05", "measure": "Weird", "section_and_group": "Labor Costs", "index": 1}]}
    f = tmp_path / "rev.json"
    f.write_text(json.dumps(revised))
    r = index_import.import_index_file(seeded, ESTIMATOR, f, dataset="gastat-cci", source_id="SRC-002", source_name="x")
    assert r["revised"] == 1 and r["rejected"] == 3
    vals = [x["value"] for x in seeded.all("SELECT value FROM index_observations WHERE series='Labor Costs' AND measure='Index' AND period='2026-04' ORDER BY id")]
    assert vals == ["104.8951126515095", "105.5"]                                                 # history retained
    assert index_import.latest_observations(seeded, "gastat-cci")[0] is not None
    api_err = tmp_path / "err.json"
    api_err.write_text(json.dumps({"error_code": "IncompatibleTypesInComparisonFilter", "message": "x"}))
    with pytest.raises(ValidationFailed) as e:
        index_import.import_index_file(seeded, ESTIMATOR, api_err, dataset="d", source_id=None, source_name="x")
    assert e.value.code == "API_ERROR_RESPONSE"


def test_index_csv_long_format(seeded, tmp_path):
    f = tmp_path / "cci.csv"
    f.write_text("Period,Series,Measure,Value\n2026-03,General Index,Index,103.2\n2026-04,General Index,Index,103.7\n")
    r = index_import.import_index_file(seeded, ESTIMATOR, f, dataset="x", source_id=None, source_name="manual")
    assert r["new"] == 2


# ---------------------------------------------------------------------------------------------- providers
def _kapsarc_page(offset, total=3, n=2):
    rows = [{"time_period": "2026-04", "measure": "Index", "section_and_group": f"S{offset + i}", "index": 100 + offset + i} for i in range(n)]
    return json.dumps({"total_count": total, "results": rows[: max(0, min(n, total - offset))]})


def test_kapsarc_provider_pages_and_imports_via_mock_only(seeded):
    p = providers.KapsarcGastatCci()
    p.page_size = 2
    f = FakeFetcher(default=lambda url: ok(url, _kapsarc_page(int(url.split("offset=")[1].split("&")[0]))))
    rep = p.refresh(seeded, ESTIMATOR, f, dry_run=False)
    assert rep["outcome"] == "OK" and rep["new"] == 3 and len(f.calls) == 2
    assert all(c.startswith("https://datasource.kapsarc.org/api/explore/v2.1/") for c in f.calls)
    again = p.refresh(seeded, ESTIMATOR, FakeFetcher(default=lambda url: ok(url, _kapsarc_page(int(url.split("offset=")[1].split("&")[0])))), dry_run=False)
    assert again["outcome"] == "NOT_MODIFIED"                                                    # identical payload -> not re-imported
    assert seeded.scalar("SELECT COUNT(*) FROM index_observations") == 3


def test_provider_dry_run_makes_no_request_and_blocked_source_is_reported_honestly(seeded):
    p = providers.KapsarcGastatCci()
    never = FakeFetcher()                                                                          # raises if used
    rep = p.refresh(seeded, ESTIMATOR, never, dry_run=True)
    assert rep["outcome"] == "DRY_RUN" and not never.calls and rep["planned_requests"]
    blocked = FakeFetcher(default=FetchResult("u", "BLOCKED", None, error="URLError: Tunnel connection failed: 403 Forbidden"))
    rep = p.refresh(seeded, ESTIMATOR, blocked, dry_run=False)
    assert rep["outcome"] == "BLOCKED" and "Tunnel" in rep["error"]
    assert seeded.scalar("SELECT COUNT(*) FROM index_observations") == 0
    last = seeded.one("SELECT * FROM source_access_log WHERE outcome='BLOCKED' ORDER BY id DESC LIMIT 1")
    assert last["provider"] == "kapsarc-gastat-cci" and "Tunnel" in last["error"]
    junk = FakeFetcher(default=lambda url: ok(url, "<html>not json</html>", "text/html"))
    assert p.refresh(seeded, ESTIMATOR, junk, dry_run=False)["outcome"] == "ERROR"
    api_err = FakeFetcher(default=lambda url: ok(url, json.dumps({"error_code": "X", "message": "bad"})))
    assert p.refresh(seeded, ESTIMATOR, api_err, dry_run=False)["outcome"] == "ERROR"


def test_gastat_site_and_etimad_are_manual_only_and_never_pretend_to_fetch(seeded):
    for pid in ("gastat-site", "etimad-awards"):
        f = FakeFetcher()
        rep = providers.PROVIDERS[pid].refresh(seeded, ESTIMATOR, f, dry_run=False)
        assert rep["outcome"] == "SKIPPED" and "Manual-only" in rep["error"] and not f.calls
        assert providers.PROVIDERS[pid].info.status == "manual-only" and not providers.PROVIDERS[pid].policy.automated_access
    assert seeded.scalar("SELECT COUNT(*) FROM index_observations") == 0 and seeded.scalar("SELECT COUNT(*) FROM rate_candidates") == 0


def test_refresh_sources_isolates_providers_and_reports(seeded, monkeypatch):
    class Boom:
        info = providers.KapsarcGastatCci.info
        policy = AccessPolicy()

        def refresh(self, *a, **k):
            raise RuntimeError("provider exploded")

    monkeypatch.setitem(providers.PROVIDERS, "kapsarc-gastat-cci", Boom())
    rep = refresh.refresh_sources(seeded, ESTIMATOR, dry_run=True, fetcher=FakeFetcher())
    by = {r["provider"]: r for r in rep["providers"]}
    assert by["kapsarc-gastat-cci"]["outcome"] == "ERROR" and "exploded" in by["kapsarc-gastat-cci"]["error"]
    assert by["web-price-watchlist"]["outcome"] == "SKIPPED"                                      # empty watchlist: nothing invented
    assert rep["summary"]["errors"] == 1 and rep["summary"]["skipped_or_dry_run"] >= 3 and rep["summary"]["new_candidates"] == 0
    st = refresh.status_report(seeded)
    assert {s["id"] for s in st["sources"]} == set(providers.PROVIDERS) and "queue" in st


def test_refresh_dry_run_on_real_registry_does_not_touch_data_or_network(seeded):
    rep = refresh.refresh_sources(seeded, ESTIMATOR, dry_run=True, fetcher=FakeFetcher())        # FakeFetcher would raise on any request
    assert rep["dry_run"] and rep["summary"]["new_candidates"] == 0 and rep["summary"]["new_or_revised_index_rows"] == 0
    assert rep["summary"]["errors"] == 0


# ---------------------------------------------------------------------------------------------- fetcher policy
class _Resp:
    def __init__(self, body=b"{}", status=200):
        self._b, self.status, self.headers = body, status, {"Content-Type": "application/json"}

    def read(self, n=-1):
        return self._b

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


class _Opener:
    def __init__(self, script):
        self.script, self.calls = list(script), []

    def open(self, req, timeout=None):
        self.calls.append(req.full_url)
        item = self.script.pop(0)
        if isinstance(item, Exception):
            raise item
        return item


def _http(code):
    return urllib.error.HTTPError("http://x", code, "msg", {}, None)


def test_fetcher_backoff_retries_only_transient_errors_and_caches(tmp_path):
    sleeps = []
    op = _Opener([_http(429), _http(503), _Resp(b'{"a":1}')])
    f = HttpFetcher(cache_dir=tmp_path / "c", sleep=sleeps.append, opener=op)
    pol = AccessPolicy(respect_robots=False, min_interval_s=0, max_retries=2, backoff_base_s=2)
    r = f.fetch("https://h.example/a", pol)
    assert r.ok and r.body == b'{"a":1}' and len(op.calls) == 3 and sleeps == [2.0, 4.0]          # exponential backoff 2s, 4s
    again = f.fetch("https://h.example/a", pol)                                                    # served from cache: no new request
    assert again.from_cache and len(op.calls) == 3
    op2 = _Opener([_http(403), _Resp()])
    r2 = HttpFetcher(cache_dir=tmp_path / "c2", sleep=lambda s: None, opener=op2).fetch("https://h.example/b", pol)
    assert r2.outcome == "BLOCKED" and r2.status == 403 and len(op2.calls) == 1                    # access refusal is never retried / circumvented


def test_fetcher_respects_manual_only_allowlist_and_robots(tmp_path):
    op = _Opener([])
    f = HttpFetcher(cache_dir=tmp_path, sleep=lambda s: None, opener=op)
    assert f.fetch("https://a.example/x", AccessPolicy(automated_access=False)).outcome == "SKIPPED"
    assert f.fetch("https://a.example/x", AccessPolicy(allowed_hosts=("b.example",))).outcome == "SKIPPED"
    assert not op.calls
    robots = _Resp(b"User-agent: *\nDisallow: /private\n")
    op = _Opener([robots, _Resp(b"ok")])
    f = HttpFetcher(cache_dir=tmp_path / "r", sleep=lambda s: None, opener=op)
    pol = AccessPolicy(min_interval_s=0)
    assert f.fetch("https://a.example/private/p", pol).outcome == "BLOCKED"
    assert f.fetch("https://a.example/public/p", pol).ok


def test_fetcher_marks_proxy_tunnel_failures_as_blocked(tmp_path):
    err = urllib.error.URLError("Tunnel connection failed: 403 Forbidden")
    r = HttpFetcher(cache_dir=tmp_path, sleep=lambda s: None, opener=_Opener([err])).fetch(
        "https://stats.example/x", AccessPolicy(respect_robots=False, min_interval_s=0, max_retries=0))
    assert r.outcome == "BLOCKED" and "Tunnel" in r.error


# ---------------------------------------------------------------------------------------------- manual rate import -> candidates
def test_rates_import_dry_run_then_real_creates_candidates_and_manifest(res, tmp_path):
    f = write_csv(tmp_path / "list.csv", [
        ("T-MAT-001", "Ready-mix C30", "m3", "310.50", "SAR", "2026-09-15", "Excluded", "Delivered", "https://example.invalid/price-list"),
        ("T-MAT-002", "Rebar B500", "ton", "2650", "SAR", "2026-09-15", "Excluded", "Supply only", "https://example.invalid/price-list"),
        ("NOPE-1", "Unknown thing", "m3", "5", "SAR", "2026-09-15", "Excluded", "", "https://example.invalid/x"),
        ("T-MAT-001", "Bad price", "m3", "abc", "SAR", "2026-09-15", "Excluded", "", "https://example.invalid/x"),
        ("T-MAT-001", "No currency", "m3", "5", "", "2026-09-15", "Excluded", "", "https://example.invalid/x")])
    dry = importer.import_rates_file(res, ESTIMATOR, f, source_name="Test list", dry_run=True)
    assert dry["dry_run"] and dry["accepted"] == 2 and dry["rejected"] == 3 and res.scalar("SELECT COUNT(*) FROM rate_candidates") == 0
    reasons = " ".join(" ".join(e["errors"]) for e in dry["errors"])
    assert "does not exist" in reasons and "not numeric" in reasons and "currency is required" in reasons
    real = importer.import_rates_file(res, ESTIMATOR, f, source_name="Test list", description="unit test", source_url="https://example.invalid/price-list")
    assert real["accepted"] == 2 and len(real["candidate_ids"]) == 2
    m = res.one("SELECT * FROM import_manifests WHERE id=?", (real["manifest_id"],))
    assert m["accepted_rows"] == 2 and m["rejected_rows"] == 3 and len(m["sha256"]) == 64 and json.loads(m["candidate_ids_json"]) == real["candidate_ids"]
    assert Path(m["stored_path"]).read_bytes() == f.read_bytes() and Path(m["manifest_path"]).exists()
    with pytest.raises(Conflict):
        importer.import_rates_file(res, ESTIMATOR, f, source_name="Test list")                    # same file twice -> refused
    c1 = cand.get(res, real["candidate_ids"][0])
    assert c1["original_price"] == "310.50" and c1["queue"] == "New" and c1["status"] == "Pending Review" and c1["manifest_id"] == real["manifest_id"]
    assert c1["retrieved_at"] and c1["source_url"] == "https://example.invalid/price-list" and c1["source_type"] == "Manual dataset"
    # nothing reached the rate library
    assert res.scalar("SELECT COUNT(*) FROM resource_rates WHERE unit_price IS NOT NULL") == 0


def test_candidate_requires_provenance_and_rejects_index_style_data(res):
    base = dict(source_type="Manual dataset", source_url="https://example.invalid/p", retrieved_at=TODAY, original_price="10", currency="SAR", unit="m3",
                resource_id="T-MAT-001", description="x", tax_basis="Excluded")
    cand.create_candidate(res, ESTIMATOR, base)
    for bad, key in ((dict(base, source_url=None, description="a"), "source URL"), (dict(base, retrieved_at=None, description="b"), "retrieval date"),
                     (dict(base, original_price="-3", description="c"), "greater than zero"), (dict(base, source_type="Official statistics", description="d"), "not unit prices"),
                     (dict(base, resource_id="MISSING", description="e"), "does not exist")):
        with pytest.raises(ValidationFailed) as e:
            cand.create_candidate(res, ESTIMATOR, bad)
        assert key in e.value.message
    with pytest.raises(Conflict):
        cand.create_candidate(res, ESTIMATOR, base)                                                # exact duplicate
    with pytest.raises(PermissionDenied):
        cand.create_candidate(res, AI, dict(base, description="from ai"))                          # AI cannot author candidates


# ---------------------------------------------------------------------------------------------- queue classification
def test_queue_classes_new_updated_conflicting_expired_requires_verification(res):
    base = dict(source_type="Supplier website", source_url="https://example.invalid/p", retrieved_at=TODAY, original_price="100", currency="SAR", unit="m3",
                resource_id="T-MAT-001", tax_basis="Excluded", delivery_basis="Delivered", price_kind="List price", effective_date=TODAY)
    a = cand.create_candidate(res, ESTIMATOR, dict(base, description="A"))
    assert cand.get(res, a)["queue"] == "New"
    b = cand.create_candidate(res, ESTIMATOR, dict(base, description="B", original_price="130"))
    assert cand.get(res, a)["queue"] == "Conflicting" or cand.recompute(res, a)["queue"] == "Conflicting"
    assert cand.get(res, b)["queue"] == "Conflicting"
    v = cand.create_candidate(res, ESTIMATOR, dict(base, description="V", resource_id="T-MAT-002", unit="kg", tax_basis="Unknown"))
    assert cand.get(res, v)["queue"] == "Requires Verification" and "VAT_UNKNOWN" in cand.get(res, v)["flags"]
    assert cand.get(res, v)["normalisation"]["complete"] is False
    e = cand.create_candidate(res, ESTIMATOR, dict(base, description="E", resource_id="T-MAT-002", unit="kg", valid_until="2020-01-01"))
    assert cand.get(res, e)["queue"] == "Expired"
    rid0 = rates.propose_rate(res, ESTIMATOR, "T-MAT-001", unit_price="90", rate_type="Published Market Price", price_date="2026-09-01", confidence="Medium",
                              evidence="TEST evidence", submit=True)
    rates.approve_rate(res, PROCUREMENT, rid0)
    u = cand.create_candidate(res, ESTIMATOR, dict(base, description="U", original_price="95", effective_date="2026-09-20"))
    assert cand.recompute(res, u)["queue"] in ("Updated", "Conflicting")
    counts = cand.queue_counts(res)
    assert counts["by_queue"].get("Conflicting", 0) >= 2 and sum(counts["by_queue"].values()) >= 4


# ---------------------------------------------------------------------------------------------- normalisation
def test_normalisation_requires_approved_inputs_and_never_changes_the_original(res):
    base = dict(source_type="Supplier website", source_url="https://example.invalid/p", retrieved_at=TODAY, effective_date="2026-09-15", unit="kg",
                resource_id="T-MAT-002", description="Rebar", original_price="2.50", currency="USD", tax_basis="Excluded", price_kind="List price")
    c = cand.get(res, cand.create_candidate(res, ESTIMATOR, base))
    assert c["normalised_price"] is None and "NORMALISATION_INCOMPLETE" in c["flags"] and any("FX" in u for u in c["normalisation"]["unresolved"])
    # the seeded USD/SAR 3.75 is Pending Review (rate_date 2020) -> still not usable
    fx = commercial.propose_fx(res, ESTIMATOR, "USD", "SAR", "3.75", "2026-01-01", "TEST FX source (synthetic)")
    assert cand.recompute(res, c["id"])["normalised_price"] is None
    commercial.approve_fx(res, PROCUREMENT, fx)
    c = cand.recompute(res, c["id"])
    assert c["normalised_price"] == "9.3750" and "x3.75" in c["normalisation_method"]
    assert c["original_price"] == "2.50" and c["original_currency"] == "USD" and c["original_unit"] == "kg"      # original untouched
    # unit conversion (ton -> kg) with the registry's fixed factor
    n = normalise.normalise(res, price="2000", currency="SAR", unit="t", target_unit="kg", tax_basis="Excluded", price_date="2026-09-15")
    assert n.complete and n.price == Decimal("2.0000") and n.steps[0]["step"] == "unit"
    # non-convertible units are NOT guessed
    n = normalise.normalise(res, price="5", currency="SAR", unit="hr", target_unit="m3", tax_basis="Excluded", price_date="2026-09-15")
    assert not n.complete and n.price is None and "not convertible" in n.unresolved[0]
    # VAT-inclusive needs an approved tax category valid on the price date
    n = normalise.normalise(res, price="115", currency="SAR", unit="kg", target_unit="kg", tax_basis="Included", price_date="2026-09-15")
    assert not n.complete and "no APPROVED" in n.unresolved[0]
    approve_vat(res)
    n = normalise.normalise(res, price="115", currency="SAR", unit="kg", target_unit="kg", tax_basis="Included", price_date="2026-09-15")
    assert n.complete and n.price == Decimal("100.0000") and n.steps[0]["step"] == "tax"
    assert not normalise.normalise(res, price="5", currency="SAR", unit="kg", target_unit="kg", tax_basis="Unknown", price_date="2026-09-15").complete
    assert not normalise.normalise(res, price="5", currency="SAR", unit="kg", target_unit="kg", tax_basis="Excluded", price_date=None).complete


# ---------------------------------------------------------------------------------------------- review, publication, versioning
def _cand(res, **kw):
    d = dict(source_type="Supplier website", source_name="Test Supplier", source_url="https://example.invalid/p", retrieved_at=TODAY, effective_date="2026-09-15",
             original_price="310.50", currency="SAR", unit="m3", resource_id="T-MAT-001", description="Ready-mix C30", tax_basis="Excluded",
             delivery_basis="Delivered", price_kind="List price", confidence="Low")
    d.update(kw)
    return cand.create_candidate(res, ESTIMATOR, d)


def test_only_authorised_humans_review_and_nothing_publishes_without_approval(res):
    cid = _cand(res)
    for actor in (AI, VIEWER, ESTIMATOR):                                                        # Estimators propose; they do not approve global rates
        with pytest.raises(PermissionDenied):
            cand.review(res, actor, cid, "approve", "ok")
    with pytest.raises(Conflict):
        cand.publish(res, ESTIMATOR, cid)                                                         # not yet approved for publication
    with pytest.raises(ValidationFailed):
        cand.review(res, PROCUREMENT, cid, "reject", "")                                          # comment required
    r = cand.review(res, PROCUREMENT, cid, "approve", "Matches supplier list dated 15 Sep")
    assert r["status"] == "Approved for Publication" and r["reviewer"] == "proc"
    assert res.scalar("SELECT COUNT(*) FROM resource_rates WHERE unit_price IS NOT NULL") == 0     # still nothing in the rate library
    ap = res.one("SELECT * FROM approvals WHERE entity_type='rate_candidate' ORDER BY id DESC LIMIT 1")
    assert ap["reviewer"] == "proc" and ap["action"] == "approve" and ap["evidence"] == "https://example.invalid/p"


def test_unsafe_candidates_cannot_be_approved_for_publication(res):
    unknown_vat = _cand(res, tax_basis="Unknown", description="vat?")
    with pytest.raises(ValidationFailed) as e:
        cand.review(res, PROCUREMENT, unknown_vat, "approve", "x")
    assert e.value.code == "NOT_PUBLISHABLE" and "VAT" in e.value.message
    promo = _cand(res, price_kind="Promotional", description="promo")
    with pytest.raises(ValidationFailed):
        cand.review(res, PROCUREMENT, promo, "approve", "x")
    unmatched = _cand(res, resource_id=None, description="Ready-mix concrete grade C30 supply")
    c = cand.get(res, unmatched)
    assert c["suggested_resource_id"] == "T-MAT-001" and "RESOURCE_UNMATCHED" in c["flags"]        # a suggestion, never applied automatically
    with pytest.raises(ValidationFailed):
        cand.review(res, PROCUREMENT, unmatched, "approve", "x")
    spec_only = cand.create_candidate(res, ESTIMATOR, dict(source_type="Manufacturer website", source_url="https://example.invalid/spec", retrieved_at=TODAY,
                                                          price_kind="Specification only", description="Anchor bolt M12 stainless", specification="A4-70, 120 mm"))
    sc = cand.get(res, spec_only)
    assert sc["original_price"] is None and sc["normalised_price"] is None and sc["queue"] == "Requires Verification"
    with pytest.raises(ValidationFailed):
        cand.review(res, PROCUREMENT, spec_only, "approve", "x")
    foreign = _cand(res, currency="USD", original_price="80", description="usd no fx")
    with pytest.raises(ValidationFailed) as e:
        cand.review(res, PROCUREMENT, foreign, "approve", "x")
    assert "normalisation incomplete" in e.value.message
    r = cand.review(res, PROCUREMENT, unknown_vat, "needs_verification", "Ask supplier whether VAT is included")
    assert r["status"] == "Needs Verification"
    fixed = cand.edit(res, ESTIMATOR, unknown_vat, tax_basis="Excluded")                          # a human supplies the missing fact
    assert fixed["status"] == "Needs Verification" and fixed["original_price"] == "310.50"
    with pytest.raises(ValidationFailed) as e:
        cand.edit(res, ESTIMATOR, unknown_vat, original_price="1")                                 # originals are immutable
    assert e.value.code == "IMMUTABLE_FIELDS"


def test_publication_creates_new_pending_version_keeps_old_and_does_not_mutate_frozen_estimate(res, tmp_path):
    db = res
    pid = build_mini(db, tmp_path)                                                                 # uses DEMO-MAT-001 (SYNTHETIC)
    vid = estimates.freeze_version(db, ESTIMATOR, pid, "v1")
    frozen = estimates.get_version(db, vid)
    old_rate = db.scalar("SELECT id FROM resource_rates WHERE resource_id='DEMO-MAT-001' AND status='Approved'")
    cid = _cand(db, resource_id="DEMO-MAT-001", original_price="123.45", description="Ready-mix C30 (candidate)")
    cand.review(db, PROCUREMENT, cid, "approve", "supplier list checked")
    new_rate = cand.publish(db, ESTIMATOR, cid)
    r = rates.get_rate(db, new_rate)
    assert r["status"] == "Pending Review" and r["unit_price"] == "123.45" and r["tax_basis"] == "Excluded" and r["version"] > 1
    assert "RC-" in r["evidence"] and "https://example.invalid/p" in r["evidence"] and "supplier list checked" in r["evidence"]
    assert cand.get(db, cid)["status"] == "Published" and cand.get(db, cid)["published_rate_id"] == new_rate
    # a Pending Review rate is NOT eligible: the strict resolver still uses the old approved version
    live = resolve_rate(rates.rates_for(db, "DEMO-MAT-001"), project_id=pid, project_region="KSA-ALL", as_of=date.today())
    assert live.eligibility == ELIGIBLE and live.rate["id"] == old_rate
    pend_only = [x for x in rates.rates_for(db, "DEMO-MAT-001") if x["id"] == new_rate]
    assert resolve_rate(pend_only, project_id=pid, project_region="KSA-ALL", as_of=date.today()).eligibility == UNAPPROVED
    # human approves the rate -> old version superseded (retained), new one approved
    rates.approve_rate(db, PROCUREMENT, new_rate, comments="verified")
    assert rates.get_rate(db, old_rate)["status"] == "Superseded" and rates.get_rate(db, old_rate)["unit_price"] == "100"
    assert rates.get_rate(db, new_rate)["status"] == "Approved"
    # the frozen estimate version is untouched and still reproducible
    after = estimates.get_version(db, vid)
    assert after["totals"] == frozen["totals"] and after["inputs_hash"] == frozen["inputs_hash"] and after["result_json"] == frozen["result_json"]
    assert snapshots.verify_snapshot(db, vid)["ok"] and snapshots.replay_direct_cost(db, vid)["ok"]
    assert snapshots.verify_snapshot(db, vid)["newer_rates_available"]
    # a fresh calculation does use the new rate (that is a *new* estimate state, not a mutation of v1)
    from app.estimators.estimate import compute_estimate
    assert compute_estimate(db, pid).fingerprint != frozen["inputs_hash"]


def test_vat_inclusive_candidate_is_published_vat_exclusive_using_the_approved_tax_rate(res):
    approve_vat(res)
    cid = _cand(res, original_price="115", tax_basis="Included", description="Incl VAT")
    c = cand.get(res, cid)
    assert c["normalised_price"] == "100.0000"
    cand.review(res, PROCUREMENT, cid, "approve", "VAT inclusive confirmed")
    rid = cand.publish(res, ESTIMATOR, cid)
    r = rates.get_rate(res, rid)
    assert Decimal(r["unit_price"]) == Decimal("100") and r["currency"] == "SAR" and r["tax_basis"] == "Excluded"
    assert cand.get(res, cid)["original_price"] == "115"


# ---------------------------------------------------------------------------------------------- web evidence extraction
JSONLD = """<html><head><title>Ready Mix C30 | Example Building Supplies Riyadh</title>
<meta property="article:modified_time" content="{mod}"/>
<script type="application/ld+json">{{"@context":"https://schema.org","@type":"Product","name":"Ready-mix concrete C30","sku":"RM-C30","brand":{{"@type":"Brand","name":"ExampleCo"}},
"description":"Grade C30 ready-mixed concrete, 20 mm aggregate.","offers":{{"@type":"Offer","price":"310.50","priceCurrency":"SAR","priceValidUntil":"{until}",
"availability":"https://schema.org/InStock","priceSpecification":{{"valueAddedTaxIncluded":false,"unitText":"m3"}}}}}}</script></head>
<body><p>Delivered to site in Riyadh. Prices exclude VAT.</p></body></html>"""


def test_web_extraction_json_ld_keeps_source_evidence_and_flags_nothing_it_does_not_know():
    from app.sources import webprice
    html = JSONLD.format(mod=TODAY, until=(date.today() + timedelta(days=30)).isoformat())
    out = webprice.extract_candidates(html, "https://shop.example.sa/p/rm-c30", retrieved_on=date.today())
    assert len(out) == 1
    c = out[0]
    assert (c.price, c.currency, c.unit_text, c.tax_basis, c.sku) == ("310.50", "SAR", "m3", "Excluded", "RM-C30")
    assert c.price_kind == "List price" and c.delivery_basis == "Delivered" and c.geography.startswith("Saudi Arabia") and "STALE" not in c.flags
    assert "Ready-mix" in c.snippet and c.method.startswith("JSON-LD")


def test_web_extraction_edge_cases_spec_only_quote_required_foreign_stale_promo_contradictory():
    from app.sources import webprice
    today = date.today()
    spec = "<html><head><title>Anchor bolt M12</title><meta property='og:description' content='A4-70 stainless, 120 mm. Technical datasheet.'/></head><body>M12 anchor</body></html>"
    c = webprice.extract_candidates(spec, "https://maker.example.com/p", retrieved_on=today)[0]
    assert c.price is None and c.price_kind == "Specification only" and "NO_PRICE" in c.flags
    quote = "<html><head><title>Cable tray</title></head><body>Heavy duty cable tray. Request a quote for pricing. Riyadh warehouse.</body></html>"
    c = webprice.extract_candidates(quote, "https://x.example.sa/p", retrieved_on=today)[0]
    assert c.price is None and c.price_kind == "Quote required"
    usd = '<script type="application/ld+json">{"@type":"Product","name":"Pump","offers":{"price":"250","priceCurrency":"USD"}}</script>'
    c = webprice.extract_candidates(usd, "https://maker.example.com/pump", retrieved_on=today)[0]
    assert c.currency == "USD" and "FOREIGN_CURRENCY" in c.flags and "GEOGRAPHY_UNKNOWN" in c.flags and "VAT_UNKNOWN" in c.flags and "DELIVERY_UNKNOWN" in c.flags
    stale = JSONLD.format(mod="2023-01-01", until="2023-06-01")
    c = webprice.extract_candidates(stale, "https://shop.example.sa/p", retrieved_on=today)[0]
    assert "STALE" in c.flags and "STALE_EXPIRED" in c.flags
    promo = ("<html><body>Saudi store. <del>SAR 120.00</del> <span>Sale SAR 99.00 per m2</span> prices include VAT</body></html>")
    cs = webprice.extract_candidates(promo, "https://store.example.sa/x", retrieved_on=today)
    assert cs and all("TEXT_PATTERN_ONLY" in c.flags for c in cs) and any(c.price == "99.00" and c.tax_basis == "Included" for c in cs)
    dup = ('<script type="application/ld+json">[{"@type":"Product","name":"Pipe","sku":"P1","offers":{"price":"10","priceCurrency":"SAR"}},'
           '{"@type":"Product","name":"Pipe","sku":"P1","offers":{"price":"12","priceCurrency":"SAR"}}]</script>')
    cs = webprice.extract_candidates(dup, "https://s.example.sa/p", retrieved_on=today)
    assert len(cs) == 2 and all("CONTRADICTORY" in c.flags for c in cs)


def test_watchlist_refresh_creates_pending_candidates_only_and_handles_blocked_pages(res):
    html = JSONLD.format(mod=TODAY, until=(date.today() + timedelta(days=30)).isoformat())
    wl = [{"url": "https://shop.example.sa/p/rm-c30", "resource_id": "T-MAT-001", "supplier": "Example Building Supplies"},
          {"url": "https://blocked.example.sa/p"}]
    f = FakeFetcher({"shop.example.sa": lambda u: ok(u, html, "text/html"), "blocked.example.sa": FetchResult("u", "BLOCKED", 403, error="HTTP 403 Forbidden")})
    p = providers.WebPriceWatchlist()
    dry = p.refresh(res, ESTIMATOR, FakeFetcher(), dry_run=True, watchlist=wl)
    assert dry["outcome"] == "DRY_RUN" and res.scalar("SELECT COUNT(*) FROM rate_candidates") == 0
    rep = p.refresh(res, ESTIMATOR, f, dry_run=False, watchlist=wl)
    assert rep["candidates_created"] == 1 and len(rep["blocked"]) == 1 and rep["blocked"][0]["error"] == "HTTP 403 Forbidden"
    c = cand.get(res, rep["candidate_ids"][0])
    assert c["status"] == "Pending Review" and c["source_url"] == "https://shop.example.sa/p/rm-c30" and c["original_price"] == "310.50"
    assert c["source_name"] == "Example Building Supplies" and c["extraction_method"].startswith("JSON-LD") and c["retrieved_at"] and "Ready-mix" in c["evidence_snippet"]
    assert res.scalar("SELECT COUNT(*) FROM resource_rates WHERE unit_price IS NOT NULL") == 0
    rep2 = p.refresh(res, ESTIMATOR, f, dry_run=False, watchlist=wl)                                # same page again: duplicate, no second candidate
    assert rep2["candidates_created"] == 0 and rep2["duplicates"] == 1


def test_generated_rate_dataset_template_can_be_used_for_candidate_import(res, tmp_path):
    """The shipped templates/resource_rate_template.xlsx layout (hidden key row 4, headers row 5) is accepted by `rates import`."""
    import openpyxl
    from app.excel.templates import generate_all
    out = generate_all(tmp_path / "tpl", res)
    f = next(p for p in out if p.name == "resource_rate_template.xlsx")
    wb = openpyxl.load_workbook(f)
    ws = wb["Rate Dataset"]
    hdr = {str(c.value): c.column for c in ws[5] if c.value}
    row = 6
    vals = {"Resource ID": "T-MAT-001", "Unit": "m3", "Unit price": 301.25, "Currency": "SAR", "Price date": "2026-09-10", "Tax basis": "Excluded",
            "Delivery basis": "Delivered", "Evidence / reference": "Supplier letter ref 77 (SYNTHETIC test)"}
    for k, v in vals.items():
        ws.cell(row=row, column=hdr[k], value=v)
    filled = tmp_path / "filled.xlsx"
    wb.save(filled)
    rep = importer.import_rates_file(res, ESTIMATOR, filled, source_name="Template test", dry_run=False)
    assert rep["accepted"] == 1 and rep["rejected"] == 0, rep
    c = cand.get(res, rep["candidate_ids"][0])
    assert c["original_price"] == "301.25" and c["evidence_snippet"].startswith("Supplier letter ref 77") and c["tax_basis"] == "Excluded"


def test_index_data_cannot_be_imported_or_approved_as_unit_prices(res, tmp_path):
    """Index points are reference data only: refused as rate candidates by source, by source type and by shape; observations never create rates."""
    from app.core.errors import ValidationFailed
    from app.rate_updates import importer
    f = tmp_path / "prices.csv"
    f.write_text("description,unit,unit price,currency\nRebar,t,104.89,SAR\n", encoding="utf-8")
    for sid in ("SRC-001", "SRC-002", "SRC-003"):                                                    # GASTAT / KAPSARC index sources
        with pytest.raises(ValidationFailed) as e:
            importer.import_rates_file(res, ESTIMATOR, f, source_id=sid, dry_run=True)
        assert e.value.code == "SOURCE_NOT_PRICE"
    r = importer.import_rates_file(res, ESTIMATOR, f, source_name="Statistics file", source_type="Official statistics", dry_run=True)
    assert r["accepted"] == 0 and "cannot be rate candidates" in r["errors"][0]["errors"][0]
    cci = tmp_path / "cci.csv"
    cci.write_text("time_period,section_and_group,measure,index\n2026-04,Labor Costs,Index,104.89\n", encoding="utf-8")
    with pytest.raises(ValidationFailed) as e:
        importer.import_rates_file(res, ESTIMATOR, cci, dry_run=True)                               # an index table has no price column
    assert e.value.code == "NO_HEADER"
    before = res.scalar("SELECT COUNT(*) FROM resource_rates WHERE unit_price IS NOT NULL")
    from app.rate_updates import index_import
    index_import.import_index_file(res, ESTIMATOR, cci, dataset="test-cci", source_id="SRC-002", source_name="CCI test")
    assert res.scalar("SELECT COUNT(*) FROM index_observations WHERE dataset='test-cci'") == 1
    assert res.scalar("SELECT COUNT(*) FROM resource_rates WHERE unit_price IS NOT NULL") == before and res.scalar("SELECT COUNT(*) FROM rate_candidates") == 0


def test_pilot_rate_template_headers_are_recognised_and_placeholders_accepted_as_nothing(res):
    """The shipped pilot template has headers the importer understands; its placeholder row must not become a candidate."""
    from pathlib import Path
    from app.rate_updates import importer
    r = importer.import_rates_file(res, ESTIMATOR, Path("templates/pilot/rates_template.csv"), source_name="Pilot template", dry_run=True)
    assert r["accepted"] == 0
    assert res.scalar("SELECT COUNT(*) FROM rate_candidates") == 0
