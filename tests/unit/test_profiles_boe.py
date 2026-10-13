"""Client-profile framework, estimate classification and Basis-of-Estimate generation: governed, evidence-bearing, nothing claimed without loaded requirements."""
from __future__ import annotations

from decimal import Decimal as D

import pytest

from app.api.actions import run
from app.core.errors import NotFound, PermissionDenied, ValidationFailed
from app.estimators import profile_checks as pc
from app.estimators.estimate import compute_estimate
from app.services import boe, governed, profiles, wbs
from tests.conftest import AI, COMMERCIAL, ESTIMATOR, VIEWER
from tests.helpers import build_mini


# ------------------------------------------------------------------ pure checks
def test_check_evaluation_and_unknown_facts_are_not_assessable():
    facts = {"a": 5, "b": True, "c": ["X", "Y"], "d": "Complete", "e": None}
    assert pc.evaluate({"fact": "a", "op": "gte", "value": "5"}, facts)[0] == "Met"
    assert pc.evaluate({"fact": "a", "op": "lt", "value": "5"}, facts)[0] == "Not met"
    assert pc.evaluate({"fact": "b", "op": "true"}, facts)[0] == "Met" and pc.evaluate({"fact": "b", "op": "false"}, facts)[0] == "Not met"
    assert pc.evaluate({"fact": "c", "op": "in", "value": "Y"}, facts)[0] == "Met"
    assert pc.evaluate({"fact": "d", "op": "eq", "value": "Complete"}, facts)[0] == "Met" and pc.evaluate({"fact": "d", "op": "ne", "value": "Complete"}, facts)[0] == "Not met"
    assert pc.evaluate({"fact": "e", "op": "gte", "value": "1"}, facts)[0] == "Not assessable"          # None is never silently Met
    assert pc.evaluate({"fact": "zz", "op": "true"}, facts)[0] == "Not assessable"
    assert pc.evaluate({"fact": "d", "op": "gte", "value": "1"}, facts)[0] == "Not assessable"          # non-numeric
    for bad in ({"fact": "nope", "op": "eq", "value": "1"}, {"fact": "a", "op": "weird", "value": "1"}, {"fact": "a", "op": "gte"}):
        with pytest.raises(ValidationFailed):
            pc.validate_check(bad)


# ------------------------------------------------------------------ profile lifecycle
REQS = [
    dict(req_id="R1", category="rule", title="Estimate covers the whole BOQ", source_ref="TEST-DOC-1 rev A cl 3.1", check={"fact": "items_total", "op": "gte", "value": "2"}),
    dict(req_id="R2", category="procedure", title="WBS assigned for at least half of the items", source_ref="TEST-DOC-1 rev A cl 4.2",
         check={"fact": "wbs_assigned_pct", "op": "gte", "value": "50"}),
    dict(req_id="R3", category="form", title="Form 7 completed and signed", source_ref="TEST-FORM-7 rev B"),
    dict(req_id="R4", category="other", title="Optional courtesy item", source_ref="TEST-DOC-2 cl 9", mandatory=0),
]


def _approved_profile(db):
    pid = profiles.create_profile(db, ESTIMATOR, "TESTCLIENT", "Test client profile (SYNTHETIC)", client="Test Client", requirements=REQS,
                                  source_documents=[{"doc_id": "TEST-DOC-1", "title": "Test estimating requirements", "revision": "A", "issuer": "Test Client"}])
    profiles.submit_profile(db, ESTIMATOR, pid)
    profiles.approve_profile(db, COMMERCIAL, pid)
    return pid


def test_profile_validation_and_authority(seeded):
    with pytest.raises(ValidationFailed):                                         # no source reference
        profiles.create_profile(seeded, ESTIMATOR, "X", "x", requirements=[dict(req_id="1", category="rule", title="t", source_ref="")])
    with pytest.raises(ValidationFailed):
        profiles.create_profile(seeded, ESTIMATOR, "X", "x", requirements=[dict(req_id="1", category="rule", title="t", source_ref="s"),
                                                                          dict(req_id="1", category="rule", title="t2", source_ref="s")])
    with pytest.raises(ValidationFailed):
        profiles.create_profile(seeded, ESTIMATOR, "X", "x", requirements=[dict(req_id="1", category="nonsense", title="t", source_ref="s")])
    empty = profiles.create_profile(seeded, ESTIMATOR, "EMPTY", "empty")
    with pytest.raises(ValidationFailed) as e:
        profiles.submit_profile(seeded, ESTIMATOR, empty)
    assert e.value.code == "PROFILE_EMPTY"
    pid = profiles.create_profile(seeded, ESTIMATOR, "T", "t", requirements=REQS)
    profiles.submit_profile(seeded, ESTIMATOR, pid)
    for who in (ESTIMATOR, AI, VIEWER):
        with pytest.raises(PermissionDenied):
            profiles.approve_profile(seeded, who, pid)
    profiles.approve_profile(seeded, COMMERCIAL, pid)
    with pytest.raises(NotFound):
        profiles.link_project(seeded, ESTIMATOR, "PRJ-X", "NOPE")


def test_compliance_matrix_never_claims_beyond_loaded_requirements(seeded, tmp_path):
    proj = build_mini(seeded, tmp_path)
    assert "no compliance statement" in profiles.compliance_report(seeded, proj)["statement"]            # nothing linked
    _approved_profile(seeded)
    profiles.link_project(seeded, ESTIMATOR, proj, "TESTCLIENT")
    r = profiles.compliance_report(seeded, proj)
    st = {x["req_id"]: x["status"] for x in r["requirements"]}
    assert st == {"R1": "Met", "R2": "Not met", "R3": "Not assessed", "R4": "Not assessed"}
    assert r["statement"].startswith("NOT COMPLIANT") and r["mandatory_unmet"] == ["R2", "R3"]
    with pytest.raises(ValidationFailed):                                                                # auto requirement cannot be attested
        profiles.attest(seeded, ESTIMATOR, proj, "R1", "Met", "evidence")
    with pytest.raises(ValidationFailed):                                                                # evidence required
        profiles.attest(seeded, ESTIMATOR, proj, "R3", "Met", "")
    profiles.attest(seeded, ESTIMATOR, proj, "R3", "Met", "Form 7 signed 2026-10-01, ref FRM-7-0042")
    wbs.add_node(seeded, ESTIMATOR, proj, "A1", "Area 1")
    wbs.assign_items(seeded, ESTIMATOR, proj, {"1.1": "A1", "1.2": "A1"})
    r = profiles.compliance_report(seeded, proj)
    assert r["statement"].startswith("All 3 mandatory") and "ONLY the requirements loaded" in r["statement"] and r["counts"]["Not assessed"] == 1       # optional R4 still open
    profiles.attest(seeded, ESTIMATOR, proj, "R3", "Not met", "Form 7 withdrawn, re-issue pending")                                    # latest attestation wins
    assert profiles.compliance_report(seeded, proj)["statement"].startswith("NOT COMPLIANT")


# ------------------------------------------------------------------ classification
SCHEME = {"name": "Test scheme (SYNTHETIC bands)", "source_document": "TEST-CLASS-1 rev A", "classes": [
    {"class_id": "C3", "name": "Low definition", "min_definition_pct": "0", "max_definition_pct": "40", "low_pct": "-30", "high_pct": "50"},
    {"class_id": "C2", "name": "Medium definition", "min_definition_pct": "40", "max_definition_pct": "80", "low_pct": "-15", "high_pct": "25"},
    {"class_id": "C1", "name": "High definition", "min_definition_pct": "80", "max_definition_pct": "100", "low_pct": "-5", "high_pct": "10"}]}


def _approve(db, kind, code, data, **kw):
    i = governed.propose(db, ESTIMATOR, kind, code, data, source="Test source document", submit=True, **kw)
    governed.approve(db, COMMERCIAL, i)
    return i


def test_classification_requires_scheme_and_maturity_and_follows_them(seeded, tmp_path):
    proj = build_mini(seeded, tmp_path)
    r = profiles.classification_report(seeded, proj)
    assert r["class_id"] is None and "No approved classification scheme" in r["notes"][0]
    bad = {**SCHEME, "classes": [SCHEME["classes"][0], {**SCHEME["classes"][1], "min_definition_pct": "30"}]}
    with pytest.raises(ValidationFailed):                                                                # overlapping bands refused
        governed.propose(seeded, ESTIMATOR, "classification_scheme", "S", bad, source="doc")
    _approve(seeded, "classification_scheme", "S", SCHEME)
    assert "No approved maturity input" in profiles.classification_report(seeded, proj)["notes"][0]
    _approve(seeded, "maturity_input", "M", {"definition_pct": "55", "basis": "Drawings IFC for civil, others at 30% (estimator judgement)"}, project_id=proj)
    r = profiles.classification_report(seeded, proj)
    est = compute_estimate(seeded, proj)
    assert r["class_id"] == "C2" and r["range"]["low_pct"] == D(-15) and r["range"]["base"] == est.commercial.total_cost
    assert r["range"]["low"] == est.commercial.total_cost * D("0.85") and r["range"]["high"] == est.commercial.total_cost * D("1.25")
    _approve(seeded, "maturity_input", "M", {"definition_pct": "80", "basis": "re-assessed"}, project_id=proj)      # band edge belongs to the upper band
    assert profiles.classification_report(seeded, proj)["class_id"] == "C1"


# ------------------------------------------------------------------ BoE
def test_boe_contains_facts_flags_synthetic_and_leaves_judgement_blank(seeded, tmp_path):
    proj = build_mini(seeded, tmp_path)
    out = run(seeded, ESTIMATOR, "generate_boe", {"project_id": proj})
    assert out["placeholders"] == 6 and len(out["files"]) == 2
    from pathlib import Path
    d = Path(out["directory"])
    md = (d / out["files"][0]).read_text(encoding="utf-8")
    assert proj in md and "SYNTHETIC DATA" in md and md.count(boe.PLACEHOLDER) == 6
    assert "Derived quantities" in md
    assert "DEMO-CONC" in md and "Synthetic Demonstration Rate" in md                  # assembly + rate type evidence
    from docx import Document
    assert any(boe.PLACEHOLDER in para.text for para in Document(str(d / out["files"][1])).paragraphs)
    assert profiles.project_facts(seeded, proj)["boe_generated"] is True
    with pytest.raises(PermissionDenied):
        run(seeded, VIEWER, "generate_boe", {"project_id": proj})
